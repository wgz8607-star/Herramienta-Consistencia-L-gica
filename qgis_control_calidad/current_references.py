"""Read municipal references from sentencias.r1_r2; never modify their source."""
from contextlib import closing
import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import tempfile

from openpyxl import load_workbook
try:
    import psycopg2
    from psycopg2 import sql
except ImportError:
    psycopg2 = sql = None
from .reference_data import PreparedReference, REQUIRED, column_name, clean

SCHEMA = 'r1_r2'
NAME = re.compile(r'^(\d{5})_(\d{8})_(r[12])$')


def municipal_names(filename):
    if not Path(filename).is_file():
        return {}
    book = load_workbook(filename, read_only=True, data_only=True)
    try:
        return {str(code).strip().zfill(5): (str(department).strip(), str(municipality).strip())
                for code, department, municipality in book.active.iter_rows(min_row=2, max_col=3, values_only=True)
                if code is not None and department is not None and municipality is not None}
    finally:
        book.close()


def catalog_from_names(table_names, municipalities=None):
    municipalities = municipalities or {}
    groups, ignored = {}, []
    for table in table_names:
        match = NAME.fullmatch(table)
        if not match:
            ignored.append(table); continue
        code, stamp, kind = match.groups()
        try:
            date = datetime.strptime(stamp, '%Y%m%d').date()
        except ValueError:
            ignored.append(table); continue
        if code not in groups or stamp > groups[code]['stamp']:
            department, municipality = municipalities.get(code, ('', 'Municipio no registrado en el catálogo local'))
            groups[code] = {'divipola': code, 'stamp': stamp, 'date': date.isoformat(),
                            'department': department, 'municipality': municipality, 'tables': {}}
        if stamp == groups[code]['stamp']:
            groups[code]['tables'][kind.upper()] = table
    return sorted(groups.values(), key=lambda entry: entry['divipola']), ignored


def read_catalog(params, divipola_file):
    if psycopg2 is None:
        raise RuntimeError('Se requiere psycopg2 para consultar las referencias vigentes.')
    with closing(psycopg2.connect(**params, connect_timeout=8)) as connection:
        connection.set_session(readonly=True)
        with connection, connection.cursor() as cursor:
            cursor.execute("SET LOCAL statement_timeout = '15s'")
            cursor.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = %s AND table_type IN ('BASE TABLE', 'VIEW', 'FOREIGN TABLE') ORDER BY table_name", (SCHEMA,))
            tables = [row[0] for row in cursor.fetchall()]
    return catalog_from_names(tables, municipal_names(divipola_file))


def needed_references(rules, report=None):
    texts = [rule[2] for rule in rules]
    if report:
        texts.append(report)
    return {kind for kind in REQUIRED if any('{tabla_' + kind.lower() + '}' in text or '{esquema_' + kind.lower() + '}' in text for text in texts)}


def prepare_current_references(params, selection, kinds):
    """Snapshot required tables consistently, keeping leading zeroes as stored text."""
    if psycopg2 is None:
        raise RuntimeError('Se requiere psycopg2 para leer las referencias vigentes.')
    prepared = {}
    try:
        for kind in kinds:
            table = selection.get('tables', {}).get(kind, '')
            expected = f"{selection.get('divipola', '')}_{selection.get('stamp', '')}_{kind.lower()}"
            if kind not in REQUIRED or not NAME.fullmatch(table) or table != expected:
                raise ValueError(f'{kind} no está disponible para el municipio y fecha seleccionados.')
        if not kinds:
            return prepared
        with closing(psycopg2.connect(**params, connect_timeout=8)) as connection:
            connection.set_session(readonly=True, isolation_level='REPEATABLE READ')
            with connection:
                for kind in sorted(kinds):
                    table = selection['tables'][kind]
                    directory = tempfile.TemporaryDirectory(prefix='qgis_vigente_')
                    try:
                        with connection.cursor() as cursor:
                            cursor.execute("SET LOCAL statement_timeout = '60s'")
                            cursor.execute(sql.SQL('SELECT * FROM {}.{}').format(sql.Identifier(SCHEMA), sql.Identifier(table)))
                            fields = [item.name for item in cursor.description]
                            columns = tuple(column_name(field, index) for index, field in enumerate(fields, 1))
                            if len(columns) != len(set(columns)):
                                raise ValueError(f'{table}: hay columnas duplicadas después de normalizarlas.')
                            missing = REQUIRED[kind] - set(columns)
                            if missing:
                                raise ValueError(f'{table}: faltan columnas requeridas: ' + ', '.join(sorted(missing)))
                            path = Path(directory.name) / 'reference.csv'
                            count, total, squares = 0, 0, 0
                            modulus = 1 << 256
                            with path.open('w', encoding='utf-8', newline='') as output:
                                writer = csv.writer(output, quoting=csv.QUOTE_ALL)
                                while True:
                                    rows = cursor.fetchmany(1000)
                                    if not rows:
                                        break
                                    for row in rows:
                                        values = [clean(value) for value in row]
                                        writer.writerow(values)
                                        value = int.from_bytes(hashlib.sha256(json.dumps(values, ensure_ascii=False, separators=(',', ':')).encode()).digest(), 'big')
                                        total = (total + value) % modulus
                                        squares = (squares + value * value) % modulus
                                        count += 1
                        # A database may return unchanged rows in a different order.
                        fingerprint = hashlib.sha256(json.dumps((columns, count, total, squares)).encode()).hexdigest()
                        prepared[kind] = PreparedReference(kind, SCHEMA + '.' + table, columns, count, fingerprint, path, directory)
                    except Exception:
                        directory.cleanup()
                        raise
        return prepared
    except Exception:
        for reference in prepared.values():
            reference.close()
        raise
