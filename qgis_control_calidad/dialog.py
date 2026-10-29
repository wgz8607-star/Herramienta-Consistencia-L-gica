from __future__ import annotations

import csv
import shutil
import traceback
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from html import escape
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from zoneinfo import ZoneInfo
from configparser import ConfigParser
from pathlib import Path

from qgis.PyQt.QtCore import QObject, QEventLoop, QSize, QThread, QTimer, QUrl, Qt, pyqtSignal
from qgis.PyQt.QtGui import QIcon, QTextCursor, QColor
from qgis.PyQt.QtWidgets import (
    QApplication, QAbstractItemView, QComboBox, QDialog, QFileDialog, QFormLayout,
    QFrame, QGridLayout, QGroupBox, QHBoxLayout, QInputDialog, QLabel, QLineEdit,
    QMessageBox, QPushButton, QProgressBar, QScrollArea, QSplitter,
    QHeaderView, QMenu, QSizePolicy, QStackedWidget, QTableWidget, QTableWidgetItem, QTabWidget, QTextBrowser, QToolButton, QVBoxLayout, QWidget,
)
from qgis.core import QgsApplication, QgsAuthMethodConfig, QgsSettings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .ingestion import CicaIngestionDialog
from .cica_data import install_csv, resolve_csv_tables
from .cica_gdb import install_gdb
from .territorial_scope import SCOPES, install_scope, resolve_scope
from .territorial_filter import TerritorialFilterDialog, TerritorialOverview
from .territorial_statistics import read_filter_statistics
from .templates import validation_format, DOWNLOAD_NAMES
from .report_identifiers import PREDIO_HEADER, is_predio_field, predio_value, rule_sheet_label, rule_sort_key, sql_report_fields
from .reference import ReferenceDialog
from .reference_data import TEMP_TABLES, install_reference
from .current_references import read_catalog, needed_references, prepare_current_references
from .review_imports import ReviewImports
from .review_files import fingerprint_file
from .review import build_cases, digest, save_review, load_review, REVIEW_FIELDS, display, review_statistics, statistics_row, STAT_HEADERS, model_component_number, model_component_sort_key
from .report_exceptions import HEADERS as EXCEPTION_HEADERS, FIELDS as EXCEPTION_FIELDS, DEFAULTS as EXCEPTION_DEFAULTS, add_controls, link_rule_reviews, configure_report_workbook, internal_detail_link, prepare_validation_index

try:
    import psycopg2
except ImportError:
    psycopg2 = None


PLUGIN_DIR = Path(__file__).resolve().parent
SQL_ROOT = PLUGIN_DIR / "resources" / "sql"
DIVIPOLA_FILE = PLUGIN_DIR / "resources" / "data" / "divipola.xlsx"
MODELS = {"Modelo Interno de Levantamiento 1.0": "rules_int_model", "CICA": "rules_cica_db"}
HELP_URL = "https://github.com/wgz8607-star/Herramienta-Consistencia-L-gica/blob/main/qgis_control_calidad/ayuda/GUIA_USO.md"
SNC_CATEGORY = "Cargue SNC"
SNC_ENVIRONMENT = "Validaciones para cargue a SNC"
EXCEPTION_NOTICE = "Solo las excepciones validadas y marcadas por DT descuentan errores."

SNC_REPOSITORY_TABLES = {
    "CICA": "reglas_cargue_snc_cica",
    "Modelo Interno de Levantamiento 1.0": "reglas_cargue_snc_xtf",
}

SQL_REPOSITORIES = {
    "Modelo Interno de Levantamiento 1.0": ("sql_modelo_interno", "sql_modelo_interno_prueba"),
    "CICA": ("sql_cica", "sql_cica_prueba"),
}
SQL_REPOSITORY_TABLES = {
    "Modelo Interno de Levantamiento 1.0": ("reglas_interno_v3_07_09_2026", "reglas_interno_v2_29_07_2024"),
    "CICA": ("reglas_cica_v3_07_09_2026", "reglas_cica_v2_29_07_2024"),
}
POSTGRES_HOST = "172.19.3.225"
REPOSITORY_CONNECTION = {
    "host": POSTGRES_HOST,
    "port": "5432",
    "dbname": "sentencias",
    "user": "admin_ant",
    "password": "ant12726",
}


def published_versions(params, source):
    with closing(psycopg2.connect(**params, connect_timeout=8)) as raw, raw as connection:
        with connection.cursor() as cursor:
            if source["kind"] == "tables":
                rows = []
                for table in source["tables"]:
                    published = "NULL" if source.get("snc") else "max(fecha_actualizacion) OVER (PARTITION BY version)"
                    cursor.execute(
                        f"WITH cabeceras AS ("
                        f"SELECT DISTINCT ON (version) version, descripcion_version, version_publicada, "
                        f"{published} AS publicada FROM public.{table} "
                        "ORDER BY version, orden NULLS LAST, componente, codigo_regla"
                        ") SELECT cabecera.version, COALESCE(cabecera.descripcion_version, ''), "
                        "'PUBLICADA', cabecera.publicada FROM cabeceras AS cabecera "
                        "WHERE cabecera.version_publicada IS TRUE AND EXISTS ("
                        f"SELECT 1 FROM public.{table} AS regla "
                        "WHERE regla.version = cabecera.version AND regla.activa) "
                        "ORDER BY cabecera.publicada DESC NULLS LAST, cabecera.version DESC"
                    )
                    rows.extend(({"table": table, "version": item[0]}, item[0], item[1], item[2], item[3]) for item in cursor.fetchall())
            else:
                cursor.execute(
                    f"SELECT id, codigo, COALESCE(descripcion, ''), estado, fecha_publicacion "
                    f"FROM {source['schema']}.versiones WHERE estado = 'PUBLICADA' ORDER BY fecha_publicacion DESC, codigo DESC"
                )
                rows = cursor.fetchall()
    return rows


@dataclass
class RuleResult:
    rule: str
    category: str
    status: str
    findings: int | str
    detail: str = ""
    rows: list[dict] | None = None
    description: str = ""
    criterion_hash: str = ""
    identity_safe: bool = True


class SqlVersionDialog(QDialog):
    """Busca y permite seleccionar una versión publicada del repositorio SQL."""

    def __init__(self, parent, params, repository_source):
        super().__init__(parent)
        self.params = params
        self.repository_source = repository_source
        self.selected_version = None
        self.setWindowTitle("Seleccionar versión de sentencias SQL")
        self.resize(720, 390)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(f"Versiones publicadas en el repositorio: {repository_source['label']}"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("Buscar por versión o descripción…")
        self.search.textChanged.connect(self.filter_rows)
        layout.addWidget(self.search)
        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["Versión", "Descripción", "Estado", "Publicada"])
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.itemDoubleClicked.connect(lambda _: self.accept_selection())
        layout.addWidget(self.table)
        controls = QHBoxLayout()
        refresh = QPushButton("Actualizar")
        refresh.clicked.connect(self.load_versions)
        choose = QPushButton("Usar esta versión")
        choose.clicked.connect(self.accept_selection)
        cancel = QPushButton("Cancelar")
        cancel.clicked.connect(self.reject)
        controls.addWidget(refresh); controls.addStretch(); controls.addWidget(cancel); controls.addWidget(choose)
        layout.addLayout(controls)
        self.load_versions()

    def load_versions(self):
        self.table.setRowCount(0)
        try:
            rows = published_versions(self.params, self.repository_source)
            for item in rows:
                row = self.table.rowCount()
                self.table.insertRow(row)
                for column, value in enumerate((item[1], item[2], item[3], item[4] or "")):
                    cell = QTableWidgetItem(str(value))
                    if column == 0:
                        cell.setData(Qt.UserRole, item[0])
                    self.table.setItem(row, column, cell)
            self.filter_rows()
        except Exception as error:
            QMessageBox.warning(self, "No fue posible consultar versiones", str(error))

    def filter_rows(self):
        text = self.search.text().strip().lower()
        for row in range(self.table.rowCount()):
            values = " ".join(self.table.item(row, column).text().lower() for column in range(2))
            self.table.setRowHidden(row, bool(text) and text not in values)

    def accept_selection(self):
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.information(self, "Seleccione una versión", "Seleccione una versión publicada para continuar.")
            return
        self.selected_version = {
            "id": self.table.item(row, 0).data(Qt.UserRole),
            "codigo": self.table.item(row, 0).text(),
            "source": self.repository_source,
        }
        self.accept()


class RuleWorker(QObject):
    result = pyqtSignal(object)
    general_report = pyqtSignal(object, str)
    reference_context = pyqtSignal(object)
    progress = pyqtSignal(int, int)
    fatal_error = pyqtSignal(str)
    finished = pyqtSignal()
    message = pyqtSignal(str)

    def __init__(self, params, schema, prefix, r1_table, rules, general_report_path=None, references=None, cica_references=None, cica_gdb=None, current_reference=None, repository_params=None, territorial_scope="all", territorial_prefixes=(), model="CICA"):
        super().__init__()
        self.params = params
        self.schema = schema
        self.prefix = prefix
        self.r1_table = r1_table
        self.current_reference = current_reference
        self.territorial_prefixes = tuple(territorial_prefixes)
        self.validation_model = model
        self.territorial_scope = territorial_scope
        self.scope_views = {}
        self.repository_params = repository_params
        self.references = dict(references or {})
        self.cica_references = dict(cica_references or {})
        self.cica_gdb = cica_gdb
        self.rules = rules
        self.general_report_path = general_report_path

    def resolve_sql(self, sql):
        for kind in ('R1', 'R2'):
            suffix = kind.lower()
            markers = (f'{{esquema_{suffix}}}.{{tabla_{suffix}}}',
                       f'{{esquema_{suffix}}}{{punto}}{{tabla_{suffix}}}',
                       f'{{esquema_{suffix}}}{{tabla_{suffix}}}')
            if any(marker in sql for marker in (f'{{esquema_{suffix}}}', f'{{tabla_{suffix}}}')):
                if kind in self.references:
                    source = f'pg_temp.{TEMP_TABLES[kind]}'
                else:
                    raise ValueError(f'Esta regla requiere {kind}: cargue el archivo temporal o seleccione R1 y R2 Vigente.')
                for marker in markers:
                    sql = sql.replace(marker, source)
                sql = sql.replace(f'{{esquema_{suffix}}}', source.rsplit('.', 1)[0])
                sql = sql.replace(f'{{tabla_{suffix}}}', source.rsplit('.', 1)[1])
        schema_prefix = self.schema.rstrip('.') + '.'
        sql = sql.replace('{esquema}.{punto}', schema_prefix)
        sql = sql.replace('{esquema}{punto}.', schema_prefix)
        sql = sql.replace('{esquema}{punto}', schema_prefix)
        sql = sql.replace('{esquema}.', schema_prefix).replace('{esquema}', schema_prefix).replace('{punto}', self.prefix)
        return resolve_scope(resolve_csv_tables(sql, self.schema, self.cica_references, self.cica_gdb.layers if self.cica_gdb else ()), self.scope_views)

    def run(self):
        owned_references = {}
        try:
            if self.current_reference:
                kinds = needed_references(self.rules, self.general_report_path)
                self.message.emit('Leyendo R1 y R2 Vigente del municipio seleccionado…')
                owned_references = prepare_current_references(self.repository_params, self.current_reference, kinds)
                self.references.update(owned_references)
                self.reference_context.emit({'reference_' + kind.lower(): ref.fingerprint for kind, ref in owned_references.items()})
            with closing(psycopg2.connect(**self.params)) as raw_connection, raw_connection as connection:
                for kind, reference in sorted(self.references.items()):
                    self.message.emit(f'Cargando {kind} en tabla temporal de la sesión…')
                    install_reference(connection, reference)
                if self.cica_gdb:
                    self.message.emit('Preparando geodatabase temporal para validar…')
                    install_gdb(connection, self.cica_gdb, exclude=self.cica_references)
                for table, reference in sorted(self.cica_references.items()):
                    self.message.emit(f'Cargando {table} en tabla temporal de la sesión…')
                    install_csv(connection, reference, self.schema, self.cica_gdb.layers.get(table) if self.cica_gdb else None)
                if self.territorial_scope != 'all' or self.territorial_prefixes:
                    temporary_tables = set(self.cica_references) | (set(self.cica_gdb.layers) if self.cica_gdb else set())
                    self.scope_views = install_scope(connection, self.schema, temporary_tables,
                                                     [TEMP_TABLES[kind] for kind in self.references], self.territorial_scope, self.territorial_prefixes, self.validation_model)
                if self.references or self.cica_references or self.cica_gdb or self.scope_views:
                    # Commit setup once so a failing rule cannot roll back the tables.
                    connection.commit()
                for index, (category, rule_name, raw_sql, description) in enumerate(self.rules, 1):
                    try:
                        sql = self.resolve_sql(raw_sql)
                        with connection.cursor() as cursor:
                            cursor.execute(sql)
                            rows = cursor.fetchall() if cursor.description else []
                            fields = [column.name for column in cursor.description] if cursor.description else []
                        unique_fields = []
                        occurrences = {}
                        for name in fields:
                            occurrences[name] = occurrences.get(name, 0) + 1
                            unique_name = name if occurrences[name] == 1 else f"{name}__{occurrences[name]}"
                            while unique_name in unique_fields:
                                unique_name += '_'
                            unique_fields.append(unique_name)
                        detail_rows = [dict(zip(unique_fields, row)) for row in rows]
                        self.result.emit(RuleResult(rule_name, category, "Correcta", len(rows), rows=detail_rows, description=description,
                                                    criterion_hash=digest(sql), identity_safe=(
                                                        len(fields) == len(set(fields))
                                                        and not re.search(r"\b(?:row_number|rank|dense_rank)\s*\(", sql, re.IGNORECASE)
                                                    )))
                    except Exception as error:
                        connection.rollback()
                        detail = re.sub(r"\s+", " ", str(error)).strip()[:240]
                        self.result.emit(RuleResult(rule_name, category, "Error", "—", detail, description=description))
                    self.progress.emit(index, len(self.rules) + bool(self.general_report_path))
                if self.general_report_path:
                    try:
                        sql = self.resolve_sql(self.general_report_path)
                        with connection.cursor() as cursor:
                            cursor.execute(sql)
                            rows = cursor.fetchall() if cursor.description else []
                            fields = [column.name for column in cursor.description] if cursor.description else []
                        self.general_report.emit([dict(zip(fields, row)) for row in rows], "")
                    except Exception as error:
                        connection.rollback()
                        self.general_report.emit([], re.sub(r"\s+", " ", str(error)).strip()[:240])
                    self.progress.emit(len(self.rules) + 1, len(self.rules) + 1)
        except Exception as error:
            self.fatal_error.emit(str(error))
        finally:
            for reference in owned_references.values():
                reference.close()
            self.finished.emit()


class ControlCalidadDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Validador Consistencia Lógica Alfanumérica")
        self.setWindowFlag(Qt.WindowMinimizeButtonHint, True)
        self.setWindowFlag(Qt.WindowMaximizeButtonHint, True)
        self.setWindowIcon(QIcon(str(PLUGIN_DIR / "icon.svg")))
        self.resize(1200, 820)
        self.results: list[RuleResult] = []
        self.review_cases = []
        self.review_imports = ReviewImports(self.review_cases)
        self.review_import_snapshot = {}
        self.detail_review_rows = {}
        self.detail_import_errors = {}
        self.review_dirty = False
        self.review_snapshot = []
        self.review_context = {}
        self.execution_info = {}
        self.review_file_issues = {}
        self.general_report_rows = []
        self.general_report_error = ""
        self.validation_context = "Ubicación territorial pendiente de identificar."
        self.selected_sql_version = None
        self.repository_rules = []
        self.thread = None
        self.worker = None
        self.category_checks = {}
        self.controls_running = False
        self.exporting = False
        self.export_created_paths = []
        self._repository_password = None
        self.references_ready = False
        self.reference_lookup_running = False
        self.filter_reading = False
        self.temporary_references = {}
        self.cica_references = {}
        self.cica_gdb = None
        self.pending_format = None
        self.completed_format = None
        self._build()
        self.refresh_categories()

    def _build(self):
        root_layout = QVBoxLayout(self)
        root_layout.setContentsMargins(0, 0, 0, 0)
        self.main_tabs = QTabWidget()
        root_layout.addWidget(self.main_tabs)
        validation_page = QWidget()
        self.main_tabs.addTab(validation_page, "Validación")
        layout = QVBoxLayout(validation_page)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(8)
        self.setStyleSheet("QDialog { background: #F5F7FA; font-size: 12px; } QWidget { font-size: 12px; } QLabel { color: #575756; } "
                           "QLineEdit, QComboBox { color: #575756; background: white; border: 1px solid #3FA4F5; border-radius: 4px; padding: 5px; } "
                           "QComboBox QAbstractItemView { color: #575756; background: white; selection-color: white; selection-background-color: #068CBB; border: 1px solid #3FA4F5; } "
                           "QPushButton, QGroupBox QPushButton { background: #068CBB; color: white; border: none; border-radius: 5px; padding: 6px 10px; font-weight: 600; } "
                           "QPushButton:hover { background: #0468A3; } "
                           "QProgressBar { border: 1px solid #3FA4F5; border-radius: 5px; text-align: center; background: white; } "
                           "QProgressBar::chunk { background: #00C7CB; border-radius: 4px; }")

        def style_section(widget, color):
            widget.setStyleSheet(
                "QGroupBox { background: white; border: none; border-top: 1px solid #DCE3EB; "
                "margin-top: 18px; padding: 8px; color: #075680; } "
                "QGroupBox::title { subcontrol-origin: margin; left: 8px; padding: 0 5px; font-size: 15px; font-weight: 700; }"
            )

        heading = QLabel("Validador Consistencia Lógica Alfanumérica")
        heading.setStyleSheet("font-size: 22px; font-weight: 700; color: #0468A3;")
        heading_row = QHBoxLayout()
        heading_row.addWidget(heading, 1)
        self.help_button = QPushButton("Ayuda")
        self.help_button.setToolTip("Abrir guía de uso del plugin")
        self.help_button.clicked.connect(self.open_help)
        heading_row.addWidget(self.help_button)
        layout.addLayout(heading_row)
        subtitle = QLabel("Modelo Interno de Levantamiento 1.0 y CICA · Validaciones alfanuméricas")
        subtitle.setStyleSheet("color: #52666a;")
        layout.addWidget(subtitle)

        splitter = QSplitter(Qt.Horizontal)
        self.main_splitter = splitter
        splitter.setChildrenCollapsible(False)
        layout.addWidget(splitter, 1)
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 4, 0)
        left_layout.setSpacing(6)
        connection = QGroupBox("1. Conexión PostgreSQL / PostGIS")
        style_section(connection, "#575756")
        connection.setStyleSheet(connection.styleSheet() +
            " QWidget { font-size: 12px; } QGroupBox { padding: 6px; } QLineEdit, QComboBox { padding: 3px; } "
            "QPushButton { padding: 4px 6px; }")
        form = QFormLayout(connection)
        form.setContentsMargins(6, 16, 6, 6)
        form.setSpacing(5)
        form.setRowWrapPolicy(QFormLayout.WrapAllRows)
        form.setFormAlignment(Qt.AlignTop)
        self.model = QComboBox()
        self.model.addItems(MODELS)
        self.model.currentTextChanged.connect(self.reset_model_selection)
        self.model.activated.connect(self.reset_model_selection)
        self.connection_mode = QComboBox()
        self.connection_mode.addItems(["Conexión guardada en QGIS", "Conexión manual a PostgreSQL"])
        self.connection_mode.currentIndexChanged.connect(self.change_connection_mode)
        self.connection_stack = QStackedWidget()
        self.connection_stack.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        self.qgis_connection = QComboBox()
        self.qgis_connection.currentTextChanged.connect(self.describe_qgis_connection)
        qgis_page = QWidget()
        qgis_layout = QFormLayout(qgis_page)
        qgis_layout.setContentsMargins(0, 0, 0, 0)
        qgis_layout.setRowWrapPolicy(QFormLayout.WrapAllRows)
        qgis_row = QHBoxLayout()
        refresh_qgis = QPushButton("Actualizar")
        refresh_qgis.clicked.connect(self.load_qgis_connections)
        qgis_row.addWidget(refresh_qgis)
        qgis_row.addWidget(self.qgis_connection, 1)
        qgis_layout.addRow("Conexión de QGIS", qgis_row)
        self.qgis_description = QLabel("Seleccione una conexión guardada en QGIS.")
        self.qgis_description.setWordWrap(True)
        self.qgis_description.setStyleSheet("color: #52666a;")
        qgis_layout.addRow("Base de datos", self.qgis_description)
        self.connection_stack.addWidget(qgis_page)
        manual_page = QWidget()
        manual_form = QFormLayout(manual_page)
        manual_form.setContentsMargins(0, 0, 0, 0)
        self.host = QLineEdit(POSTGRES_HOST)
        self.host.setReadOnly(True)
        self.host.setToolTip("Dirección de PostgreSQL preestablecida; no se puede modificar.")
        self.port = QLineEdit("5432")
        self.database = QLineEdit()
        self.user = QLineEdit()
        self.password = QLineEdit()
        self.password.setEchoMode(QLineEdit.Password)
        manual_form.addRow("Servidor", self.host)
        manual_form.addRow("Puerto", self.port)
        manual_form.addRow("Base de datos", self.database)
        manual_form.addRow("Usuario", self.user)
        manual_form.addRow("Contraseña", self.password)
        self.connection_stack.addWidget(manual_page)
        self.schema = QComboBox()
        self.schema.setEditable(False)
        self.schema.currentTextChanged.connect(self.clear_territorial_prefixes)
        form.addRow("Modelo", self.model)
        self.scope_selection = 'all'
        self.scope_prefixes = ()
        self.advanced_filter_summary = QLabel('Todos · Sin filtro de manzana/vereda')
        self.advanced_filter_summary.setWordWrap(True)
        self.advanced_filter_summary.setStyleSheet('color: #344054; background: white; border: 1px solid #DCE4EC; border-radius: 7px; padding: 10px; font-weight: 600;')
        self.environment = QComboBox()
        self.environment.addItems(["Seleccione un ambiente…", "Validaciones del modelo", SNC_ENVIRONMENT])
        self.environment_label = QLabel("Ambiente")
        self.environment.currentTextChanged.connect(self.model_changed)
        form.addRow(self.environment_label, self.environment)
        self.sql_version_label = QLabel("Seleccione una versión de SQL")
        choose_sql_version = QPushButton("Buscar versión de sentencias…")
        self.choose_sql_version = choose_sql_version
        choose_sql_version.setEnabled(False)
        choose_sql_version.clicked.connect(self.open_sql_version_dialog)
        version_row = QVBoxLayout()
        version_row.setSpacing(4)
        version_row.addWidget(self.sql_version_label, 1)
        version_row.addWidget(choose_sql_version)
        form.addRow("Versión de SQL", version_row)
        self.ingestion_label = QLabel("Archivos CICA")
        self.ingestion_button = QPushButton("Cargue archivos CICA…")
        self.ingestion_button.clicked.connect(self.open_ingestion)
        form.addRow(self.ingestion_label, self.ingestion_button)
        form.addRow("Forma de conexión", self.connection_mode)
        form.addRow("", self.connection_stack)
        test = QPushButton("Listar Esquemas")
        test.clicked.connect(self.test_connection)
        self.schema_controls = QWidget()
        schema_form = QFormLayout(self.schema_controls)
        schema_form.setContentsMargins(0, 0, 0, 0)
        schema_form.setSpacing(5)
        schema_form.setRowWrapPolicy(QFormLayout.WrapAllRows)
        schema_form.addRow(test)
        schema_form.addRow("Esquema a validar", self.schema)
        qgis_layout.addRow(self.schema_controls)
        left_layout.addWidget(connection)
        left_layout.addStretch()
        for combo in left.findChildren(QComboBox):
            combo.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
            combo.setMinimumContentsLength(10)
            combo.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
            combo.setMinimumWidth(100)
        for label in left.findChildren(QLabel):
            label.setWordWrap(True)
        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setFrameShape(QFrame.NoFrame)
        left_scroll.setMinimumWidth(250)
        left_scroll.setWidget(left)
        self.left_tabs = QTabWidget()
        self.left_tabs.setMinimumWidth(250)
        self.left_tabs.setStyleSheet('QTabBar::tab { padding: 6px 10px; min-height: 24px; }')
        self.left_tabs.addTab(left_scroll, "Conexión")
        left_panel = QWidget()
        left_panel_layout = QVBoxLayout(left_panel)
        left_panel_layout.setContentsMargins(0, 0, 0, 0)
        left_panel_layout.setSpacing(6)
        left_panel_layout.addWidget(self.left_tabs, 1)
        splitter.addWidget(left_panel)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        group = QWidget()
        group.setObjectName('controlSelection')
        group.setStyleSheet('QWidget#controlSelection { background: white; border: 1px solid #DCE3EB; border-radius: 6px; }')
        group_layout = QHBoxLayout(group)
        group_layout.setContentsMargins(8, 6, 8, 6)
        group_layout.setSpacing(8)
        control_title = QLabel("Controles a ejecutar")
        control_title.setStyleSheet('color: #0468A3; font-weight: 700; font-size: 12px;')
        group_layout.addWidget(control_title)
        self.category_menu = QMenu(self)
        self.category_menu.setToolTipsVisible(True)
        self.category_button = QToolButton()
        self.category_button.setText("Seleccionar grupos")
        self.category_button.setStyleSheet('QToolButton { background: #068CBB; color: white; border: none; border-radius: 4px; padding: 6px 14px; font-size: 12px; font-weight: 600; }')
        self.category_button.setPopupMode(QToolButton.InstantPopup)
        self.category_button.setMenu(self.category_menu)
        group_layout.addWidget(self.category_button)
        self.category_summary = QLabel()
        self.category_summary.setStyleSheet('color: #575756; font-size: 12px;')
        group_layout.addWidget(self.category_summary, 1)
        right_layout.addWidget(group)

        result_box = QGroupBox("Resultados")
        style_section(result_box, "#3FA4F5")
        result_layout = QVBoxLayout(result_box)
        self.summary = QLabel("Aún no se han ejecutado controles.")
        self.summary.setWordWrap(True)
        self.location_summary = QLabel("Ubicación territorial: se identificará al ejecutar los controles.")
        self.location_summary.setWordWrap(True)
        self.location_summary.setStyleSheet("color: #52666a; font-size: 12px; padding: 3px;")
        result_layout.addWidget(self.location_summary)
        metrics_row = QHBoxLayout()
        self.kpi_values = {}
        for field, title in (('findings', 'Hallazgos'), ('dt', 'Excepciones DT'), ('pending', 'Errores pendientes')):
            card = QFrame()
            card.setStyleSheet('QFrame { background: #FAFCFE; border: 1px solid #DCE3EB; border-radius: 5px; } QLabel { border: none; }')
            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(10, 5, 10, 5)
            card_layout.addWidget(QLabel(title))
            value = QLabel('0')
            value.setStyleSheet('font-size: 22px; color: #075680; font-weight: 700; border: none;')
            card_layout.addWidget(value)
            self.kpi_values[field] = value
            metrics_row.addWidget(card)
        result_layout.addLayout(metrics_row)
        self.table = QTableWidget(0, len(STAT_HEADERS))
        self.table.setHorizontalHeaderLabels([
            header.replace('Reglas ejecutadas', 'Reglas\nejecutadas').replace('Errores técnicos', 'Errores\ntécnicos').replace('Excepciones ', 'Excepciones\n')
            for header in STAT_HEADERS
        ])
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.table.verticalHeader().setDefaultSectionSize(24)
        self.table.setStyleSheet("QTableWidget { font-size: 12px; background: white; gridline-color: #D7EAF7; border: 1px solid #3FA4F5; } "
                                 "QHeaderView::section { background: #068CBB; color: white; border: none; padding: 4px; font-size: 12px; font-weight: 700; } "
                                 "QTableWidget::item:selected { background: #BFEAF0; color: #575756; }")
        self.table.setMinimumHeight(215)
        result_layout.addWidget(self.table)
        self.summary.setStyleSheet('font-size: 11px; color: #52666a;')
        result_layout.addWidget(self.summary)
        detail_group = QGroupBox("Gestión de Excepciones")
        style_section(detail_group, "#3FA4F5")
        detail_group_layout = QVBoxLayout(detail_group)
        detail_note = QLabel(EXCEPTION_NOTICE)
        detail_note.setToolTip(EXCEPTION_NOTICE)
        detail_note.setWordWrap(True)
        detail_group_layout.addWidget(detail_note)
        self.detail_review_table = QTableWidget(0, 5)
        self.detail_review_table.setHorizontalHeaderLabels(['Detalle', 'Archivo', 'Última carga', 'Estado', 'Acción'])
        self.detail_review_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.detail_review_table.setSelectionMode(QAbstractItemView.NoSelection)
        self.detail_review_table.verticalHeader().setVisible(False)
        self.detail_review_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.detail_review_table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeToContents)
        self.detail_review_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        self.detail_review_table.setMinimumHeight(195)
        self.detail_review_table.setStyleSheet('QTableWidget { background: white; font-size: 12px; } QLabel { font-size: 12px; } QPushButton { font-size: 12px; padding: 4px 6px; } QHeaderView::section { background: #068CBB; color: white; padding: 5px; font-size: 12px; }')
        self.detail_filename_labels = {}
        detail_group_layout.addWidget(self.detail_review_table, 1)
        self.review_file_alert = QLabel()
        self.review_file_alert.setWordWrap(True)
        self.review_file_alert.setStyleSheet('color: #955900; background: #FFF4DA; padding: 6px; border-radius: 4px;')
        self.review_file_alert.hide()
        detail_group_layout.addWidget(self.review_file_alert)
        self.incidents_toggle = QToolButton()
        self.incidents_toggle.setText('▸ Ver incidencias de la última carga')
        self.incidents_toggle.setCheckable(True)
        self.incidents_browser = QTextBrowser()
        self.incidents_browser.setMaximumHeight(150)
        self.incidents_browser.hide()
        self.incidents_toggle.toggled.connect(self.toggle_incidents)
        detail_group_layout.addWidget(self.incidents_toggle)
        detail_group_layout.addWidget(self.incidents_browser)
        self.format_download = QPushButton('Descargar formato homologado V2')
        self.format_download.setEnabled(False)
        self.format_download.setStyleSheet('QPushButton { background: transparent; color: #0468A3; text-decoration: underline; text-align: left; padding: 4px; } QPushButton:disabled { color: #8B929A; }')
        self.format_download.setToolTip('Se habilita al finalizar una validación del modelo CICA V2 o Modelo interno V2.')
        self.format_download.clicked.connect(self.download_validation_format)
        detail_group_layout.addWidget(self.format_download)
        self.last_execution_label = QLabel('Sin ejecución')
        self.last_execution_label.setStyleSheet('font-size: 11px; color: #52666a;')
        detail_group_layout.addWidget(self.last_execution_label)
        right_layout.addWidget(result_box, 2)
        right_layout.addWidget(detail_group, 3)
        right_scroll = QScrollArea()
        right_scroll.setWidgetResizable(True)
        right_scroll.setFrameShape(QFrame.NoFrame)
        right_scroll.setWidget(right)
        splitter.addWidget(right_scroll)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        splitter.setSizes([300, 740])

        controls = QHBoxLayout()
        self.progress = QProgressBar(self)
        self.progress.setVisible(False)
        self.progress.setMinimumWidth(230)
        controls.addWidget(self.progress)
        self.status = QLabel("Configure la conexión y seleccione los controles.")
        self.status.setWordWrap(True)
        controls.addWidget(self.status, 1)
        self.export_button = QPushButton("Generar reporte detallado")
        self.export_button.clicked.connect(self.export_excel)
        self.export_button.setEnabled(False)
        self.run_button = QPushButton("Ejecutar validación")
        self.run_button.setEnabled(False)
        self.run_button.setStyleSheet("background: #0468A3; color: white; font-weight: 700; padding: 8px 14px;")
        self.run_button.clicked.connect(self.run_rules)
        controls.addWidget(self.export_button)
        left_panel_layout.addWidget(self.run_button)
        layout.addLayout(controls)
        references = QGroupBox("Referencias R1 y R2")
        references.setStyleSheet('QWidget { font-size: 12px; }')
        reference_layout = QVBoxLayout(references)
        reference_layout.setContentsMargins(8, 14, 8, 6)
        self.reference_mode = QComboBox()
        self.reference_mode.addItems(["Archivos R1 y R2 cargados por usuario", "R1 y R2 Vigente"])
        self.reference_mode.currentIndexChanged.connect(self.reference_mode_changed)
        self.reference_mode.setSizeAdjustPolicy(QComboBox.AdjustToMinimumContentsLengthWithIcon)
        self.reference_mode.setMinimumContentsLength(22)
        self.reference_mode.setMinimumContentsLength(12)
        self.reference_mode.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        reference_layout.addWidget(self.reference_mode)
        self.reference_pages = QStackedWidget()
        temporary_page = QWidget()
        temporary_layout = QVBoxLayout(temporary_page)
        temporary_layout.setContentsMargins(0, 0, 0, 0)
        self.load_reference_button = QPushButton("Cargar R1 / R2…")
        self.load_reference_button.clicked.connect(self.open_references)
        self.reference_status = QLabel("Sin referencias temporales")
        self.reference_status.setWordWrap(True)
        temporary_layout.addWidget(self.load_reference_button)
        temporary_layout.addWidget(self.reference_status, 1)
        self.reference_pages.addWidget(temporary_page)
        current_page = QWidget()
        current_layout = QVBoxLayout(current_page)
        current_layout.setContentsMargins(0, 0, 0, 0)
        current_layout.addWidget(QLabel('R1 y R2 Vigente · sentencias.r1_r2'))
        self.municipality = QComboBox()
        self.municipality.setEditable(True)
        self.municipality.setInsertPolicy(QComboBox.NoInsert)
        self.municipality.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        self.municipality.setMinimumContentsLength(12)
        self.municipality.completer().setCaseSensitivity(Qt.CaseInsensitive)
        self.municipality.completer().setFilterMode(Qt.MatchContains)
        self.municipality.lineEdit().setPlaceholderText('Seleccione o busque DIVIPOLA / municipio')
        self.municipality.currentTextChanged.connect(self.show_current_reference)
        current_layout.addWidget(self.municipality)
        self.reference_date = QLabel('Fecha vigente: sin selección')
        self.current_reference_status = QLabel('Consulte los municipios disponibles.')
        self.current_reference_status.setWordWrap(True)
        current_layout.addWidget(self.reference_date)
        current_layout.addWidget(self.current_reference_status)
        self.refresh_references = QPushButton('Consultar R1 y R2 Vigente')
        self.refresh_references.clicked.connect(self.load_current_references)
        current_layout.addWidget(self.refresh_references)
        self.reference_pages.addWidget(current_page)
        self.schema.currentTextChanged.connect(self.select_reference_municipality)
        self.reference_pages.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Maximum)
        reference_layout.addWidget(self.reference_pages)
        reference_layout.addStretch()
        reference_scroll = QScrollArea()
        reference_scroll.setWidgetResizable(True)
        reference_scroll.setFrameShape(QFrame.NoFrame)
        reference_scroll.setWidget(references)
        self.left_tabs.addTab(reference_scroll, "Referencias R1/R2")
        filter_page = QWidget()
        filter_layout = QVBoxLayout(filter_page)
        filter_layout.addWidget(self.advanced_filter_summary)
        self.filter_statistics = None
        self.filter_overview = TerritorialOverview(filter_page)
        filter_layout.addWidget(self.filter_overview)
        self.open_filter_window_button = QPushButton('Ver estadísticas y seleccionar filtros…')
        self.open_filter_window_button.clicked.connect(self.open_advanced_filter_window)
        self.open_filter_window_button.setStyleSheet('QPushButton { background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #217346, stop:1 #0066B3); color: white; border: none; border-radius: 7px; padding: 11px 12px; font-weight: 700; } QPushButton:hover { background: #0066B3; }')
        filter_layout.addWidget(self.open_filter_window_button)
        filter_layout.addStretch()
        filter_scroll = QScrollArea()
        filter_scroll.setWidgetResizable(True)
        filter_scroll.setFrameShape(QFrame.NoFrame)
        filter_scroll.setWidget(filter_page)
        self.advanced_filter_tab_index = self.left_tabs.addTab(filter_scroll, 'Estadística y filtro avanzado')
        self.left_tabs.setTabEnabled(self.advanced_filter_tab_index, False)
        self.filter_source_signature = None
        self.filter_needs_refresh = True
        self.left_tabs.currentChanged.connect(self.advanced_filter_tab_changed)
        tab_bar = self.left_tabs.tabBar()
        tab_font = tab_bar.font()
        tab_font.setPointSize(9)
        tab_bar.setFont(tab_font)
        tab_bar.setElideMode(Qt.ElideNone)
        tab_bar.setUsesScrollButtons(False)
        tab_bar.setExpanding(True)
        metrics = tab_bar.fontMetrics()
        tab_width = sum(metrics.horizontalAdvance(self.left_tabs.tabText(i)) + 28 for i in range(self.left_tabs.count()))
        tab_bar.setMinimumWidth(tab_width)
        tab_bar.setMinimumHeight(metrics.height() + 20)
        self.left_tabs.setMinimumWidth(max(250, tab_width + 4))
        help_page = QWidget()
        help_layout = QVBoxLayout(help_page)
        search_row = QHBoxLayout()
        self.help_search = QLineEdit()
        self.help_search.setPlaceholderText("Buscar en la ayuda…")
        help_search_button = QPushButton("Buscar siguiente")
        self.help_search.returnPressed.connect(self.find_help_text)
        help_search_button.clicked.connect(self.find_help_text)
        search_row.addWidget(self.help_search, 1)
        search_row.addWidget(help_search_button)
        help_layout.addLayout(search_row)
        self.help_browser = QTextBrowser()
        self.help_browser.setOpenExternalLinks(False)
        self.help_browser.setOpenLinks(False)
        self.help_browser.anchorClicked.connect(self.navigate_help)
        help_layout.addWidget(self.help_browser, 1)
        self.main_tabs.addTab(help_page, "Ayuda")
        self.load_help_content()
        self.load_qgis_connections()
        self.update_ingestion_button()
        self.resize_connection_stack()
        self.file_watch_timer = QTimer(self)
        self.file_watch_timer.setInterval(2000)
        self.file_watch_timer.timeout.connect(self.poll_review_files)
        self.file_watch_timer.start()

    def load_help_content(self):
        guide = PLUGIN_DIR / 'ayuda' / 'GUIA_USO.md'
        try:
            text = guide.read_text(encoding='utf-8')
        except OSError:
            self.help_browser.setPlainText('No se pudo cargar la guía incluida en el plugin. Reinstale el ZIP completo para recuperar la ayuda local.')
            return
        self.help_source = text
        self.help_browser.document().setBaseUrl(QUrl.fromLocalFile(str(guide.parent) + '/'))
        if hasattr(self.help_browser, 'setMarkdown'):
            self.help_browser.setMarkdown(text)
        else:
            self.help_browser.setPlainText(text)
        self.help_browser.moveCursor(QTextCursor.Start)

    def open_help(self):
        self.main_tabs.setCurrentIndex(1)
        self.help_search.setFocus()

    def find_help_text(self):
        text = self.help_search.text().strip()
        if text and not self.help_browser.find(text):
            self.help_browser.moveCursor(QTextCursor.Start)
            self.help_browser.find(text)

    def navigate_help(self, url):
        fragment = url.fragment()
        if not fragment:
            return  # Never launch an external browser from embedded help.
        # Resolve the guide's table of contents without network access.
        for line in getattr(self, 'help_source', '').splitlines():
            if line.startswith('#'):
                title = line.lstrip('#').strip()
                slug = re.sub(r'[^\w\s-]', '', title.lower()).replace(' ', '-')
                if slug == fragment:
                    self.help_browser.moveCursor(QTextCursor.Start)
                    while self.help_browser.find(title):
                        block = self.help_browser.textCursor().block()
                        if block.blockFormat().headingLevel() > 0:
                            break
                    self.help_browser.ensureCursorVisible()
                    return

    def resize_connection_stack(self):
        self.update_connection_height()
        # Reparented fields receive style/font metrics in the next Qt layout pass.
        QTimer.singleShot(0, self.update_connection_height)

    def update_connection_height(self):
        page = self.connection_stack.currentWidget()
        if page is not None:
            page.ensurePolished()
            page.layout().activate()
            required = max(page.sizeHint().height(), page.layout().minimumSize().height(),
                           page.layout().totalHeightForWidth(max(1, page.width())))
            self.connection_stack.setFixedHeight(required + 4)

    def update_ingestion_button(self, *_):
        supports_snc = self.model.currentText() in SNC_REPOSITORY_TABLES
        self.environment.setVisible(supports_snc)
        self.environment_label.setVisible(supports_snc)
        if not supports_snc:
            self.environment.blockSignals(True)
            self.environment.setCurrentIndex(0)
            self.environment.blockSignals(False)
        visible = self.model.currentText() == "CICA"
        self.ingestion_label.setVisible(visible)
        self.ingestion_button.setVisible(visible)

    def reset_model_selection(self, *_):
        self.format_download.setEnabled(False)
        self.environment.blockSignals(True)
        self.environment.setCurrentIndex(0)
        self.environment.blockSignals(False)
        self.scope_selection = "all"
        self.clear_territorial_prefixes()
        self.update_ingestion_button()
        self.choose_sql_version.setEnabled(False)
        self.clear_repository_rules()

    def model_changed(self, *_):
        self.format_download.setEnabled(False)
        self.clear_repository_rules()
        snc = self.environment.currentText() == SNC_ENVIRONMENT
        self.choose_sql_version.setEnabled(self.environment.currentIndex() > 0 and not snc)
        self.choose_sql_version.setToolTip('Cargue SNC usa automáticamente su única versión publicada.' if snc else '')
        if self.environment.currentIndex() > 0:
            self.open_sql_version_dialog()

    def repository_connection_params(self):
        params = dict(REPOSITORY_CONNECTION)
        if params.get("password"):
            return params
        if self._repository_password is None:
            password, accepted = QInputDialog.getText(
                self, "Acceso al repositorio de sentencias",
                "Contraseña del repositorio de sentencias (solo para esta sesión):",
                QLineEdit.Password,
            )
            if not accepted or not password:
                self.status.setText("Ejecución bloqueada: falta autenticar el acceso al repositorio de sentencias.")
                return None
            self._repository_password = password
        params["password"] = self._repository_password
        return params

    def repository_source(self, params):
        if self.environment.currentText() == SNC_ENVIRONMENT:
            table = SNC_REPOSITORY_TABLES[self.model.currentText()]
            with psycopg2.connect(**params, connect_timeout=8) as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT to_regclass(%s)", (f"public.{table}",))
                    if not cursor.fetchone()[0]:
                        raise ValueError(f"La tabla public.{table} aún no está disponible para el cargue SNC.")
                    cursor.execute("SELECT has_table_privilege(current_user, %s, 'SELECT')", (f"public.{table}",))
                    if not cursor.fetchone()[0]:
                        raise ValueError(f"Se requiere permiso SELECT sobre public.{table}.")
            return {"kind": "tables", "schema": "public", "tables": [table], "label": f"public.{table}", "snc": True}
        with psycopg2.connect(**params, connect_timeout=8) as connection:
            with connection.cursor() as cursor:
                tables = []
                for table in SQL_REPOSITORY_TABLES[self.model.currentText()]:
                    cursor.execute("SELECT to_regclass(%s), has_table_privilege(current_user, %s, 'SELECT')", (f"public.{table}", f"public.{table}"))
                    exists, readable = cursor.fetchone()
                    if exists and readable:
                        tables.append(table)
                if tables:
                    return {"kind": "tables", "schema": "public", "tables": tables, "label": "public." + ", public.".join(tables)}
        candidates = SQL_REPOSITORIES[self.model.currentText()]
        with psycopg2.connect(**params, connect_timeout=8) as connection:
            with connection.cursor() as cursor:
                for schema in candidates:
                    cursor.execute("SELECT to_regnamespace(%s)", (schema,))
                    if cursor.fetchone()[0]:
                        return {"kind": "schema", "schema": schema, "label": f"{schema}.reglas"}
        raise ValueError("No se encontró el repositorio centralizado de SQL para el modelo seleccionado.")

    def clear_repository_rules(self):
        self.format_download.setEnabled(False)
        self.selected_sql_version = None
        self.repository_rules = []
        self.sql_version_label.setText("Seleccione una versión de SQL")
        self.refresh_categories()
        self.status.setText("Ejecución bloqueada: seleccione una versión con controles activos del repositorio SQL.")

    def update_run_availability(self):
        self.update_advanced_filter_availability()
        selected = {name for name, card in self.category_checks.items() if card.isChecked()}
        ready = bool(self.selected_sql_version) and any(
            item["type"] == "CONTROL" and item["component"] in selected
            and item["sql"] and item["sql"].strip()
            for item in self.repository_rules
        )
        self.run_button.setEnabled(ready and not self.controls_running and not self.exporting and not self.reference_lookup_running and not self.filter_reading)
        self.run_button.setToolTip("" if ready else "Cargue una versión del repositorio con controles activos y seleccione al menos un grupo.")

    def open_sql_version_dialog(self):
        self.clear_repository_rules()
        if self.environment.currentIndex() == 0:
            QMessageBox.information(self, "Seleccione un ambiente", "Seleccione el ambiente antes de buscar una versión de SQL.")
            return
        if psycopg2 is None:
            QMessageBox.warning(self, "Dependencia faltante", "No se puede consultar el repositorio SQL sin psycopg2.")
            return
        try:
            params = self.repository_connection_params()
            if params is None:
                return
            source = self.repository_source(params)
        except Exception as error:
            if getattr(error, "pgcode", None) in ("28P01", "28000"):
                self._repository_password = None
            QMessageBox.information(self, "No fue posible conectar al repositorio", str(error))
            return
        if source.get('snc'):
            self.choose_sql_version.setEnabled(False)
            try:
                versions = published_versions(params, source)
                if len(versions) != 1:
                    message = ('No hay una versión publicada con reglas activas para Cargue SNC.' if not versions else
                               'Cargue SNC debe tener una sola versión publicada. Despublique las versiones adicionales.')
                    raise ValueError(message)
                item = versions[0]
                version = {'id': item[0], 'codigo': item[1], 'source': source}
            except Exception as error:
                self.sql_version_label.setText('Cargue SNC: versión no disponible')
                QMessageBox.warning(self, 'No se cargaron las sentencias SNC', str(error))
                return
        else:
            dialog = SqlVersionDialog(self, params, source)
            if dialog.exec() != QDialog.Accepted or not dialog.selected_version:
                return
            version = dialog.selected_version
        try:
            with psycopg2.connect(**params, connect_timeout=8) as connection:
                with connection.cursor() as cursor:
                    if source["kind"] == "tables":
                        table = version["id"]["table"]
                        cursor.execute(
                            f"SELECT componente, codigo_regla, tipo, sentencia_sql, COALESCE(descripcion_regla, '') FROM public.{table} AS regla "
                            "WHERE version = %s AND activa AND ("
                            f"SELECT cabecera.version_publicada FROM public.{table} AS cabecera "
                            "WHERE cabecera.version = regla.version "
                            "ORDER BY cabecera.orden NULLS LAST, cabecera.componente, cabecera.codigo_regla LIMIT 1"
                            ") IS TRUE ORDER BY componente, orden, codigo_regla",
                            (version["id"]["version"],),
                        )
                    else:
                        cursor.execute(
                            f"SELECT componente, codigo_regla, tipo, sentencia_sql, COALESCE(descripcion, '') FROM {source['schema']}.reglas "
                            "WHERE version_id = %s AND activa ORDER BY componente, orden, codigo_regla",
                            (version["id"],),
                        )
                    self.repository_rules = [
                        {"component": SNC_CATEGORY if source.get("snc") and row[2] == "CONTROL" else row[0],
                         "source_component": row[0], "rule": row[1], "type": row[2], "sql": row[3], "description": row[4]}
                        for row in cursor.fetchall()
                    ]
            self.selected_sql_version = dict(version, source=source)
            self.sql_version_label.setText(f"{version['codigo']} · {len(self.repository_rules)} sentencias")
            self.refresh_categories()
            if any(item["type"] == "CONTROL" and item["sql"] and item["sql"].strip() for item in self.repository_rules):
                self.status.setText(f"Versión SQL seleccionada: {version['codigo']}.")
            else:
                self.status.setText("Ejecución bloqueada: la versión seleccionada no contiene controles activos.")
        except Exception as error:
            self.clear_repository_rules()
            QMessageBox.warning(self, "No fue posible cargar reglas", str(error))

    def update_advanced_filter_summary(self):
        detail = f'{len(self.scope_prefixes)} manzanas/veredas' if self.scope_prefixes else 'Sin filtro de manzana/vereda'
        self.advanced_filter_summary.setText(SCOPES[self.scope_selection] + ' · ' + detail)

    def clear_territorial_prefixes(self, *_):
        self.scope_prefixes = ()
        self.update_advanced_filter_summary()
        self.update_advanced_filter_availability()

    def update_advanced_filter_availability(self):
        if not hasattr(self, 'advanced_filter_tab_index'):
            return
        has_cica = self.model.currentText() == "CICA" and (
            "cica_predio" in self.cica_references or
            bool(self.cica_gdb and "cica_predio" in self.cica_gdb.layers))
        data_ready = bool(self.schema.currentText().strip()) or has_cica
        current = self.selected_current_reference() if self.reference_mode.currentIndex() == 1 else None
        r1_ready = bool(self.temporary_references.get('R1')) if self.reference_mode.currentIndex() == 0 else bool(current and current.get('tables', {}).get('R1'))
        signature = (self.model.currentText(), self.schema.currentText(), self.connection_mode.currentIndex(),
                     self.qgis_connection.currentText(), self.database.text(), self.user.text(),
                     id(self.cica_references.get('cica_predio')), id(self.cica_gdb), self.reference_mode.currentIndex(),
                     id(self.temporary_references.get('R1')), current.get('tables', {}).get('R1') if current else None)
        if signature != self.filter_source_signature:
            if self.filter_source_signature is not None:
                self.scope_prefixes = ()
                self.update_advanced_filter_summary()
            self.filter_source_signature = signature
            self.filter_needs_refresh = True
            self.filter_statistics = None
            self.filter_overview.set_statistics(None)
        ready = data_ready and r1_ready
        enabled = ready and not self.controls_running and not self.exporting and not self.reference_lookup_running
        if not ready and self.left_tabs.currentIndex() == self.advanced_filter_tab_index:
            self.left_tabs.setCurrentIndex(0)
        self.left_tabs.setTabEnabled(self.advanced_filter_tab_index, enabled)
        self.left_tabs.setTabToolTip(self.advanced_filter_tab_index, '' if ready else 'Seleccione la base catastral a validar y cargue R1 temporal o seleccione un municipio con R1 vigente.')

    def advanced_filter_tab_changed(self, index):
        if index != self.advanced_filter_tab_index or self.filter_reading:
            return
        self.update_advanced_filter_availability()
        if self.left_tabs.isTabEnabled(index) and self.filter_needs_refresh:
            self.filter_needs_refresh = False
            QTimer.singleShot(0, self.load_advanced_filter_tab)

    def load_advanced_filter_tab(self):
        self.update_advanced_filter_availability()
        if self.left_tabs.currentIndex() == self.advanced_filter_tab_index and self.left_tabs.isTabEnabled(self.advanced_filter_tab_index):
            try:
                self.read_territorial_prefixes()
            except Exception as error:
                self.filter_statistics = None
                self.filter_overview.set_statistics({'source_error': str(error), 'r1_status': 'No disponible.'})
                self.status.setText('No se pudieron consultar las estadísticas: ' + str(error))

    def apply_advanced_filter(self, scope, prefixes):
        self.scope_selection = scope
        self.scope_prefixes = tuple(prefixes)
        self.update_advanced_filter_summary()
        self.status.setText('Filtro aplicado: ' + self.advanced_filter_summary.text())

    def open_advanced_filter_window(self):
        self.update_advanced_filter_availability()
        if not self.left_tabs.isTabEnabled(self.advanced_filter_tab_index) or self.filter_reading:
            return
        dialog = TerritorialFilterDialog(self.scope_selection, self.scope_prefixes, self.read_territorial_prefixes, self)
        if self.filter_statistics is not None:
            dialog.set_statistics(self.filter_statistics)
        else:
            QTimer.singleShot(0, dialog.load_codes)
        if dialog.exec_() == QDialog.Accepted:
            self.apply_advanced_filter(dialog.scope.currentData(), dialog.selected_prefixes())

    def read_territorial_prefixes(self):
        self.filter_statistics = None
        self.filter_overview.set_statistics(None)
        snapshot = None
        if self.model.currentText() == 'CICA':
            snapshot = self.cica_references.get('cica_predio')
            if snapshot is None and self.cica_gdb:
                snapshot = self.cica_gdb.layers.get('cica_predio')
        params = {} if snapshot else self.connection_params()
        current = self.selected_current_reference() if self.reference_mode.currentIndex() == 1 else None
        repository_params = self.repository_connection_params() if current else None
        self.filter_reading = True
        self.main_tabs.setEnabled(False)
        self.update_run_availability()
        try:
            data = self.report_task(read_filter_statistics, params, self.schema.currentText().strip(), self.model.currentText(), snapshot,
                                    self.temporary_references.get('R1') if self.reference_mode.currentIndex() == 0 else None,
                                    'temporary' if self.reference_mode.currentIndex() == 0 else 'current', current, repository_params)
            self.filter_statistics = data
            self.filter_overview.set_statistics(data)
            return data
        finally:
            self.filter_reading = False
            self.main_tabs.setEnabled(True)
            self.update_run_availability()

    def open_ingestion(self):
        if self.controls_running or self.exporting:
            return
        dialog = CicaIngestionDialog(self, self.cica_references, self.cica_gdb)
        try:
            if dialog.exec() == QDialog.Accepted:
                self.clear_territorial_prefixes()
                for table, reference in dialog.take_prepared().items():
                    old = self.cica_references.get(table)
                    self.cica_references[table] = reference
                    if old:
                        old.close()
                prepared_gdb = dialog.take_geodatabase()
                if prepared_gdb:
                    if self.cica_gdb:
                        self.cica_gdb.close()
                    self.cica_gdb = prepared_gdb
                self.ingestion_button.setText('Cargue archivos CICA…')
                self.ingestion_button.setToolTip(f'{len(self.cica_references)} CSV · ' + (f'{len(self.cica_gdb.layers)} capas/tablas temporales' if self.cica_gdb else 'Sin geodatabase temporal'))
        finally:
            dialog.clear_prepared()
            self.update_advanced_filter_availability()

    def download_validation_format(self):
        if not self.completed_format or not self.format_download.isEnabled():
            return
        label, source = self.completed_format
        destination, _ = QFileDialog.getSaveFileName(self, label, DOWNLOAD_NAMES.get(source.name, source.name), 'Formato Excel (*.xlsx)')
        if not destination:
            return
        if not destination.lower().endswith('.xlsx'):
            destination += '.xlsx'
        try:
            if Path(destination).resolve() == source.resolve():
                raise ValueError('Seleccione una ubicación distinta de los archivos internos del plugin.')
            shutil.copyfile(source, destination)
            self.status.setText('Formato guardado: ' + destination)
        except Exception as error:
            QMessageBox.warning(self, 'No se pudo descargar el formato', str(error))

    def reference_mode_changed(self, index):
        self.reference_pages.setCurrentIndex(index)
        self.update_advanced_filter_availability()
        if index == 1 and not self.municipality.count():
            self.load_current_references()

    def open_references(self):
        if self.controls_running:
            return
        dialog = ReferenceDialog(self)
        if dialog.exec() == QDialog.Accepted:
            prepared = dialog.take_references()
            if prepared:
                self.release_references()
                self.temporary_references = prepared
                self.reference_status.setText(' · '.join(f'{kind}: {ref.filename} ({ref.row_count} registros)' for kind, ref in prepared.items()))
        dialog.clear_prepared()
        self.update_advanced_filter_availability()

    def release_references(self):
        if self.controls_running:
            return
        for reference in self.temporary_references.values():
            reference.close()
        self.temporary_references = {}
        self.reference_status.setText('Sin referencias temporales')
        self.update_advanced_filter_availability()

    def selected_current_reference(self):
        index = self.municipality.currentIndex()
        if index < 0 or self.municipality.currentText() != self.municipality.itemText(index):
            return None
        return self.municipality.currentData()

    def show_current_reference(self, *_):
        selection = self.selected_current_reference()
        if selection is None:
            self.reference_date.setText('Fecha vigente: sin selección')
            self.current_reference_status.setText('Seleccione un municipio del catálogo.')
            self.update_advanced_filter_availability()
            return
        date = datetime.strptime(selection['date'], '%Y-%m-%d').strftime('%d/%m/%Y')
        self.reference_date.setText('Fecha vigente: ' + date)
        tables = selection['tables']
        self.current_reference_status.setText('\n'.join(kind + ': ' + tables.get(kind, 'No disponible para esta fecha') for kind in ('R1', 'R2')))
        self.update_advanced_filter_availability()

    def select_reference_municipality(self, *_):
        codes = set(re.findall(r'(?<!\d)(\d{5})(?!\d)', self.schema.currentText()))
        matches = [index for index in range(self.municipality.count()) if self.municipality.itemData(index) and self.municipality.itemData(index)['divipola'] in codes]
        if len(matches) == 1:
            self.municipality.setCurrentIndex(matches[0])
            self.show_current_reference()

    def load_current_references(self):
        if self.controls_running or self.exporting or self.reference_lookup_running:
            return
        params = self.repository_connection_params()
        if params is None:
            return
        selected = self.selected_current_reference()
        previous = selected['divipola'] if selected else None
        self.reference_lookup_running = True
        self.main_tabs.setEnabled(False)
        self.update_run_availability()
        self.status.setText('Consultando R1 y R2 Vigente en sentencias.r1_r2…')
        try:
            entries, ignored = self.report_task(read_catalog, params, DIVIPOLA_FILE)
            self.municipality.blockSignals(True)
            self.municipality.clear()
            self.municipality.addItem('Seleccione un municipio…', None)
            for entry in entries:
                label = entry['divipola'] + ' · ' + entry['municipality']
                if entry['department']:
                    label += ' · ' + entry['department']
                self.municipality.addItem(label, entry)
            self.municipality.blockSignals(False)
            for index in range(1, self.municipality.count()):
                if self.municipality.itemData(index)['divipola'] == previous:
                    self.municipality.setCurrentIndex(index)
                    break
            self.select_reference_municipality()
            self.show_current_reference()
            self.status.setText(f'{len(entries)} municipios disponibles en R1 y R2 Vigente.' + (f' {len(ignored)} tablas sin nomenclatura reconocida.' if ignored else ''))
        except Exception as error:
            self.municipality.blockSignals(False)
            self.municipality.clear(); self.show_current_reference()
            self.status.setText('No se pudo consultar R1 y R2 Vigente. Revise la conexión al repositorio.')
            QMessageBox.warning(self, 'No se consultaron las referencias vigentes', str(error))
        finally:
            self.reference_lookup_running = False
            self.main_tabs.setEnabled(True)
            self.update_run_availability()

    def change_connection_mode(self, index):
        self.connection_stack.setCurrentIndex(index)
        destination = self.connection_stack.currentWidget()
        if self.schema_controls.parentWidget() is not destination:
            previous_layout = self.schema_controls.parentWidget().layout()
            previous_row, _ = previous_layout.getWidgetPosition(self.schema_controls)
            previous_layout.takeRow(previous_row)
            destination.layout().addRow(self.schema_controls)
        self.resize_connection_stack()
        self.schema.clear()
        self.status.setText("Pulse Listar Esquemas para cargar los esquemas disponibles.")

    def load_qgis_connections(self):
        settings = QgsSettings()
        settings.beginGroup("PostgreSQL/connections")
        names = sorted(settings.childGroups())
        settings.endGroup()
        current = self.qgis_connection.currentText()
        self.qgis_connection.blockSignals(True)
        self.qgis_connection.clear()
        self.qgis_connection.addItems(names)
        if current in names:
            self.qgis_connection.setCurrentText(current)
        self.qgis_connection.blockSignals(False)
        self.describe_qgis_connection(self.qgis_connection.currentText())

    def describe_qgis_connection(self, name):
        if not name:
            self.qgis_description.setText("No se encontraron conexiones PostgreSQL guardadas en este perfil de QGIS.")
            return
        settings = QgsSettings()
        group = f"PostgreSQL/connections/{name}"
        host = settings.value(f"{group}/host", "")
        port = settings.value(f"{group}/port", "5432")
        database = settings.value(f"{group}/database", "")
        self.qgis_description.setText(f"{host}:{port} · {database}")

    def qgis_connection_params(self):
        name = self.qgis_connection.currentText()
        if not name:
            raise ValueError("Seleccione una conexión PostgreSQL guardada en QGIS.")
        settings = QgsSettings()
        group = f"PostgreSQL/connections/{name}"
        params = {
            "host": settings.value(f"{group}/host", ""),
            "port": settings.value(f"{group}/port", "5432"),
            "dbname": settings.value(f"{group}/database", ""),
            "user": settings.value(f"{group}/username", ""),
            "password": settings.value(f"{group}/password", ""),
        }
        authcfg = settings.value(f"{group}/authcfg", "")
        if authcfg:
            auth_config = QgsAuthMethodConfig()
            if QgsApplication.authManager().loadAuthenticationConfig(authcfg, auth_config, True):
                params["user"] = auth_config.config("username") or params["user"]
                params["password"] = auth_config.config("password") or params["password"]
        return params

    def refresh_categories(self, *_):
        if not hasattr(self, "category_menu"):
            return
        self.category_menu.clear()
        self.category_checks.clear()
        if self.repository_rules:
            counts = defaultdict(int)
            for item in self.repository_rules:
                if item["type"] == "CONTROL" and item["sql"] and item["sql"].strip():
                    counts[item["component"]] += 1
            self.populate_category_cards(sorted(counts.items()))
            return
        self.populate_category_cards([])

    def populate_category_cards(self, categories):
        categories = list(categories)
        if not self.selected_sql_version and self.model.currentText() in SNC_REPOSITORY_TABLES and not any(name == SNC_CATEGORY for name, _ in categories):
            categories.append((SNC_CATEGORY, 0))
        self.category_menu.addAction("Seleccionar todos", self.select_all_categories)
        self.category_menu.addAction("Desmarcar todos", self.clear_category_selection)
        self.category_menu.addSeparator()
        for category_name, count in categories:
            text = "+ Cargue SNC" if category_name == SNC_CATEGORY else f"{category_name} · {count} reglas"
            action = self.category_menu.addAction(text)
            action.setCheckable(True)
            optional_snc = category_name == SNC_CATEGORY
            action.setChecked(not optional_snc or (count > 0 and bool(self.selected_sql_version)))
            action.setEnabled(not optional_snc or count > 0)
            if optional_snc and not count:
                action.setToolTip("No hay reglas SNC cargadas para este ambiente.")
            action.toggled.connect(self.update_category_summary)
            self.category_checks[category_name] = action
        self.update_category_summary()

    def clear_category_selection(self):
        for action in self.category_checks.values():
            action.setChecked(False)
        self.update_category_summary()

    def select_all_categories(self):
        for name, action in self.category_checks.items():
            if name != SNC_CATEGORY and action.isEnabled():
                action.setChecked(True)
        self.update_category_summary()

    def update_category_summary(self, *_):
        selected = [name for name, action in self.category_checks.items() if action.isChecked()]
        if self.repository_rules:
            count = sum(1 for item in self.repository_rules if item["type"] == "CONTROL" and item["component"] in selected and item["sql"] and item["sql"].strip())
        else:
            count = sum(len(list((SQL_ROOT / MODELS[self.model.currentText()] / name).glob("*.sql"))) for name in selected)
        self.category_summary.setText(f"{len(selected)} grupos · {count} reglas")
        self.category_summary.setToolTip(', '.join(selected))
        self.category_button.setEnabled(any(action.isEnabled() for action in self.category_checks.values()))
        self.update_run_availability()

    def connection_params(self):
        if self.connection_mode.currentIndex() == 0:
            return self.qgis_connection_params()
        return {"host": POSTGRES_HOST, "port": self.port.text().strip(), "dbname": self.database.text().strip(), "user": self.user.text().strip(), "password": self.password.text()}

    def resolve_validation_context(self, params, schema):
        """Relaciona los cinco primeros dígitos de ILC_Predio con el DIVIPOLA interno."""
        if not DIVIPOLA_FILE.exists():
            return "Ubicación territorial no disponible: falta la referencia DIVIPOLA del plugin."
        if not re.fullmatch(r"[A-Za-z0-9_]+", schema):
            return "Ubicación territorial no disponible: el nombre del esquema no es válido."
        try:
            workbook = load_workbook(DIVIPOLA_FILE, read_only=True, data_only=True)
            sheet = workbook.active
            divipola = {
                str(code).strip().zfill(5): (str(department).strip(), str(municipality).strip())
                for code, department, municipality in sheet.iter_rows(min_row=2, values_only=True)
                if code not in (None, "") and department not in (None, "") and municipality not in (None, "")
            }
            with psycopg2.connect(**params, connect_timeout=8) as connection:
                with connection.cursor() as cursor:
                    cursor.execute(
                        "SELECT column_name FROM information_schema.columns "
                        "WHERE table_schema = %s AND table_name = 'ilc_predio' "
                        "AND column_name IN ('numero_predial', 'numero_predial_nacional')",
                        (schema,),
                    )
                    available = {row[0] for row in cursor.fetchall()}
                    field = "numero_predial" if "numero_predial" in available else "numero_predial_nacional"
                    if field not in available:
                        return "Ubicación territorial no disponible: ILC_Predio no tiene un campo de número predial reconocido."
                    cursor.execute(
                        f"SELECT DISTINCT left({field}::text, 5) FROM {schema}.ilc_predio "
                        f"WHERE {field} IS NOT NULL AND length({field}::text) >= 5 ORDER BY 1"
                    )
                    codes = [str(row[0]).zfill(5) for row in cursor.fetchall()]
            matches = [(code, *divipola[code]) for code in codes if code in divipola]
            if not codes:
                return "Ubicación territorial no identificada: ILC_Predio no contiene números prediales válidos."
            if not matches:
                return "Ubicación territorial no identificada: los códigos DIVIPOLA no están en la referencia interna."
            labels = [f"{department} · {municipality} ({code})" for code, department, municipality in matches]
            visible = labels[:4]
            suffix = f" · y {len(labels) - len(visible)} más" if len(labels) > len(visible) else ""
            return "Ubicación territorial validada: " + "; ".join(visible) + suffix
        except Exception as error:
            return "Ubicación territorial no disponible: " + re.sub(r"\s+", " ", str(error)).strip()[:160]

    def test_connection(self):
        if psycopg2 is None:
            QMessageBox.critical(self, "Dependencia faltante", "El QGIS instalado no incluye psycopg2. Instale psycopg2-binary en el Python de QGIS.")
            return
        try:
            with psycopg2.connect(**self.connection_params(), connect_timeout=8) as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT version()")
                    version = cursor.fetchone()[0]
                    cursor.execute("SELECT nspname FROM pg_namespace WHERE nspname NOT LIKE 'pg_%' AND nspname <> 'information_schema' ORDER BY nspname")
                    schemas = [row[0] for row in cursor.fetchall()]
            self.schema.clear()
            self.schema.addItems(schemas)
            if "public" in schemas:
                self.schema.setCurrentText("public")
            self.status.setText("Conexión exitosa: " + version.split(",")[0])
            self.references_ready = True
            if self.reference_mode.currentIndex() == 1:
                self.load_current_references()
        except Exception as error:
            QMessageBox.critical(self, "No fue posible conectar", str(error))

    def validation_confirmation_metadata(self, params, schema, rules, current_reference, general_report_path=None):
        """Only display selected validation inputs, never credentials or SQL text."""
        rows = [
            ('Modelo', self.model.currentText()), ('Ambiente', self.environment.currentText()),
            ('Versión SQL', (self.selected_sql_version or {}).get('codigo', 'No seleccionada')),
            ('Base catastral a validar', params.get('dbname', '') or 'No disponible'),
            ('Esquema', schema), ('Alcance territorial', SCOPES[self.scope_selection]),
            ('Manzanas / veredas (NPN 1–17)', ', '.join(self.scope_prefixes) or 'Todas las del alcance seleccionado'),
        ]
        if self.model.currentText() == 'CICA':
            if self.cica_gdb:
                rows.append(('Geodatabase CICA', self.cica_gdb.filename))
            if self.cica_references:
                rows.append(('Archivos CICA', '\n'.join(table + ': ' + reference.filename for table, reference in sorted(self.cica_references.items()))))
        rows.append(('Origen de R1/R2', self.reference_mode.currentText()))
        required = needed_references(rules, general_report_path)
        for kind in ('R1', 'R2'):
            if self.reference_mode.currentIndex() == 0:
                reference = self.temporary_references.get(kind)
                value = (reference.filename + f' · {reference.row_count:,} registros cargados') if reference else 'No cargado'
            else:
                table = current_reference.get('tables', {}).get(kind) if current_reference else None
                value = 'r1_r2.' + table if table else 'No seleccionado o no disponible'
            rows.append((kind, value + (' · requerido' if kind in required else ' · no requerido por las reglas seleccionadas')))
        if current_reference:
            rows.extend([('Municipio de R1/R2', current_reference['divipola'] + ' · ' + current_reference.get('municipality', '')),
                         ('Fecha R1/R2 vigente', current_reference['date'])])
        counts = defaultdict(int)
        for component, _name, _sql, _description in rules:
            counts[component] += 1
        rows.append(('Controles a ejecutar', '\n'.join(f'{name}: {count} reglas' for name, count in sorted(counts.items(), key=lambda item: model_component_sort_key(item[0])))))
        rows.append(('Total de reglas', str(len(rules))))
        rows.append(('Reporte general', 'Sí' if general_report_path else 'No disponible en la versión seleccionada'))
        return rows

    def confirm_validation(self, params, schema, rules, current_reference, general_report_path=None):
        dialog = QDialog(self)
        dialog.setWindowTitle('Confirmar validación')
        dialog.resize(760, 650)
        layout = QVBoxLayout(dialog)
        title = QLabel('Revise los metadatos antes de ejecutar la validación')
        title.setWordWrap(True)
        title.setStyleSheet('font-size: 15px; font-weight: 700; color: #0468A3;')
        layout.addWidget(title)
        browser = QTextBrowser()
        browser.setOpenExternalLinks(False); browser.setOpenLinks(False)
        rows = self.validation_confirmation_metadata(params, schema, rules, current_reference, general_report_path)
        html = '<table width="100%" cellspacing="0" cellpadding="8">' + ''.join(
            '<tr><td width="32%" bgcolor="#EAF3F8"><b>' + escape(str(name)) + '</b></td><td>' +
            escape(str(value)).replace('\n', '<br>') + '</td></tr>' for name, value in rows) + '</table>'
        browser.setHtml(html); layout.addWidget(browser, 1)
        buttons = QHBoxLayout(); buttons.addStretch()
        cancel = QPushButton('Cancelar'); cancel.clicked.connect(dialog.reject)
        execute = QPushButton('Ejecutar validación'); execute.clicked.connect(dialog.accept)
        execute.setStyleSheet('background: #0468A3; color: white; padding: 8px 14px; font-weight: 700;')
        buttons.addWidget(cancel); buttons.addWidget(execute); layout.addLayout(buttons)
        cancel.setDefault(True); cancel.setFocus()
        return dialog.exec_() == QDialog.Accepted

    def run_rules(self):
        scope = self.scope_selection
        if scope not in SCOPES:
            QMessageBox.warning(self, 'Seleccione el alcance', 'Seleccione Todos, Urbano, Rural o Centro poblado. Se aplicará el mismo filtro a CICA, R1 y R2.')
            return
        self.update_run_availability()
        if not self.run_button.isEnabled():
            QMessageBox.warning(self, "Controles no disponibles", "Cargue una versión del repositorio SQL con controles activos y seleccione los grupos a ejecutar.")
            return
        if psycopg2 is None:
            QMessageBox.critical(self, "Dependencia faltante", "El QGIS instalado no incluye psycopg2.")
            return
        selected = [name for name, check in self.category_checks.items() if check.isChecked()]
        if not selected:
            QMessageBox.warning(self, "Seleccione controles", "Seleccione al menos una categoría de reglas.")
            return
        rules = [
            (item["component"], item["rule"], item["sql"], item.get("description", ""))
            for item in self.repository_rules
            if item["type"] == "CONTROL" and item["component"] in selected
            and item["sql"] and item["sql"].strip()
        ]
        general_report_path = next(
            (item["sql"] for item in self.repository_rules if item["type"] == "REPORTE_GENERAL"), None
        )
        if not rules:
            QMessageBox.critical(self, "Reglas no encontradas", "No hay reglas SQL en la selección actual.")
            return
        schema = self.schema.currentText().strip()
        if not schema:
            QMessageBox.warning(self, "Seleccione un esquema", "Pulse Listar Esquemas y seleccione el esquema que desea validar.")
            return
        try:
            params = self.connection_params()
        except Exception as error:
            QMessageBox.warning(self, "Conexión incompleta", str(error))
            return
        missing = self.missing_references(rules + ([('', '', general_report_path, '')] if general_report_path else []))
        if missing:
            QMessageBox.warning(self, 'Faltan referencias', 'Prepare o seleccione antes de ejecutar: ' + ', '.join(missing))
            return
        current_reference = self.selected_current_reference() if self.reference_mode.currentIndex() == 1 else None
        repository_params = self.repository_connection_params() if current_reference else None
        if current_reference and repository_params is None:
            return
        if current_reference:
            schema_codes = set(re.findall(r'(?<!\d)(\d{5})(?!\d)', schema))
            if len(schema_codes) == 1 and current_reference['divipola'] not in schema_codes:
                QMessageBox.warning(self, 'Municipio diferente', 'El DIVIPOLA seleccionado para R1 y R2 Vigente no coincide con el código del esquema a validar.')
                return
        if not self.confirm_validation(params, schema, rules, current_reference, general_report_path):
            return
        if not self.commit_review_editor() or not self.confirm_review_changes():
            return
        metadata = ConfigParser()
        metadata.read(PLUGIN_DIR / 'metadata.txt', encoding='utf-8')
        self.pending_format = validation_format(self.model.currentText(), self.environment.currentText(), self.selected_sql_version)
        self.completed_format = None
        self.format_download.setEnabled(False)
        self.execution_info = {
            'Fecha de ejecución': datetime.now(ZoneInfo('America/Bogota')).isoformat(timespec='seconds'),
            'Versión plugin': metadata.get('general', 'version', fallback='Desconocida'),
            'Modelo': self.model.currentText(), 'Ambiente': self.environment.currentText(),
            'Versión SQL': self.selected_sql_version.get('codigo', ''),
            'Base de datos': params.get('dbname', ''), 'Esquema': schema,
        }
        self.execution_info["Alcance territorial"] = SCOPES[scope]
        self.execution_info["Manzanas / veredas (NPN 1–17)"] = ", ".join(self.scope_prefixes) or "Todas"
        self.results.clear()
        self.review_cases = []
        self.review_imports = ReviewImports(self.review_cases)
        self.review_import_snapshot = {}
        self.detail_import_errors = {}
        self.refresh_detail_review_controls()
        self.review_snapshot = []
        self.review_dirty = False
        self.review_context = {
            "model": self.model.currentText(), "environment": self.environment.currentText(),
            "host": params.get("host", ""), "port": str(params.get("port", "5432")),
            "database": params.get("dbname", ""), "schema": schema,
            "reference_r1": (self.temporary_references['R1'].fingerprint if 'R1' in self.temporary_references else '') if self.reference_mode.currentIndex() == 0 else (current_reference['divipola'] + '/' + current_reference['stamp'] if current_reference else ''),

            "repository": self.selected_sql_version.get("source", {}).get("label", ""),
        }
        self.review_context["territorial_scope"] = scope
        self.review_context["territorial_prefixes"] = list(self.scope_prefixes)
        if current_reference:
            self.review_context['reference_source'] = 'sentencias.r1_r2'
            self.review_context['reference_municipality'] = current_reference['divipola']
            self.review_context['reference_date'] = current_reference['date']
            self.execution_info['R1 y R2 Vigente'] = current_reference['divipola'] + ' · ' + current_reference['date']
        if self.reference_mode.currentIndex() == 0 and 'R2' in self.temporary_references:
            self.review_context['reference_r2'] = self.temporary_references['R2'].fingerprint
        cica = dict(self.cica_references) if self.model.currentText() == 'CICA' else {}
        cica_gdb = self.cica_gdb if self.model.currentText() == 'CICA' else None
        if cica_gdb:
            self.review_context['cica_gdb'] = cica_gdb.fingerprint
        if cica:
            self.review_context['cica_csv'] = {table: reference.fingerprint for table, reference in sorted(cica.items())}
        self.table.setRowCount(0)
        self.general_report_rows = []
        self.general_report_error = ""
        self.validation_context = self.resolve_validation_context(params, schema)
        self.location_summary.setText(self.validation_context)
        self.progress.setRange(0, len(rules) + bool(general_report_path))
        self.progress.setValue(0)
        self.progress.setVisible(True)
        self.status.setText(f"Ejecutando {len(rules)} reglas y reporte general…")
        self.controls_running = True
        self.ingestion_button.setEnabled(False)
        self.export_button.setEnabled(False)
        self.run_button.setEnabled(False)
        self.thread = QThread(self)
        temporary = dict(self.temporary_references) if self.reference_mode.currentIndex() == 0 else {}
        self.worker = RuleWorker(params, schema, "", "", rules, general_report_path, references=temporary, cica_references=cica, cica_gdb=cica_gdb, current_reference=current_reference, repository_params=repository_params, territorial_scope=scope, territorial_prefixes=self.scope_prefixes, model=self.model.currentText())
        self.worker.reference_context.connect(self.review_context.update)
        self.worker.message.connect(self.status.setText)
        self.reference_mode.setEnabled(False)
        self.load_reference_button.setEnabled(False)
        self.refresh_references.setEnabled(False)
        self.municipality.setEnabled(False)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.result.connect(self.add_result)
        self.worker.general_report.connect(self.add_general_report)
        self.worker.progress.connect(self.set_progress)
        self.worker.fatal_error.connect(lambda text: QMessageBox.critical(self, "No fue posible ejecutar", text))
        self.worker.finished.connect(self.finished)
        self.worker.finished.connect(self.thread.quit)
        self.worker.finished.connect(self.worker.deleteLater)
        self.thread.finished.connect(self.thread.deleteLater)
        self.thread.start()

    def add_result(self, result):
        self.results.append(result)
        self.review_cases.extend(build_cases([result], self.review_context))

    def add_general_report(self, rows, error):
        self.general_report_rows = rows
        self.general_report_error = error

    def set_progress(self, current, total):
        self.progress.setRange(0, total)
        self.progress.setValue(current)
        self.status.setText(f"Ejecutando regla {current} de {total}…")

    def finished(self):
        self.controls_running = False
        self.ingestion_button.setEnabled(True)
        self.completed_format = self.pending_format if any(result.status != 'Error' for result in self.results) else None
        self.format_download.setEnabled(bool(self.completed_format))
        if self.completed_format:
            self.format_download.setText('Descargar · ' + self.completed_format[0])
        for control in (self.reference_mode, self.load_reference_button, self.refresh_references, self.municipality):
            control.setEnabled(True)
        self.update_run_availability()
        self.progress.setValue(self.progress.maximum())
        self.status.setText("Ejecución terminada.")
        self.last_execution_label.setText('Última ejecución: ' + self.execution_info.get('Fecha de ejecución', 'sesión actual'))
        self.review_imports = ReviewImports(self.review_cases)
        self.detail_import_errors = {}
        self.review_import_snapshot = self.review_imports.snapshot()
        self.refresh_detail_review_controls()
        self.review_snapshot = [case.tracking() + case.exception_tracking() for case in self.review_cases]


    def commit_review_editor(self):
        # Reviews are edited in Excel; no record-review window is created.
        return True

    def model_validation_results(self):
        return self.execution_info.get('Ambiente', self.environment.currentText()) == 'Validaciones del modelo'

    def ordered_result_groups(self, groups):
        return sorted(groups.items(), key=lambda item: model_component_sort_key(item[0])) if self.model_validation_results() else sorted(groups.items())

    def refresh_review_statistics(self):
        groups, totals = review_statistics(self.results, self.review_cases)
        self.table.setRowCount(0)
        rows = self.ordered_result_groups(groups) + ([('TOTAL', totals)] if groups else [])
        for component, values in rows:
            row = self.table.rowCount()
            self.table.insertRow(row)
            if self.model_validation_results():
                number = model_component_number(component)
                self.table.setVerticalHeaderItem(row, QTableWidgetItem('Σ' if component == 'TOTAL' else str(number) if number is not None else ''))
            for column, value in enumerate(statistics_row(component, values)):
                item = QTableWidgetItem(str(value))
                if column:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                if component == 'TOTAL':
                    font = item.font(); font.setBold(True); item.setFont(font)
                    item.setBackground(QColor('#E5F2FC'))
                self.table.setItem(row, column, item)
        for field, widget in self.kpi_values.items():
            widget.setText(f"{totals[field]:,}".replace(',', '.'))
        self.summary.setText(
            f"{totals['rules']} reglas · {totals['operator']} excepciones operador · "
            f"{totals['pending']} errores · {totals['errors']} errores técnicos. "
            "Errores = Hallazgos − Excepciones DT."
        )

    def refresh_detail_review_controls(self):
        self.refresh_review_statistics()
        components = sorted({case.component for case in self.review_cases})
        if set(components) != set(self.detail_review_rows):
            self.detail_review_table.setRowCount(0)
            self.detail_review_rows = {}
            self.detail_filename_labels = {}
            for row, component in enumerate(components):
                self.detail_review_table.insertRow(row)
                self.detail_review_table.setRowHeight(row, 30)
                self.detail_review_table.setItem(row, 0, QTableWidgetItem(component))
                filename = QLabel('—')
                filename.setWordWrap(True)
                load = QPushButton("Cargar")
                state = QLabel("No cargado")
                state.setWordWrap(True)
                load.clicked.connect(lambda checked=False, c=component: self.recover_detail(c))
                self.detail_review_table.setCellWidget(row, 1, filename)
                self.detail_review_table.setItem(row, 2, QTableWidgetItem('—'))
                self.detail_review_table.setCellWidget(row, 4, load)
                self.detail_review_table.setCellWidget(row, 3, state)
                self.detail_review_rows[component] = (load, state)
                self.detail_filename_labels[component] = filename
        for row, (component, (load, state)) in enumerate(self.detail_review_rows.items()):
            source = self.review_imports.sources.get(component)
            load.setEnabled(not self.controls_running and not self.exporting)
            load.setText("Recargar" if source else "Cargar")
            if source:
                counts = source['counts']
                warnings = sum(counts[name] for name in ('sin_correspondencia', 'ambiguas', 'conflictos', 'excepciones_invalidas', 'estados_invalidos', 'datos_invalidos', 'operador_bloqueado'))
                label = ('Cargado con observaciones' if warnings else 'Cargado') if counts['coincidencias'] else 'Sin coincidencias'
                self.detail_filename_labels[component].setText(Path(source['path']).name)
                self.detail_filename_labels[component].setToolTip(source['path'])
                state.setText('✓ ' + label)
                self.detail_review_table.item(row, 2).setText(datetime.fromisoformat(source['loaded_at']).strftime('%H:%M:%S'))
                self.detail_review_table.item(row, 2).setToolTip(source['loaded_at'])
                state.setToolTip(f"{source['path']}\n{counts['incorporadas']} incorporadas · {counts['conflictos']} conflictos")
                color = '#087F23' if counts['coincidencias'] and not warnings else '#A05A00'
                state.setStyleSheet(f'color: {color}; font-weight: 600;')
            else:
                self.detail_filename_labels[component].setText('—')
                self.detail_filename_labels[component].setToolTip('')
                state.setText('○ No cargado')
                state.setToolTip('')
                state.setStyleSheet('color: #575756;')
            if component in self.detail_import_errors:
                state.setText('✕ Error de carga')
                load.setText('Reintentar')
                state.setStyleSheet('color: #B00020; font-weight: 600;')
                state.setToolTip(self.detail_import_errors[component])
        self.check_review_files()

    def toggle_incidents(self, visible):
        self.incidents_browser.setVisible(visible)
        self.incidents_toggle.setText(('▾' if visible else '▸') + ' Ver incidencias de la última carga')

    def poll_review_files(self):
        if self.isVisible() and not self.controls_running and not self.exporting:
            self.check_review_files()

    def check_review_files(self, force=False):
        previous = set(self.review_file_issues)
        self.review_file_issues = self.review_imports.check_files(force)
        for component in previous - set(self.review_file_issues):
            if component in self.detail_review_rows and component in self.review_imports.sources and component not in self.detail_import_errors:
                source = self.review_imports.sources[component]
                counts = source['counts']
                warnings = sum(counts.get(name, 0) for name in ('sin_correspondencia', 'ambiguas', 'conflictos', 'excepciones_invalidas', 'estados_invalidos', 'datos_invalidos'))
                state = self.detail_review_rows[component][1]
                state.setText('✓ Cargado con observaciones' if warnings else '✓ Cargado')
                state.setStyleSheet('color: #A05A00;' if warnings else 'color: #087F23;')
                state.setToolTip(source['path'])
        for component, message in self.review_file_issues.items():
            if component in self.detail_review_rows:
                load, state = self.detail_review_rows[component]
                state.setText('⚠ Archivo modificado' if self.review_imports.sources[component].get('file_state') == 'modified' else '✕ Archivo no disponible')
                state.setStyleSheet('color: #A05A00; font-weight: 600;')
                state.setToolTip(message)
                load.setText('Recargar')
        self.review_file_alert.setVisible(bool(self.review_file_issues))
        self.review_file_alert.setText('⚠ Recargue los detalles modificados o no disponibles antes de generar el reporte: ' + ', '.join(self.review_file_issues) if self.review_file_issues else '')
        self.export_button.setEnabled(bool(self.results) and not self.controls_running and not self.exporting and not self.review_file_issues)
        return self.review_file_issues

    def missing_references(self, rules):
        missing = []
        for kind in ('R1', 'R2'):
            marker = kind.lower()
            needed = any('{tabla_' + marker + '}' in rule[2] or '{esquema_' + marker + '}' in rule[2] for rule in rules)
            if not needed:
                continue
            if self.reference_mode.currentIndex() == 0:
                ref = self.temporary_references.get(kind)
                if ref is None or not ref.csv_path.is_file():
                    missing.append(kind + ' temporal')
            else:
                selection = self.selected_current_reference()
                if selection is None or kind not in selection['tables']:
                    missing.append(kind + ' vigente (municipio/fecha sin referencia disponible)')
        return missing

    def recovery_message(self, counts):
        return (
            "Resumen de excepciones\n\n"
            "DT\n"
            f"• Incorporadas: {counts['dt_incorporadas']}\n"
            f"• Descartadas sin observación: {counts['dt_sin_observacion']}\n\n"
            "Operador\n"
            f"• Incorporadas: {counts['operador_incorporadas']}\n"
            f"• Descartadas sin observación: {counts['operador_sin_observacion']}\n"
            f"• Aprobadas por DT: {counts['operador_aprobadas_dt']}\n\n"
            "Incidencias de carga\n"
            f"• Sin coincidencia: {counts['sin_correspondencia']}\n"
            f"• Coincidencias ambiguas: {counts['ambiguas']}\n"
            f"• Conflictos (se conservó la revisión actual): {counts['conflictos']}\n"
            f"• Registros no recuperables: {counts['datos_invalidos']}\n"
            f"  — Por identidad insuficiente: {counts['identidad_insuficiente']}\n"
            f"  — Por datos o identidad alterados: {counts['identidad_alterada']}\n\n"
            + EXCEPTION_NOTICE
        )

    def recover_detail(self, component):
        if self.controls_running or self.exporting or not self.review_cases:
            return
        path, _ = QFileDialog.getOpenFileName(self, f"Cargar detalle: {component}", "", "Excel de revisión (*.xlsx)")
        if not path:
            return
        try:
            before = fingerprint_file(path)
            records = load_review(path)
            after = fingerprint_file(path)
            if before['sha256'] != after['sha256']:
                raise ValueError('El archivo cambió durante la carga. Guárdelo y recargue.')
            previous = [case.tracking() + case.exception_tracking() for case in self.review_cases]
            counts = self.review_imports.load(component, path, records, file_snapshot=after)
        except Exception as error:
            self.detail_import_errors[component] = str(error)
            self.refresh_detail_review_controls()
            self.incidents_browser.setPlainText(f'{component}: no se cargó el archivo.\n{error}\nSeleccione un detalle válido y vuelva a intentarlo.')
            QMessageBox.warning(self, "No se cargó el detalle", str(error))
            return
        self.review_dirty = self.review_dirty or previous != [case.tracking() + case.exception_tracking() for case in self.review_cases]
        self.detail_import_errors.pop(component, None)
        self.refresh_detail_review_controls()
        self.incidents_browser.setPlainText(component + '\n\n' + self.recovery_message(counts))
        QMessageBox.information(self, f"Detalle {component}", self.recovery_message(counts))

    def confirm_review_changes(self):
        if not self.review_dirty:
            return True
        choice = QMessageBox.question(
            self, "Revisión de la sesión", "Se descartará el seguimiento cargado en esta sesión. Los archivos de detalle no se modifican y podrá volver a cargarlos. ¿Continuar?",
            QMessageBox.Discard | QMessageBox.Cancel, QMessageBox.Cancel,
        )
        if choice == QMessageBox.Discard:
            for case, tracking in zip(self.review_cases, self.review_snapshot):
                for name, value in zip(REVIEW_FIELDS + EXCEPTION_FIELDS, tracking):
                    setattr(case, name, value)
            self.review_imports.restore_snapshot(self.review_import_snapshot)
            self.refresh_detail_review_controls()
            self.review_dirty = False
            return True
        return False

    def closeEvent(self, event):
        if self.controls_running or self.exporting or self.reference_lookup_running or self.filter_reading:
            QMessageBox.information(self, "Proceso en curso", "Espere a que termine el proceso actual antes de cerrar.")
            event.ignore()
        elif self.commit_review_editor() and self.confirm_review_changes():
            self.release_references()
            for reference in self.cica_references.values():
                reference.close()
            self.cica_references = {}
            if self.cica_gdb:
                self.cica_gdb.close()
                self.cica_gdb = None
            event.accept()
        else:
            event.ignore()

    def reject(self):
        self.close()

    def report_task(self, function, *args, **kwargs):
        """Run a workbook-only operation without blocking Qt's event loop."""
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(function, *args, **kwargs)
            loop = QEventLoop()
            timer = QTimer()
            timer.setInterval(50)
            timer.timeout.connect(lambda: loop.quit() if future.done() else None)
            timer.start()
            if not future.done():
                loop.exec_()
            timer.stop()
            return future.result()

    def export_excel(self):
        if self.exporting:
            return
        self.export_created_paths = []
        failure = None
        try:
            self.build_detail_reports()
        except Exception as error:
            failure = str(error) or type(error).__name__
            details = traceback.format_exc()
            cleanup_errors = []
            for path in self.export_created_paths:
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    cleanup_errors.append(str(path))
            if cleanup_errors:
                details += '\nNo se pudieron retirar archivos incompletos:\n' + '\n'.join(cleanup_errors)
            self.incidents_browser.setPlainText('Error al generar el reporte: ' + failure + '\n\n' + details)
            self.incidents_toggle.setChecked(True)
        finally:
            self.exporting = False
            self.left_tabs.setEnabled(True)
            self.category_button.setEnabled(bool(self.category_checks))
            self.update_run_availability()
            self.refresh_detail_review_controls()
            self.check_review_files()
            self.progress.hide()
        if failure is not None:
            self.status.setText('No se generó el reporte: ' + failure + '. Consulte las incidencias.')
            message = QMessageBox(self)
            message.setIcon(QMessageBox.Warning)
            message.setWindowTitle('No se generó el reporte')
            message.setText('La generación se interrumpió. Los resultados y revisiones siguen disponibles.')
            message.setInformativeText(failure)
            message.setDetailedText(self.incidents_browser.toPlainText())
            message.exec_()

    def save_generated_report(self, workbook, filename):
        if self.check_review_files(force=True):
            raise ValueError('Un detalle cambió durante la generación. Recárguelo y genere de nuevo. No se conserva una entrega incompleta.')
        path = Path(filename)
        if path.exists():
            raise FileExistsError('El archivo de destino ya existe; no se sobrescribe: ' + path.name)
        self.export_created_paths.append(path)
        self.status.setText("Guardando " + path.name + "… Espere; puede minimizar la ventana.")
        configure_report_workbook(workbook)
        self.report_task(workbook.save, path)

    def build_detail_reports(self):
        if self.controls_running:
            return
        if self.check_review_files(force=True):
            QMessageBox.warning(self, 'Revisión desactualizada', 'Hay detalles modificados o no disponibles. Guarde y recargue esos archivos antes de generar una nueva entrega.')
            return
        if not self.commit_review_editor():
            return
        if not self.results:
            QMessageBox.information(self, "Sin resultados", "Ejecute controles antes de exportar un reporte.")
            return
        directory = QFileDialog.getExistingDirectory(self, "Seleccione la carpeta para guardar el reporte detallado")
        if not directory:
            return
        self.exporting = True
        self.left_tabs.setEnabled(False)
        self.category_button.setEnabled(False)
        self.run_button.setEnabled(False)
        self.refresh_detail_review_controls()
        rule_sheet_count = len({(result.category, result.rule) for result in self.results if result.rows})
        package_count = len({result.category for result in self.results if result.rows})
        report_total = max(len(self.results) + rule_sheet_count + package_count + 3, 1)
        report_step = 0
        self.progress.setRange(0, report_total)
        self.progress.setValue(0)
        self.progress.setVisible(True)
        self.export_button.setEnabled(False)
        self.status.setText("Preparando reporte detallado…")
        QApplication.processEvents()

        def advance_report(message):
            nonlocal report_step
            report_step += 1
            self.progress.setValue(min(report_step, report_total))
            self.status.setText(message)
            QApplication.processEvents()

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        path = str(Path(directory) / f"resumen_control_calidad_{stamp}.xlsx")
        book = Workbook()
        summary = book.active
        validation_metadata = dict(self.execution_info)
        validation_metadata.setdefault('Modelo', self.review_context.get('model', self.model.currentText()))
        validation_metadata.setdefault('Ambiente', self.review_context.get('environment', self.environment.currentText()))
        validation_metadata.setdefault('Esquema', self.review_context.get('schema', ''))
        summary_header_row = prepare_validation_index(
            summary, validation_metadata, self.validation_context, title='Resumen ejecutivo', headers=STAT_HEADERS,
            widths={'A': 30, 'B': 20, 'C': 18, 'D': 20, 'E': 23, 'F': 18, 'G': 20})
        groups, totals = review_statistics(self.results, self.review_cases)
        for result in self.results:
            advance_report("Organizando resumen ejecutivo…")
        for category, values in self.ordered_result_groups(groups):
            summary.append(statistics_row(category, values))
        summary.append(statistics_row('TOTAL', totals))
        summary.auto_filter.ref = f'A{summary_header_row}:G{summary.max_row}'
        general = book.create_sheet("Reporte general")
        general.append(["ID", "Ítem evaluado", "Cantidad", "Reglas asociadas"])
        if self.general_report_rows:
            for item in self.general_report_rows:
                values = {str(key).lower(): value for key, value in item.items()}
                general.append([
                    values.get("id", ""), values.get("descripcion", ""), values.get("cantidad", ""),
                    values.get("reglas_validacion_afectadas", values.get("reglas_validacion_afectadas", "")),
                ])
        else:
            general.append(["", "No fue posible generar el reporte general", "", self.general_report_error or "No se recibieron resultados."])
        detail = book.create_sheet("Detalle Validación")
        detail.append(["Componente", "Regla incumplida", "Hallazgos", "Hoja de detalle"])
        error_rows = []
        review_by_row = {
            id(row): case for row, case in zip(
                [row for result in self.results for row in result.rows or []], self.review_cases
            )
        }
        rows_by_rule = defaultdict(list)
        rule_descriptions = {}
        for result in self.results:
            rule_descriptions[(result.category, result.rule)] = result.description or ""
            for row in result.rows or []:
                predio = self.predio_from_row(row)
                error_rows.append((predio, result.category, result.rule, row))
                rows_by_rule[(result.category, result.rule)].append((predio, row))
        green = PatternFill("solid", fgColor="068CBB")
        navy = green
        red = green
        soft_blue = PatternFill("solid", fgColor="EAF3F8")
        soft_gray = PatternFill("solid", fgColor="F3F6F8")
        description_fill = PatternFill("solid", fgColor="DDF4F5")
        warning = PatternFill("solid", fgColor="FFF2CC")
        success = PatternFill("solid", fgColor="D9EAD3")
        exception = PatternFill("solid", fgColor="FCE4D6")
        category_fills = {category: green for category, _rule in rows_by_rule}
        for sheet, header_fill in ((general, green), (detail, navy)):
            for cell in sheet[1]:
                cell.fill = header_fill; cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
            for row_index in range(2, sheet.max_row + 1):
                if row_index % 2 == 0:
                    for cell in sheet[row_index]:
                        cell.fill = soft_blue
        for row in summary.iter_rows(min_row=summary_header_row + 1):
            total_row = row[0].value == 'TOTAL'
            for cell in row:
                cell.alignment = Alignment(vertical='center', horizontal='left' if cell.column == 1 else 'right')
                cell.font = Font(name='Calibri', size=11, color='344054', bold=total_row)
                if total_row or cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='DCECF9' if total_row else 'F0F6FB')
            summary.row_dimensions[row[0].row].height = 26
        plugin_version = self.execution_info.get('Versión plugin')
        if not plugin_version:
            plugin_metadata = ConfigParser()
            plugin_metadata.read(PLUGIN_DIR / 'metadata.txt', encoding='utf-8')
            plugin_version = plugin_metadata.get('general', 'version', fallback='No disponible')
        version_row = summary.max_row + 2
        summary.cell(version_row, 1, 'Versión plugin')
        version_cell = summary.cell(version_row, 2, str(plugin_version)); version_cell.data_type = 's'
        for cell in (summary.cell(version_row, 1), version_cell):
            cell.fill = PatternFill('solid', fgColor='EAF3F8')
            cell.font = Font(color='0468A3', bold=True)
            cell.alignment = Alignment(vertical='center')
        summary.row_dimensions[version_row].height = 24
        used_sheet_names = set(book.sheetnames)
        rule_sheet_names = {}
        for category, rule in sorted(rows_by_rule, key=lambda item: (item[0], rule_sort_key(item[1]))):
            safe_rule = re.sub(r"[\\\\/*?:\[\]]", "_", rule_sheet_label(rule))
            sheet_name = f"{safe_rule}"[:31]
            suffix = 2
            while sheet_name in used_sheet_names:
                sheet_name = f"{safe_rule[:28]}_{suffix}"
                suffix += 1
            used_sheet_names.add(sheet_name)
            rule_sheet_names[(category, rule)] = sheet_name
            detail.append([category, rule, len(rows_by_rule[(category, rule)]), sheet_name])
            detail.cell(detail.max_row, 4).hyperlink = f"#{sheet_name}!A1"
            rule_sheet = book.create_sheet(sheet_name)
            fields = sql_report_fields(row for _predio, row in rows_by_rule[(category, rule)])
            rule_headers = [*fields, "clave_caso", *EXCEPTION_HEADERS]
            last_column = len(rule_headers)
            rule_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
            rule_sheet.cell(1, 1, f"Regla incumplida: {rule}")
            rule_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
            rule_sheet.cell(2, 1, rule_descriptions[(category, rule)] or "Descripción de la regla no disponible en el repositorio de sentencias.")
            rule_sheet.row_dimensions[2].height = 54
            rule_sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
            rule_sheet.cell(3, 1, self.validation_context)
            rule_sheet.row_dimensions[3].height = 30
            for column, header in enumerate(rule_headers, start=1):
                rule_sheet.cell(4, column, header)
            for index, (predio, row) in enumerate(rows_by_rule[(category, rule)], start=5):
                if index % 250 == 0:
                    self.status.setText(f"Preparando {rule}: {index - 4} registros…")
                    QApplication.processEvents()
                for column, value in enumerate([
                    *[predio_value(row) if key == PREDIO_HEADER else row.get(key, "") for key in fields],
                    review_by_row[id(row)].key if id(row) in review_by_row else "",
                    *(review_by_row[id(row)].exception_tracking() if id(row) in review_by_row else EXCEPTION_DEFAULTS),
                ], start=1):
                    cell = rule_sheet.cell(index, column, display(value))
                    cell.data_type = 's'
                    cell.number_format = '@'
            for cell in rule_sheet[1]:
                cell.fill = PatternFill("solid", fgColor="0468A3"); cell.font = Font(color="FFFFFF", bold=True, size=11)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            rule_sheet.cell(2, 1).fill = description_fill
            rule_sheet.cell(2, 1).font = Font(color="0468A3", bold=True)
            rule_sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
            rule_sheet.cell(3, 1).fill = PatternFill("solid", fgColor="EAF3F8")
            rule_sheet.cell(3, 1).font = Font(color="0468A3", italic=True)
            rule_sheet.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
            for cell in rule_sheet[4]:
                cell.fill = category_fills[category]; cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            rule_sheet.freeze_panes = "A5"; rule_sheet.auto_filter.ref = f"A4:{rule_sheet.cell(4, last_column).coordinate[:-1]}{rule_sheet.max_row}"
            for column_index, column in enumerate(rule_sheet.iter_cols(), start=1):
                rule_sheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
            for row_index, cells in enumerate(rule_sheet.iter_rows(min_row=5), 5):
                if row_index % 250 == 0:
                    QApplication.processEvents()
                if row_index % 2 == 0:
                    for cell in cells:
                        cell.fill = soft_gray
            self.report_task(add_controls, rule_sheet, 5, rule_sheet.max_row, rule_sheet.max_column - 3)
            advance_report(f"Generando detalle de {rule}…")
        by_predio = book.create_sheet("Hallazgos por predio")
        by_predio.append([PREDIO_HEADER, "Componentes", "Reglas incumplidas", "Total hallazgos", "Observación"])
        grouped_predios = defaultdict(lambda: {"components": set(), "rules": set(), "total": 0})
        for predio, component, rule, _row in error_rows:
            group = grouped_predios[predio or "Sin identificador"]
            group["components"].add(component)
            group["rules"].add(f"{component}: {rule}")
            group["total"] += 1
        for predio, values in sorted(grouped_predios.items()):
            by_predio.append([
                predio, ", ".join(sorted(values["components"])), "; ".join(sorted(values["rules"])), values["total"], "",
            ])
        for cell in by_predio[1]:
            cell.fill = navy; cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        by_predio.freeze_panes = "A2"; by_predio.auto_filter.ref = by_predio.dimensions
        for column in by_predio.columns:
            by_predio.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        for row_index, cells in enumerate(by_predio.iter_rows(min_row=2), 2):
            if row_index % 250 == 0:
                QApplication.processEvents()
            if row_index % 2 == 0:
                for cell in cells:
                    cell.fill = soft_blue
        advance_report("Generando resumen por predio…")
        technical = book.create_sheet("Errores técnicos")
        technical.append(["Componente", "Regla", "Detalle técnico"])
        technical_errors = [result for result in self.results if result.status == "Error"]
        if technical_errors:
            for result in technical_errors:
                technical.append([result.category, result.rule, result.detail])
        else:
            technical.append(["Sin errores técnicos", "", "La ejecución terminó sin errores de conexión o SQL."])
        for cell in technical[1]:
            cell.fill = red; cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        technical.freeze_panes = "A2"; technical.auto_filter.ref = technical.dimensions
        technical.column_dimensions["A"].width = 24
        technical.column_dimensions["B"].width = 18
        technical.column_dimensions["C"].width = 90
        for category in sorted({category for category, _rule in rows_by_rule}):
            package = Workbook()
            package_index = package.active
            index_header_row = prepare_validation_index(package_index, validation_metadata, self.validation_context, category)
            for rule in sorted((rule for item_category, rule in rows_by_rule if item_category == category), key=rule_sort_key):
                source = book[rule_sheet_names[(category, rule)]]
                # Each component workbook owns its sheet namespace; a matching
                # number in another component must not add an artificial suffix.
                target_name = re.sub(r"[\\\\/*?:\[\]]", "_", rule_sheet_label(rule))[:31]
                target = package.create_sheet(target_name)
                package_index.append([category, rule, len(rows_by_rule[(category, rule)]), target.title])
                internal_detail_link(package_index.cell(package_index.max_row, 4), target.title)
                for source_row in source.iter_rows():
                    if source_row[0].row % 250 == 0:
                        QApplication.processEvents()
                    for cell in source_row:
                        copied = target.cell(cell.row, cell.column, cell.value)
                        copied.data_type = cell.data_type
                        copied.number_format = cell.number_format
                for merged_range in source.merged_cells.ranges:
                    target.merge_cells(str(merged_range))
                for cell in target[1]:
                    cell.fill = PatternFill("solid", fgColor="0468A3"); cell.font = Font(color="FFFFFF", bold=True, size=11)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                target.cell(2, 1).fill = description_fill
                target.cell(2, 1).font = Font(color="0468A3", bold=True)
                target.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
                target.row_dimensions[2].height = 54
                target.cell(3, 1).fill = PatternFill("solid", fgColor="EAF3F8")
                target.cell(3, 1).font = Font(color="0468A3", italic=True)
                target.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
                target.row_dimensions[3].height = 30
                for cell in target[4]:
                    cell.fill = green; cell.font = Font(color="FFFFFF", bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                for row_index, cells in enumerate(target.iter_rows(min_row=5), 5):
                    if row_index % 250 == 0:
                        QApplication.processEvents()
                    if row_index % 2 == 0:
                        for cell in cells:
                            cell.fill = soft_gray
                for column_index, column in enumerate(target.iter_cols(), start=1):
                    target.column_dimensions[get_column_letter(column_index)].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
                target.freeze_panes = "A5"; target.auto_filter.ref = f"A4:{target.cell(4, target.max_column).coordinate[:-1]}{target.max_row}"
            package_index.auto_filter.ref = f'A{index_header_row}:D{package_index.max_row}'
            safe_category = re.sub(r"[^A-Za-z0-9_-]+", "_", category).strip("_") or "paquete"
            self.report_task(save_review, None, [case for case in self.review_cases if case.component == category],
                        [result for result in self.results if result.category == category], workbook=package)
            self.report_task(link_rule_reviews, package)
            self.save_generated_report(package, Path(directory) / f"detalle_{safe_category}_{stamp}.xlsx")
            advance_report(f"Guardando paquete {category}…")
        for sheet in [detail, *[book[name] for name in rule_sheet_names.values()]]:
            book.remove(sheet)
        advance_report("Guardando reporte Excel…")
        # The executive file remains statistical only; human follow-up lives
        # exclusively in each detail workbook, not in another editable snapshot.
        self.save_generated_report(book, path)
        if self.check_review_files(force=True):
            raise ValueError('Un detalle cambió durante la generación. Recárguelo y genere de nuevo.')
        self.progress.setValue(report_total)
        self.export_button.setEnabled(True)
        self.status.setText("Reporte exportado correctamente.")

    @staticmethod
    def predio_from_row(row):
        return predio_value(row)
