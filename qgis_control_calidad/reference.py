"""Selección y validación de los registros R1 y R2 en formato Excel."""
from pathlib import Path
import re
import unicodedata
from datetime import datetime

from openpyxl import load_workbook
from qgis.PyQt.QtWidgets import QApplication, QComboBox, QDialog, QFileDialog, QFormLayout, QHBoxLayout, QLabel, QLineEdit, QMessageBox, QProgressBar, QPushButton, QVBoxLayout

try:
    import psycopg2
    from psycopg2 import sql
except ImportError:
    psycopg2 = None


REQUIRED = {
    "R1": {"NUMERO_PREDIAL", "DESTINO_ECONOMICO", "AREA_TERRENO", "AREA_CONSTRUIDA", "AVALUO", "VIGENCIA"},
    "R2": {"NUMERO_PREDIAL", "MATRICULA_INMOBILIARIA", "AREA_TERRENO_1", "AREA_CONSTRUIDA_1", "VIGENCIA"},
}


class ReferenceDialog(QDialog):
    def __init__(self, parent=None, params_provider=None):
        super().__init__(parent)
        self.params_provider = params_provider
        self.setWindowTitle("Referencias catastrales R1 y R2")
        self.resize(680, 290)
        layout = QVBoxLayout(self)
        title = QLabel("Cargar en public R1 y R2")
        title.setStyleSheet("font-size: 18px; font-weight: 600; color: #08756a;")
        layout.addWidget(title)
        note = QLabel("Seleccione archivos Excel (.xlsx). Los TXT de ancho fijo se incorporarán en una versión posterior.")
        note.setWordWrap(True)
        layout.addWidget(note)
        form = QFormLayout()
        self.r1 = self.file_row(form, "Archivo R1")
        self.r2 = self.file_row(form, "Archivo R2")
        self.schema = QLineEdit("public")
        self.schema.setReadOnly(True)
        form.addRow("Esquema de referencia", self.schema)
        layout.addLayout(form)
        self.status = QLabel("Seleccione los archivos R1 y R2.")
        layout.addWidget(self.status)
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        layout.addWidget(self.progress)
        buttons = QHBoxLayout()
        buttons.addStretch()
        validate = QPushButton("Validar y cargar en public")
        validate.clicked.connect(self.validate)
        self.accept_button = QPushButton("Aceptar")
        self.accept_button.setEnabled(False)
        self.accept_button.clicked.connect(self.accept)
        buttons.addWidget(validate)
        buttons.addWidget(self.accept_button)
        layout.addLayout(buttons)

    def file_row(self, form, label):
        field = QLineEdit(); field.setReadOnly(True)
        choose = QPushButton("Seleccionar")
        choose.clicked.connect(lambda: self.choose(field))
        field.choose_button = choose
        row = QHBoxLayout(); row.addWidget(field, 1); row.addWidget(choose)
        form.addRow(label, row)
        return field

    def choose(self, field):
        path, _ = QFileDialog.getOpenFileName(self, "Seleccione archivo Excel", "", "Excel (*.xlsx)")
        if path:
            field.setText(path)

    def load_existing_tables(self):
        if psycopg2 is None or self.params_provider is None:
            self.status.setText("Configure una conexión PostgreSQL antes de consultar las referencias publicadas.")
            return
        try:
            with psycopg2.connect(**self.params_provider(), connect_timeout=8) as connection:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = 'public' AND (lower(table_name) LIKE 'r1%' OR lower(table_name) LIKE '%_r1%') ORDER BY table_name")
                    r1_tables = [row[0] for row in cursor.fetchall()]
                    cursor.execute("SELECT table_name FROM information_schema.tables WHERE table_schema = 'public' AND (lower(table_name) LIKE 'r2%' OR lower(table_name) LIKE '%_r2%') ORDER BY table_name")
                    r2_tables = [row[0] for row in cursor.fetchall()]
            self.existing_r1.clear(); self.existing_r1.addItems(r1_tables)
            self.existing_r2.clear(); self.existing_r2.addItems(r2_tables)
            self.status.setText(f"Se encontraron {len(r1_tables)} R1 y {len(r2_tables)} R2 publicados en public.")
        except Exception as error:
            self.status.setText("No fue posible consultar las referencias publicadas.")
            QMessageBox.warning(self, "No fue posible consultar public", str(error))

    def validate_one(self, kind, filename):
        workbook = load_workbook(filename, read_only=True, data_only=True)
        sheet = workbook.active
        headers = {str(value).strip().upper() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)) if value is not None}
        missing = REQUIRED[kind] - headers
        if missing:
            return f"{kind}: faltan {', '.join(sorted(missing))}"
        return None

    def validate(self):
        if not self.r1.text() or not self.r2.text():
            QMessageBox.warning(self, "Archivos incompletos", "Seleccione los archivos R1 y R2.")
            return
        errors = []
        for kind, field in (("R1", self.r1), ("R2", self.r2)):
            try:
                error = self.validate_one(kind, field.text())
                if error: errors.append(error)
            except Exception as exc:
                errors.append(f"{kind}: {exc}")
        if errors:
            QMessageBox.critical(self, "Estructura no válida", "\n".join(errors))
            return
        try:
            count_r1 = self.row_count(self.r1.text())
            count_r2 = self.row_count(self.r2.text())
            total = count_r1 + count_r2
            self.progress.setRange(0, max(total, 1)); self.progress.setValue(0); self.progress.setVisible(True)
            table_r1, loaded_r1 = self.load_excel_to_public("r1", self.r1.text(), 0, total)
            table_r2, loaded_r2 = self.load_excel_to_public("r2", self.r2.text(), loaded_r1, total)
        except Exception as error:
            self.progress.setVisible(False)
            self.status.setText("No fue posible cargar R1 y R2 en public.")
            QMessageBox.critical(self, "Error de carga", str(error))
            return
        self.progress.setValue(self.progress.maximum())
        self.status.setText(f"Carga terminada: public.{table_r1} ({loaded_r1} registros) y public.{table_r2} ({loaded_r2} registros).")
        self.accept_button.setEnabled(True)
        QMessageBox.information(self, "Carga finalizada", "R1 y R2 se cargaron correctamente en public. Pulse Aceptar para regresar.")

    @staticmethod
    def safe_name(prefix, filename):
        name = unicodedata.normalize("NFKD", Path(filename).stem).encode("ascii", "ignore").decode().lower()
        name = re.sub(r"[^a-z0-9]+", "_", name).strip("_")
        name = re.sub(r"(^|_)r[12]$", "", name).strip("_")
        suffix = datetime.now().strftime("%Y%m%d")
        if not name:
            return f"{prefix}_{suffix}"
        available = 60 - len(prefix) - len(suffix) - 2
        return f"{prefix}_{name[:available]}_{suffix}"

    @staticmethod
    def column_name(value, index):
        name = unicodedata.normalize("NFKD", ReferenceDialog.clean_excel_value(value)).encode("ascii", "ignore").decode().lower()
        name = re.sub(r"[^a-z0-9]+", "_", name).strip("_") or f"campo_{index}"
        return name[:60]

    @staticmethod
    def clean_excel_value(value):
        """Convierte las secuencias OOXML (_xNNNN_) al carácter original."""
        text = "" if value is None else str(value)
        return re.sub(r"_x([0-9a-fA-F]{4})_", lambda match: chr(int(match.group(1), 16)), text)

    @staticmethod
    def row_count(filename):
        workbook = load_workbook(filename, read_only=True, data_only=True)
        return max(workbook.active.max_row - 1, 0)

    def load_excel_to_public(self, prefix, filename, completed, total):
        if psycopg2 is None or self.params_provider is None:
            raise RuntimeError("Configure una conexión PostgreSQL para cargar las referencias.")
        workbook = load_workbook(filename, read_only=True, data_only=True)
        sheet = workbook.active
        source_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = [self.column_name(value, index) for index, value in enumerate(source_headers, 1)]
        if len(columns) != len(set(columns)):
            raise ValueError("El archivo contiene encabezados que se repiten al normalizarlos.")
        table = self.safe_name(prefix, filename)
        rows = ([self.clean_excel_value(value) for value in row] for row in sheet.iter_rows(min_row=2, values_only=True))
        with psycopg2.connect(**self.params_provider()) as connection:
            with connection.cursor() as cursor:
                cursor.execute(sql.SQL("DROP TABLE IF EXISTS public.{}" ).format(sql.Identifier(table)))
                definitions = sql.SQL(", ").join(sql.SQL("{} TEXT").format(sql.Identifier(column)) for column in columns)
                cursor.execute(sql.SQL("CREATE TABLE public.{} ({})").format(sql.Identifier(table), definitions))
                query = sql.SQL("INSERT INTO public.{} ({}) VALUES ({})").format(
                    sql.Identifier(table),
                    sql.SQL(", ").join(sql.Identifier(column) for column in columns),
                    sql.SQL(", ").join(sql.Placeholder() for _ in columns),
                )
                batch = []
                loaded = 0
                for row in rows:
                    batch.append(row)
                    loaded += 1
                    if len(batch) == 1000:
                        cursor.executemany(query, batch); batch = []
                        self.progress.setValue(completed + loaded)
                        self.status.setText(f"Cargando {prefix.upper()} en public: {loaded} de {total - completed} registros…")
                        QApplication.processEvents()
                if batch:
                    cursor.executemany(query, batch)
        self.progress.setValue(completed + loaded)
        QApplication.processEvents()
        return table, loaded
