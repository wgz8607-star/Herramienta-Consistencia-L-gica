"""Prepare local reference snapshots and install session-only PostgreSQL tables."""
import csv
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
import tempfile
import unicodedata

from openpyxl import load_workbook
try:
    from psycopg2 import sql
except ImportError:
    sql = None

REQUIRED = {
    'R1': {'numero_predial', 'destino_economico', 'area_terreno', 'area_construida', 'avaluo', 'vigencia'},
    'R2': {'numero_predial', 'matricula_inmobiliaria', 'area_terreno_1', 'area_construida_1', 'vigencia'},
}
TEMP_TABLES = {'R1': 'reference_r1', 'R2': 'reference_r2'}


def clean(value):
    text = '' if value is None else str(value)
    return re.sub(r'_x([0-9a-fA-F]{4})_', lambda m: chr(int(m.group(1), 16)), text)


def column_name(value, index):
    text = unicodedata.normalize('NFKD', clean(value)).encode('ascii', 'ignore').decode().lower()
    return (re.sub(r'[^a-z0-9]+', '_', text).strip('_') or f'campo_{index}')[:60]


@dataclass
class PreparedReference:
    kind: str
    filename: str
    columns: tuple
    row_count: int
    fingerprint: str
    csv_path: Path
    directory: object

    def close(self):
        self.directory.cleanup()


def prepare_reference(kind, filename, progress=None):
    if kind not in REQUIRED:
        raise ValueError('Tipo de referencia no válido.')
    directory = tempfile.TemporaryDirectory(prefix='qgis_reference_')
    workbook = None
    try:
        workbook = load_workbook(filename, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows()
        header = next(rows, None)
        if not header:
            raise ValueError(f'{kind}: el archivo no tiene encabezados.')
        columns = tuple(column_name(cell.value, index) for index, cell in enumerate(header, 1))
        if len(columns) != len(set(columns)):
            raise ValueError(f'{kind}: encabezados duplicados después de normalizarlos.')
        missing = REQUIRED[kind] - set(columns)
        if missing:
            raise ValueError(f'{kind}: faltan columnas: ' + ', '.join(sorted(missing)))
        digest = hashlib.sha256(json.dumps(columns).encode())
        path = Path(directory.name) / 'reference.csv'
        count = 0
        with path.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream, quoting=csv.QUOTE_ALL)
            for row in rows:
                if all(cell.value is None for cell in row):
                    continue
                values = []
                for cell in row:
                    value = clean(cell.value)
                    if isinstance(cell.value, (int, float)) and float(cell.value).is_integer() and re.fullmatch('0+', cell.number_format or ''):
                        value = str(int(cell.value)).zfill(len(cell.number_format))
                    values.append(value)
                writer.writerow(values)
                digest.update(json.dumps(values, ensure_ascii=False, separators=(',', ':')).encode('utf-8'))
                digest.update(b'\n')
                count += 1
                if progress and count % 1000 == 0:
                    progress(kind, count)
        return PreparedReference(kind, Path(filename).name, columns, count, digest.hexdigest(), path, directory)
    except Exception:
        directory.cleanup()
        raise
    finally:
        if workbook is not None:
            workbook.close()


def install_reference(connection, reference):
    """No DROP/CREATE in public. TEMP rows survive rule-level rollback after commit."""
    if sql is None:
        raise RuntimeError('Se requiere psycopg2 para cargar referencias temporales.')
    table = TEMP_TABLES[reference.kind]
    definitions = sql.SQL(', ').join(sql.SQL('{} TEXT').format(sql.Identifier(c)) for c in reference.columns)
    with connection.cursor() as cursor:
        cursor.execute(sql.SQL('CREATE TEMP TABLE {} ({}) ON COMMIT PRESERVE ROWS').format(sql.Identifier(table), definitions))
        query = sql.SQL('COPY pg_temp.{} ({}) FROM STDIN WITH (FORMAT CSV, ENCODING \'UTF8\')').format(
            sql.Identifier(table), sql.SQL(', ').join(sql.Identifier(c) for c in reference.columns))
        with reference.csv_path.open('r', encoding='utf-8', newline='') as stream:
            cursor.copy_expert(query.as_string(connection), stream)
