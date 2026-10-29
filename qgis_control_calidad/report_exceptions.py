"""DT/operator exception controls shared by every record-level Excel report."""
from copy import copy
import math
import unicodedata

from openpyxl.styles import PatternFill, Protection, Font, Alignment
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

HEADERS = ('Excepción DT', 'Observación DT', 'Excepción operador', 'Observación operador')
FIELDS = ('excepcion_dt', 'observacion_dt', 'excepcion_operador', 'observacion_operador')
DEFAULTS = ('No', '', 'No', '')
EXCEPTION_COLORS = {'Sí': ('C6EFCE', '006100'), 'No': ('FFC7CE', '9C0006')}


def yes(value):
    return str(value or '').strip().casefold() in ('sí', 'si')


def validate(values, valid_observation):
    dt, note_dt, operator, note_operator = values
    if str(dt) not in ('Sí', 'No') or str(operator) not in ('Sí', 'No'):
        raise ValueError('Las excepciones deben marcarse como Sí o No.')
    if yes(dt) and not valid_observation(note_dt):
        raise ValueError('Excepción DT requiere una Observación DT válida.')
    if yes(operator) and not valid_observation(note_operator):
        raise ValueError('Excepción operador requiere una Observación operador válida.')


def fit_observations(sheet, first_row, last_row, first_column):
    """Fit notes to current content, up to Excel's column/row limits."""
    def length(value):
        return sum(0 if unicodedata.combining(c) else 2 if unicodedata.east_asian_width(c) in ('W', 'F') else 1
                   for c in value.expandtabs(4))
    for column in (first_column + 1, first_column + 3):
        lines = [line for row in range(first_row, last_row+1)
                 for line in str(sheet.cell(row, column).value or '').splitlines()]
        width = min(255, max(48, max((length(line) for line in lines), default=0) + 3))
        sheet.column_dimensions[get_column_letter(column)].width = width
        for row in range(first_row, last_row+1):
            cell = sheet.cell(row, column)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = 'top'
            cell.alignment = alignment
            text = str(cell.value or '')
            count = sum(max(1, math.ceil(length(line) / max(1, width-2))) for line in (text.splitlines() or ['']))
            sheet.row_dimensions[row].height = min(409.5, max(sheet.row_dimensions[row].height or 15, count*15+4))


def add_controls(sheet, first_row, last_row, first_column, protect=True):
    """Independent visible Sí/No dropdowns for both roles."""
    if last_row < first_row:
        return
    dt, note_dt, operator, note_operator = [get_column_letter(first_column + i) for i in range(4)]
    for offset, header in enumerate(HEADERS):
        sheet.cell(first_row - 1, first_column + offset).comment = Comment(
            'Seleccione Sí o No. Las observaciones se pueden diligenciar antes o después. Las marcas y observaciones del operador permanecen visibles aunque DT marque Sí.', 'Validador')
    for column in (dt, operator):
        validation = DataValidation(type='list', formula1='"Sí,No"', allow_blank=False)
        validation.showDropDown = False
        validation.showErrorMessage = True
        validation.errorStyle = 'stop'
        validation.errorTitle = 'Valor fuera del dominio'
        validation.error = 'Seleccione Sí o No.'
        validation.showInputMessage = True
        validation.promptTitle = 'Excepción'
        validation.prompt = 'Seleccione Sí o No; la observación se diligencia libremente.'
        sheet.add_data_validation(validation)
        validation.add(f'{column}{first_row}:{column}{last_row}')
    for column in (dt, operator):
        for value, (color, text_color) in EXCEPTION_COLORS.items():
            sheet.conditional_formatting.add(
                f'{column}{first_row}:{column}{last_row}',
                FormulaRule(formula=[f'${column}{first_row}="{value}"'],
                            fill=PatternFill('solid', fgColor=color, bgColor=color),
                            font=Font(color=text_color), stopIfTrue=True),
            )
    # Keep initial colors visible even in previews that do not calculate Excel
    # conditional formatting; conditional rules still update subsequent edits.
    for row_index in range(first_row, last_row + 1):
        for column_index in (first_column, first_column + 2):
            cell = sheet.cell(row_index, column_index)
            if cell.value in EXCEPTION_COLORS:
                fill_color, text_color = EXCEPTION_COLORS[cell.value]
                cell.fill = PatternFill('solid', fgColor=fill_color)
                font = copy(cell.font)
                font.color = text_color
                cell.font = font
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=first_column, max_col=first_column+3):
        for cell in row:
            cell.protection = Protection(locked=False)
    fit_observations(sheet, first_row, last_row, first_column)
    if protect:
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False


def link_formula(sheet_name, key_column, first_row, last_row, role_column, summary_row):
    quoted = "'" + sheet_name.replace("'", "''") + "'"
    keys = f'{quoted}!${get_column_letter(key_column)}${first_row}:${get_column_letter(key_column)}${last_row}'
    values = f'{quoted}!${get_column_letter(role_column)}${first_row}:${get_column_letter(role_column)}${last_row}'
    return f'IF(COUNTIF({keys},$A{summary_row})=1,IFERROR(INDEX({values},MATCH($A{summary_row},{keys},0)),""),"")'


def link_rule_reviews(book):
    """Make detail rows editable and Revisión a key-based consolidated view."""
    summary = book['Revisión']
    origins = book.create_sheet('_OrigenRevision')
    origins.append(['hoja', 'columna_clave', 'primera_fila', 'ultima_fila', 'columna_excepcion'])
    lookup = {}
    for sheet in book:
        if sheet.title in ('Revisión', '_OrigenRevision') or sheet.max_row < 5:
            continue
        headers = {cell.value: cell.column for cell in sheet[4] if cell.value is not None}
        if 'clave_caso' not in headers or not all(name in headers for name in HEADERS):
            continue
        key_col, role_col = headers['clave_caso'], headers[HEADERS[0]]
        last_row = sheet.max_row
        origins.append([sheet.title, key_col, 5, last_row, role_col])
        sheet.column_dimensions[get_column_letter(key_col)].hidden = True
        add_controls(sheet, 5, last_row, role_col)
        for row in range(5, last_row+1):
            key = sheet.cell(row, key_col).value
            lookup.setdefault(key, []).append((sheet.title, key_col, 5, last_row, role_col))
    origins.sheet_state = 'hidden'
    for row in range(2, summary.max_row+1):
        key = summary.cell(row, 1).value
        matches = lookup.get(key, [])
        if len(matches) != 1:
            continue
        name, key_col, first, last, role_col = matches[0]
        for offset in range(4):
            cell = summary.cell(row, 9+offset)
            cell.value = '=' + link_formula(name, key_col, first, last, role_col+offset, row)
            cell.protection = Protection(locked=True)
    # Detail editing is the sole source in a workbook with linked rule sheets.
    summary.data_validations.dataValidation = []


def configure_report_workbook(book):
    """Protect report structure and data, retaining exception edits and column widths."""
    book.security.lockStructure = True
    for sheet in book:
        if sheet.sheet_state != 'visible':
            continue
        # False means the action is allowed in Excel's protection flags.
        sheet.protection.sheet = True
        sheet.protection.formatColumns = False
        sheet.protection.formatRows = True
        sheet.protection.formatCells = True
        sheet.protection.insertColumns = True
        sheet.protection.insertRows = True
        sheet.protection.deleteColumns = True
        sheet.protection.deleteRows = True
        sheet.protection.insertHyperlinks = True
        sheet.protection.sort = True
        sheet.protection.autoFilter = False
        sheet.protection.objects = True
        sheet.protection.scenarios = True
        sheet.protection.pivotTables = True
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False
        coordinate = sheet.freeze_panes
        row = coordinate_to_tuple(coordinate.coordinate if hasattr(coordinate, 'coordinate') else coordinate)[0] if coordinate else 2
        sheet.freeze_panes = 'B' + str(max(2, row))


def internal_detail_link(cell, sheet_name):
    """An explicit in-workbook destination, visually distinct from plain text."""
    text = '↪ Ir a ' + sheet_name
    quoted = "'" + sheet_name.replace("'", "''") + "'"
    cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=quoted + '!A1', display=text,
                               tooltip='Vínculo interno: abrir la hoja ' + sheet_name + ' en este mismo archivo.')
    cell.font = Font(name='Calibri', size=11, color='0066B3', underline='single', bold=True)
    cell.fill = PatternFill('solid', fgColor='EAF3FF')
    alignment = copy(cell.alignment)
    alignment.vertical = 'center'; alignment.wrap_text = True; alignment.indent = 1
    cell.alignment = alignment
    cell.parent.row_dimensions[cell.row].height = max(26, cell.parent.row_dimensions[cell.row].height or 0)


def prepare_validation_index(sheet, metadata, location, component=None, title=None, headers=None, widths=None):
    """Validation metadata above a results or navigation table, without a context tab."""
    sheet.title = title or 'Detalle Validación'
    headers = headers or ('Componente', 'Regla incumplida', 'Hallazgos', 'Hoja de detalle')
    last_column = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = sheet.cell(1, 1, sheet.title + (' — ' + str(component) if component else ''))
    title.data_type = 's'
    title.fill = PatternFill('solid', fgColor='0468A3')
    title.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    title.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 36
    fields = [(key, metadata.get(key) or 'No disponible') for key in (
        'Fecha de ejecución', 'Modelo', 'Ambiente', 'Versión SQL', 'Alcance territorial',
        'Manzanas / veredas (NPN 1–17)', 'Base de datos', 'Esquema')]
    fields.append(('Ubicación territorial', location or 'No disponible'))
    for key in ('R1 y R2 Vigente',):
        if metadata.get(key):
            fields.append((key, metadata[key]))
    for row, (label, value) in enumerate(fields, 2):
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_column)
        name = sheet.cell(row, 1, label); name.data_type = 's'
        cell = sheet.cell(row, 2, str(value)); cell.data_type = 's'
        name.font = Font(name='Calibri', size=11, bold=True, color='0468A3')
        name.fill = PatternFill('solid', fgColor='EAF3F8')
        cell.font = Font(name='Calibri', size=11, color='344054')
        name.alignment = cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[row].height = min(105, max(28, math.ceil(len(str(value)) / 75) * 15 + 8))
    header_row = len(fields) + 3
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='068CBB')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[header_row].height = 30
    for column, width in (widths or {'A': 30, 'B': 28, 'C': 16, 'D': 28}).items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'B' + str(header_row + 1)
    sheet.sheet_view.showGridLines = False
    return header_row
