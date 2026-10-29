"""Human review independent of QGIS and of the automatic validation outcome."""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
import hashlib
import json
import math
import os
import re
from pathlib import Path
import tempfile
import unicodedata
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from .report_identifiers import PREDIO_HEADER, is_predio_field, predio_value, CONDITION_HEADER, condition_text
from .report_exceptions import HEADERS as EXCEPTION_HEADERS, FIELDS as EXCEPTION_FIELDS, DEFAULTS as EXCEPTION_DEFAULTS, yes, validate, add_controls, link_formula, configure_report_workbook

PENDING = "Revisión pendiente"
EXCEPTION = "Excepción"
ADJUST = "Solicitar ajuste"
STATES = (PENDING, EXCEPTION, ADJUST)
FORMAT = "revision-catastral-1"
REVIEW_FIELDS = ("estado_revision_usuario", "observacion_usuario", "fecha_revision_usuario")
# These are explicitly positional fields, not object identifiers such as id/t_id.
POSITION_FIELDS = {"row_number", "row_num", "rownum", "numero_fila", "nro_fila", "fila_reporte"}


def valid_observation(value):
    return isinstance(value, str) and any(
        not c.isspace() and unicodedata.category(c)[0] not in "CZ" for c in value
    )


def canonical(value):
    """Retain SQL types and exact strings (including leading zeroes)."""
    if value is None:
        return ["null", None]
    if isinstance(value, bool):
        return ["bool", value]
    if isinstance(value, str):
        return ["str", value]
    if isinstance(value, int):
        return ["int", str(value)]
    if isinstance(value, Decimal):
        return ["decimal", str(value)]
    if isinstance(value, float):
        return ["float", repr(value)]
    if isinstance(value, (datetime, date, time)):
        return [type(value).__name__, value.isoformat()]
    if isinstance(value, (bytes, bytearray, memoryview)):
        return ["bytes", bytes(value).hex()]
    if isinstance(value, dict):
        return ["dict", [[str(k), canonical(v)] for k, v in sorted(value.items(), key=lambda x: str(x[0]))]]
    if isinstance(value, (list, tuple)):
        return ["list", [canonical(v) for v in value]]
    raise ValueError(f"Tipo de dato no compatible con recuperación segura: {type(value).__name__}")


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def digest(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def display(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return dumps(canonical(value))
    if isinstance(value, (bytes, bytearray, memoryview)):
        return bytes(value).hex()
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


@dataclass
class ReviewCase:
    key: str
    identity: str
    component: str
    rule: str
    automatic: str
    data: dict
    recoverable: bool = True
    estado_revision_usuario: str = PENDING
    observacion_usuario: str = ""
    fecha_revision_usuario: str = ""
    excepcion_dt: str = "No"
    observacion_dt: str = ""
    excepcion_operador: str = "No"
    observacion_operador: str = ""

    def exception_tracking(self):
        return tuple(getattr(self, name) for name in EXCEPTION_FIELDS)

    def tracking(self):
        return tuple(getattr(self, name) for name in REVIEW_FIELDS)

    def has_tracking(self):
        return self.tracking() != (PENDING, "", "") or self.exception_tracking() != EXCEPTION_DEFAULTS

    def set_review(self, state, observation, timestamp=None):
        if state not in STATES:
            raise ValueError("Estado de revisión no válido.")
        if state == EXCEPTION and not valid_observation(observation):
            raise ValueError("Una Excepción requiere una Observación con texto válido.")
        if (state, observation) != self.tracking()[:2]:
            self.estado_revision_usuario = state
            self.observacion_usuario = observation
            self.fecha_revision_usuario = timestamp or datetime.now(ZoneInfo("America/Bogota")).isoformat(timespec="seconds")
            return True
        return False


def build_cases(results, context):
    cases = []
    for result in results:
        for row in result.rows or []:
            fields = {k: v for k, v in row.items() if str(k).lower() not in POSITION_FIELDS}
            has_identifier = any(
                re.search(r"(^|_)(id|tid|uuid|guid|fid|npn|nit|documento|cedula|identificacion|predio|predial|codigo|matricula|fmi)(_|$)", str(k).lower())
                and str(k).lower() not in {'codigo_regla', 'id_regla', 'codigo_validacion', 'id_validacion'}
                and v is not None and display(v).strip() for k, v in fields.items()
            )
            safe = (bool(getattr(result, "criterion_hash", ""))
                    and getattr(result, "identity_safe", True) and has_identifier)
            try:
                payload = {
                    "format": FORMAT, "context": context,
                    "validation": [result.category, result.rule],
                    "criterion": getattr(result, "criterion_hash", ""),
                    "automatic": result.status, "data": canonical(fields),
                }
                identity = dumps(payload)
            except ValueError:
                safe = False
                identity = dumps({"format": FORMAT, "unsupported": True})
            cases.append(ReviewCase(digest(identity), identity, result.category, result.rule,
                                    result.status, dict(row), safe))
    return cases


def _text(cell, value):
    value = display(value)
    if len(value) > 32767:
        raise ValueError("Un dato supera el límite de texto de Excel (32767 caracteres); no se guardó un archivo incompleto.")
    cell.value = value
    cell.data_type = "s"  # Prevent Excel formulas and numeric conversion of identifiers.
    cell.number_format = "@"


def style_rule_component_sheet(sheet):
    sheet.sheet_properties.tabColor = '0468A3'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    widths = {'A': 25, 'B': 20, 'C': 14, 'D': 50, 'E': 85}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 32
    line = Side(style='thin', color='E1E8EF')
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='0468A3')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        fill = 'F0F6FB' if row[0].row % 2 == 0 else 'FFFFFF'
        lines = 1
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color='344054')
            cell.fill = PatternFill('solid', fgColor=fill)
            cell.border = Border(bottom=line)
            cell.alignment = Alignment(vertical='center', wrap_text=True,
                                       horizontal='center' if cell.column == 3 else 'left')
            width = widths[cell.column_letter]
            lines = max(lines, sum(max(1, math.ceil(len(part) / max(1, width - 3)))
                                   for part in str(cell.value or '').split('\n')))
        row[0].font = Font(name='Calibri', size=11, bold=True, color='0468A3')
        try:
            findings = int(row[2].value)
        except (TypeError, ValueError):
            findings = None
        if findings is not None:
            row[2].font = Font(name='Calibri', size=11, bold=True, color='9C5700' if findings else '217346')
            row[2].fill = PatternFill('solid', fgColor='FFF3D6' if findings else 'EAF5EC')
        if row[3].value:
            row[3].font = Font(name='Calibri', size=11, color='9C0006')
            row[3].fill = PatternFill('solid', fgColor='FCECEE')
        sheet.row_dimensions[row[0].row].height = min(409, max(30, lines * 15 + 8))


def save_review(path, cases, results=(), general_rows=(), workbook=None):
    for case in cases:
        validate(case.exception_tracking(), valid_observation)
        if case.estado_revision_usuario not in STATES:
            raise ValueError("Estado de revisión no válido.")
        if case.estado_revision_usuario == EXCEPTION and not valid_observation(case.observacion_usuario):
            raise ValueError("No se puede guardar una Excepción sin Observación válida.")
    book = workbook if workbook is not None else Workbook()
    sheet = book.create_sheet("Revisión") if workbook is not None else book.active
    sheet.title = "Revisión"
    fields = sorted({str(k) for case in cases for k in case.data if not is_predio_field(k)})
    headers = ["clave_caso", "componente", "validacion", "resultado_automatico", *REVIEW_FIELDS,
               "recuperable", *EXCEPTION_HEADERS, PREDIO_HEADER, *["dato:" + name for name in fields]]
    sheet.append(headers)
    for index, case in enumerate(cases, 2):
        values = [case.key, case.component, case.rule, case.automatic, *case.tracking(),
                  "Sí" if case.recoverable else "No", *case.exception_tracking(), predio_value(case.data), *[case.data.get(name) for name in fields]]
        for column, value in enumerate(values, 1):
            _text(sheet.cell(index, column), value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
    for column in ('E', 'F', 'G'):
        sheet.column_dimensions[column].hidden = True  # Legacy format compatibility.
    add_controls(sheet, 2, len(cases) + 1, 9)
    meta = book.create_sheet("_Identidades")
    meta.append([FORMAT, "fragmento", "contenido"])
    identities = {case.key: case.identity for case in cases}
    meta_row = 1
    for key, identity in identities.items():
        for part, offset in enumerate(range(0, len(identity), 30000)):
            meta.append([key, part, identity[offset:offset+30000]])
            meta_row += 1
            for cell in (meta.cell(meta_row, 1), meta.cell(meta_row, 3)):
                cell.data_type = 's'
    meta.sheet_state = "hidden"
    execution = book.create_sheet("Reglas & Componente")
    execution.append(["Componente", "Validación", "Hallazgos", "Detalle técnico", "Descripción"])
    for result in results:
        execution.append([None] * 5)
        for col, value in enumerate([result.category, result.rule, result.findings, result.detail, result.description], 1):
            _text(execution.cell(execution.max_row, col), value)
    style_rule_component_sheet(execution)
    if general_rows:
        general = book.create_sheet("Reporte general")
        original_columns = {str(k) for row in general_rows for k in row}
        columns = ([PREDIO_HEADER] if any(is_predio_field(k) for k in original_columns) else []) + sorted(k for k in original_columns if not is_predio_field(k))
        general.append(columns)
        for row_index, row in enumerate(general_rows, 2):
            for col, name in enumerate(columns, 1):
                _text(general.cell(row_index, col), predio_value(row) if name == PREDIO_HEADER else row.get(name))
    statistics = book.create_sheet("Estadística de revisión")
    statistics.append(STAT_HEADERS)
    groups, totals = review_statistics(results, cases)
    for component, values in sorted(groups.items()):
        statistics.append(statistics_row(component, values))
    statistics.append(statistics_row('TOTAL', totals))
    statistics.freeze_panes = 'A2'
    statistics.auto_filter.ref = statistics.dimensions
    sheet.sheet_state = "hidden"
    configure_report_workbook(book)
    # Write atomically: a failed save must not replace the last saved review.
    if path is None:
        return book
    path = Path(path)
    handle, temporary = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(handle)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        book.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


@dataclass
class ImportedReview:
    key: str
    state: str
    observation: str
    timestamp: str
    recoverable: bool
    valid: bool = True
    exception_values: tuple = EXCEPTION_DEFAULTS
    component: str = ""


def load_review(path):
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if "_Identidades" not in book or "Revisión" not in book:
            raise ValueError("Este archivo no es una revisión guardada por la herramienta; los reportes anteriores no contienen una clave segura.")
        meta = book["_Identidades"]
        if meta.cell(1, 1).value != FORMAT:
            raise ValueError("Formato de revisión no compatible.")
        chunks = defaultdict(dict)
        for key, part, content in meta.iter_rows(min_row=2, max_col=3, values_only=True):
            if key is None:
                continue
            if not isinstance(part, int) or part < 0 or part in chunks[key] or not isinstance(content, str):
                raise ValueError("Metadatos de identidad duplicados o dañados.")
            chunks[key][part] = content
        identities = {}
        for key, pieces in chunks.items():
            if sorted(pieces) != list(range(len(pieces))):
                raise ValueError("Identidad incompleta.")
            identity = ''.join(pieces[i] for i in range(len(pieces)))
            if digest(identity) != key:
                raise ValueError("Identidad de caso alterada.")
            identities[key] = json.loads(identity)
        sheet = book["Revisión"]
        headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
        required = ["clave_caso", "componente", "validacion", "resultado_automatico", *REVIEW_FIELDS, "recuperable"]
        if any(headers.count(name) != 1 for name in required) or len(headers) != len(set(headers)):
            raise ValueError("Faltan columnas obligatorias o hay columnas duplicadas.")
        linked = defaultdict(list)
        if '_OrigenRevision' in book:
            for name, key_col, first, last, role_col in book['_OrigenRevision'].iter_rows(min_row=2, max_col=5, values_only=True):
                if name not in book or not all(isinstance(v, int) and v > 0 for v in (key_col, first, last, role_col)):
                    raise ValueError('Referencias de hojas de regla inválidas.')
                source = book[name]
                source_headers = {cell.value: cell.column for cell in next(source.iter_rows(min_row=4, max_row=4)) if cell.value is not None}
                if source_headers.get('clave_caso') != key_col or any(source_headers.get(header) != role_col+i for i, header in enumerate(EXCEPTION_HEADERS)):
                    raise ValueError('Columnas de revisión alteradas en una hoja de regla.')
                for source_cells in source.iter_rows(min_row=first, max_row=last):
                    values = [cell.value for cell in source_cells]
                    key = values[key_col-1]
                    role_values = tuple(str(values[role_col+i-1] or default) for i, default in enumerate(EXCEPTION_DEFAULTS))
                    data = {header: values[col-1] for header, col in source_headers.items()}
                    safe = not any(cell.data_type == 'f' for cell in source_cells)
                    linked[key].append((name, key_col, first, last, role_col, role_values, data, safe))
        records = []
        for sheet_row, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            row = dict(zip(headers, [c.value for c in cells]))
            if not any(value is not None for value in row.values()):
                continue
            key = row.get("clave_caso")
            payload = identities.get(key)
            links = linked.get(key, [])
            role_values = tuple(str(row.get(name) or default) for name, default in zip(EXCEPTION_HEADERS, EXCEPTION_DEFAULTS))
            formula_headers = set()
            source_data = None
            link_valid = True
            if links:
                if len(links) != 1:
                    link_valid = False
                else:
                    name, key_col, first, last, role_col, role_values, source_data, link_valid = links[0]
                    # Consolidation formulas are presentation only. Excel rewrites
                    # quoting/function prefixes on save; never evaluate them or use
                    # their cached values. Read the unique keyed source row instead.
                    formula_headers.update(EXCEPTION_HEADERS)
            valid = payload is not None and link_valid and not any(
                cell.data_type == 'f' and header not in formula_headers for header, cell in zip(headers, cells)
            )
            if valid and not payload.get("unsupported"):
                original = dict(payload["data"][1])
                actual = {h[5:]: '' if row[h] is None else str(row[h]) for h in headers if str(h).startswith('dato:') and h[5:].lower() not in POSITION_FIELDS}
                expected = {k: _display_canonical(v) for k, v in original.items()}
                # Blank columns from other rules are harmless; changed data is not.
                # New exports coalesce displayed NPN aliases, while identities retain
                # every original SQL field so matching still changes with the data.
                unified = PREDIO_HEADER in row
                visible_expected = {k: v for k, v in expected.items() if not (unified and is_predio_field(k))}
                valid = all(actual.get(k) == v for k, v in visible_expected.items()) and not any(v for k, v in actual.items() if k not in expected)
                if unified:
                    valid = valid and ('' if row[PREDIO_HEADER] is None else str(row[PREDIO_HEADER])) == predio_value(expected)
                valid = valid and [row['componente'], row['validacion']] == payload['validation'] and row['resultado_automatico'] == payload['automatic']
                if source_data is not None and CONDITION_HEADER in source_data:
                    # The compact detail omits object attributes. Their immutable
                    # snapshot and case key are still checked in hidden Revisión.
                    value = source_data.get(CONDITION_HEADER)
                    valid = valid and ('' if value is None else str(value)) == condition_text(expected)
                elif source_data is not None:
                    if PREDIO_HEADER in source_data:
                        value = source_data[PREDIO_HEADER]
                        valid = valid and ('' if value is None else str(value)) == (predio_value(expected) or 'Sin identificador')
                    for field, expected_value in expected.items():
                        if PREDIO_HEADER in source_data and is_predio_field(field):
                            continue
                        source_field = field if field in source_data else ('Predio' if re.sub(r'[^a-z0-9]', '', field.lower()) == 'numeropredialnacional' else field)
                        actual_value = source_data.get(source_field)
                        valid = valid and ('' if actual_value is None else str(actual_value)) == expected_value
            records.append(ImportedReview(key, str(row[REVIEW_FIELDS[0]] or PENDING),
                                          str(row[REVIEW_FIELDS[1]] or ''), str(row[REVIEW_FIELDS[2]] or ''),
                                          (row.get('recuperable') == 'Sí' or _legacy_matricula_identity(payload)), valid,
                                          role_values, payload.get('validation', [''])[0] if payload else ''))
        return records
    finally:
        book.close()


def _legacy_matricula_identity(payload):
    """Old reports marked matrícula-only groups unsupported; retain their keys.

    The current case must independently pass recoverability checks; this does
    not promote duplicate or unsupported current query outputs.
    """
    if not payload or payload.get('unsupported') or not payload.get('criterion'):
        return False
    for name, value in payload.get('data', ['', []])[1]:
        if re.search(r'(^|_)(matricula|fmi)(_|$)', name.lower()) and _display_canonical(value).strip():
            return True
    return False


def _display_canonical(value):
    kind, data = value
    if kind == 'null':
        return ''
    if kind == 'bool':
        return str(data)
    if kind in ('dict', 'list'):
        return dumps(value)
    return str(data)


def restore_review(cases, records):
    current_counts = Counter(case.key for case in cases)
    old_counts = Counter(record.key for record in records)
    current = {case.key: case for case in cases}
    counts = dict(incorporadas=0, sin_correspondencia=0, ambiguas=0, conflictos=0,
                  excepciones_invalidas=0, estados_invalidos=0, datos_invalidos=0, sin_seguimiento=0, operador_bloqueado=0, identidad_insuficiente=0, identidad_alterada=0,
                  dt_incorporadas=0, dt_sin_observacion=0, operador_incorporadas=0,
                  operador_sin_observacion=0, operador_aprobadas_dt=0)
    for record in records:
        if not record.valid or not record.recoverable:
            counts['datos_invalidos'] += 1
            counts['identidad_alterada' if not record.valid else 'identidad_insuficiente'] += 1
            continue
        if record.key not in current:
            counts['sin_correspondencia'] += 1
            continue
        case = current[record.key]
        if current_counts[record.key] != 1 or old_counts[record.key] != 1 or not case.recoverable:
            counts['ambiguas'] += 1
            continue
        state = record.state
        if state not in STATES:
            state = PENDING
            counts['estados_invalidos'] += 1
        if state == EXCEPTION and not valid_observation(record.observation):
            state = PENDING
            counts['excepciones_invalidas'] += 1
        dt, note_dt, operator, note_operator = record.exception_values
        if any(str(value).strip().casefold() not in ('sí', 'si', 'no', '') for value in (dt, operator)):
            counts['estados_invalidos'] += 1
        dt = 'Sí' if yes(dt) else 'No'
        operator = 'Sí' if yes(operator) else 'No'
        operator_marked = operator == 'Sí'
        operator_justified = operator_marked and valid_observation(note_operator)
        if dt == 'Sí' and not valid_observation(note_dt):
            dt = 'No'
            counts['dt_sin_observacion'] += 1
            counts['excepciones_invalidas'] += 1
        if operator_marked and not operator_justified:
            operator = 'No'
            counts['operador_sin_observacion'] += 1
            counts['excepciones_invalidas'] += 1
        operator_approved = dt == 'Sí' and operator_justified
        exception_values = (dt, note_dt, operator, note_operator)
        tracking = (state, record.observation, record.timestamp)
        if case.has_tracking():
            if case.tracking() != tracking or case.exception_tracking() != exception_values:
                counts['conflictos'] += 1
            continue
        if tracking == (PENDING, '', '') and exception_values == EXCEPTION_DEFAULTS:
            counts['sin_seguimiento'] += 1
            continue
        for name, value in zip(REVIEW_FIELDS, tracking):
            setattr(case, name, value)
        for name, value in zip(EXCEPTION_FIELDS, exception_values):
            setattr(case, name, value)
        counts['incorporadas'] += 1
        counts['dt_incorporadas'] += int(dt == 'Sí')
        counts['operador_incorporadas'] += int(operator == 'Sí')
        counts['operador_aprobadas_dt'] += int(operator_approved)
    return counts


STAT_HEADERS = ('Componente', 'Reglas ejecutadas', 'Hallazgos',
                'Excepciones DT', 'Excepciones operador', 'Errores', 'Errores técnicos')


def review_statistics(results, cases):
    """Review changes outstanding findings, never the automatic SQL outcome."""
    grouped = {}
    for result in results:
        values = grouped.setdefault(result.category, dict(rules=0, findings=0, pending=0, dt=0, operator=0, errors=0))
        values['rules'] += 1
        values['findings'] += result.findings if isinstance(result.findings, int) else 0
        values['errors'] += result.status == 'Error'
    if not results:
        for case in cases:
            values = grouped.setdefault(case.component, dict(rules=0, findings=0, pending=0, dt=0, operator=0, errors=0))
            values['findings'] += 1
        for component, values in grouped.items():
            values['rules'] = len({case.rule for case in cases if case.component == component})
    for case in cases:
        if case.component not in grouped:
            continue
        values = grouped[case.component]
        if yes(case.excepcion_dt) and valid_observation(case.observacion_dt):
            values['dt'] += 1
        if yes(case.excepcion_operador) and valid_observation(case.observacion_operador):
            values['operator'] += 1
    for values in grouped.values():
        values['pending'] = max(0, values['findings'] - values['dt'])
    totals = {field: sum(values[field] for values in grouped.values())
              for field in ('rules', 'findings', 'pending', 'dt', 'operator', 'errors')}
    return grouped, totals


MODEL_COMPONENT_ORDER = {'estructura': 0, 'administrativas': 1, 'juridicas': 2,
                         'fisicas': 3, 'economicas': 4, 'novedades': 5,
                         'etnico': 9, 'etnicos': 9, 'etnica': 9, 'etnicas': 9}


def model_component_number(component):
    name = unicodedata.normalize('NFKD', str(component)).encode('ascii', 'ignore').decode().strip().lower()
    return MODEL_COMPONENT_ORDER.get(name)


def model_component_sort_key(component):
    number = model_component_number(component)
    return (number if number is not None else 100, str(component).casefold())


def statistics_row(component, values):
    return (component, values['rules'], values['findings'], values['dt'], values['operator'],
            values['pending'], values['errors'])
