"""Excel rule templates and reviewed, insert-only repository imports."""
from collections import defaultdict
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
try:
    from psycopg2 import sql
except ImportError:
    sql = None

COLUMNS = ('version', 'componente', 'codigo_regla', 'tipo', 'sentencia_sql',
           'descripcion_regla', 'activa', 'orden', 'descripcion_version', 'version_publicada')
TABLES = ('reglas_cargue_snc_cica', 'reglas_cargue_snc_xtf',
          'reglas_cica_v2_29_07_2024', 'reglas_cica_v3_07_09_2026',
          'reglas_interno_v2_29_07_2024', 'reglas_interno_v3_07_09_2026')
TEMPLATE = Path(__file__).parent / 'resources' / 'plantillas' / 'plantilla_reglas.xlsx'


@dataclass(frozen=True)
class RuleUpload:
    row: int
    values: tuple

    @property
    def key(self):
        return self.values[:3]


@dataclass
class UploadPlan:
    table: str
    rules: list
    signature: str
    inserts: list
    states: list


def boolean(value, optional=False):
    if value is None or isinstance(value, str) and not value.strip():
        if optional:
            return None
        raise ValueError('indique Sí o No')
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in ('sí', 'si', 'true', '1'):
        return True
    if text in ('no', 'false', '0'):
        return False
    raise ValueError('use Sí o No')


def create_template(filename, existing=()):
    for row_number, values in enumerate(existing, 2):
        for column, value in zip(COLUMNS, values):
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(f'Fila {row_number}, {column}: excede el límite de Excel (32.767 caracteres); no se exporta para evitar truncar datos.')
    book = Workbook(); sheet = book.active; sheet.title = 'Reglas'
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='0468A3')
        cell.font = Font(color='FFFFFF', bold=True)
    widths = (24, 24, 24, 22, 85, 60, 14, 12, 55, 22)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number, values in enumerate(existing, 2):
        for column, value in enumerate(values, 1):
            if column in (7, 10) and value is not None:
                value = 'Sí' if value else 'No'
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str):
                cell.data_type = 's'
    last_row = max(201, len(existing) + 1)
    for row in sheet.iter_rows(min_row=2, max_row=last_row, max_col=len(COLUMNS)):
        for index, cell in enumerate(row):
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if index != 7:
                cell.number_format = '@'
    for column in ('G', 'J'):
        validation = DataValidation(type='list', formula1='"Sí,No"', allow_blank=column == 'J')
        validation.showDropDown = False; validation.showErrorMessage = True; validation.errorStyle = 'stop'
        validation.error = 'Seleccione Sí o No.'
        sheet.add_data_validation(validation); validation.add(f'{column}2:{column}1048576')
    types = DataValidation(type='list', formula1='"CONTROL,REPORTE_GENERAL"', allow_blank=False)
    types.showDropDown = False; types.showErrorMessage = True; types.errorStyle = 'stop'
    sheet.add_data_validation(types); types.add('D2:D1048576')
    sheet.freeze_panes = 'E2'; sheet.auto_filter.ref = f'A1:J{last_row}'
    guide = book.create_sheet('Instrucciones')
    instructions = [
        'PLANTILLA DE SENTENCIAS SQL — una fila por regla',
        'Diligencie únicamente la hoja Reglas. La hoja Ejemplo no se carga.',
        'No cambie los encabezados. Puede agregar más filas; las filas vacías se omiten.',
        'version, componente, codigo_regla, tipo, sentencia_sql, activa y orden son obligatorios. descripcion_regla debe explicar el control; se permite vacía para conservar registros existentes.',
        'Conserve version y codigo_regla como texto para no perder ceros iniciales.',
        'Pegue toda la consulta SQL en una sola celda sentencia_sql, incluidos sus saltos de línea. No añada $regla$ ni INSERT INTO.',
        'Use {esquema}.tabla para consultar las fuentes elegidas por el plugin, incluidas las temporales.',
        'La consulta CONTROL debe devolver únicamente hallazgos. El importador guarda el SQL; no comprueba su sintaxis ni lo ejecuta.',
        'activa: Sí para habilitar la regla; No para guardarla en preparación. orden: entero desde 0.',
        'descripcion_version y version_publicada son opcionales. Para una versión nueva, complete ambos en al menos una fila.',
        'Una versión nueva se guarda sin publicar si version_publicada está vacío. Si se repiten metadatos de una versión, deben coincidir.',
        'En una versión existente se conservan su descripción y publicación actuales; esta carga no las modifica.',
        'Elija el destino en Administrar reglas. El archivo no contiene credenciales ni decide el esquema de destino.',
        'Se detectan duplicados por version + componente + codigo_regla. Dentro del archivo se rechazan.',
        'Las reglas que ya existen no se sobrescriben. La vista previa distingue coincidencias y contenido diferente.',
        'Revise la vista previa y pulse Incorporar reglas nuevas. Una falla revierte toda la operación.',
        'Cargue SNC permite una sola versión publicada con reglas activas. No se despublican versiones automáticamente.',
        'Se requieren permisos de administración de la tabla elegida: consulta, inserción y bloqueo de escritura durante la carga. El plugin no concede permisos.',
        'No utilice fórmulas de Excel. El límite por celda es de 32.767 caracteres; no exceda ese tamaño al pegar SQL.',
        'Destinos admitidos: ' + ', '.join('public.' + table for table in TABLES),
    ]
    for text in instructions:
        guide.append([text]); guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[guide.max_row].height = 36
    guide.column_dimensions['A'].width = 120
    example = book.create_sheet('Ejemplo')
    example.append(COLUMNS)
    example.append(['V1_SNC', 'Cargue SNC', 'snc_001', 'CONTROL',
                    "SELECT numero_predial, 'Número predial vacío' AS observacion\nFROM {esquema}.cica_predio\nWHERE numero_predial IS NULL OR BTRIM(numero_predial) = '';",
                    'Detecta predios sin número predial.', 'No', 100, 'Validaciones para cargue a SNC', 'No'])
    for index, width in enumerate(widths, 1):
        example.column_dimensions[get_column_letter(index)].width = width
    for cell in example[2]:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    example.row_dimensions[2].height = 90
    book.save(filename); book.close()


def read_rules(filename):
    book = load_workbook(filename, read_only=True, data_only=False)
    try:
        if 'Reglas' not in book.sheetnames:
            raise ValueError('El archivo debe incluir la hoja Reglas.')
        rows = book['Reglas'].iter_rows()
        header = next(rows, None)
        if not header:
            raise ValueError('La hoja Reglas no tiene encabezados.')
        names = [str(cell.value or '').strip().lower() for cell in header]
        while names and not names[-1]:
            names.pop()
        if len(names) != len(set(names)) or set(names) != set(COLUMNS):
            raise ValueError('Los encabezados deben ser: ' + ', '.join(COLUMNS))
        rules, errors, keys = [], [], set()
        for row_number, cells in enumerate(rows, 2):
            raw = {name: cells[index].value if index < len(cells) else None for index, name in enumerate(names)}
            if all(value is None or isinstance(value, str) and not value.strip() for value in raw.values()):
                continue
            try:
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError('no se admiten fórmulas de Excel')
                values = []
                for name in COLUMNS[:6]:
                    value = raw[name]
                    if name == 'descripcion_regla' and value is None:
                        value = ''
                    if not isinstance(value, str) or name != 'descripcion_regla' and not value.strip():
                        raise ValueError(name + ': debe contener texto')
                    values.append(value if name == 'sentencia_sql' else value.strip())
                if values[3] not in ('CONTROL', 'REPORTE_GENERAL'):
                    raise ValueError('tipo: use CONTROL o REPORTE_GENERAL')
                active = boolean(raw['activa'])
                order = raw['orden']
                if isinstance(order, bool) or order is None or not str(order).strip().isdigit():
                    raise ValueError('orden: indique un entero desde 0')
                description = raw['descripcion_version']
                if description is not None and not isinstance(description, str):
                    raise ValueError('descripcion_version debe ser texto')
                values += [active, int(order), description.strip() if description else None, boolean(raw['version_publicada'], optional=True)]
                key = tuple(values[:3])
                if key in keys:
                    raise ValueError('regla duplicada por version + componente + codigo_regla: ' + ' / '.join(key))
                keys.add(key)
                rules.append(RuleUpload(row_number, tuple(values)))
            except ValueError as error:
                errors.append(f'Fila {row_number}: {error}')
        if errors:
            raise ValueError('\n'.join(errors[:30]) + ('\nHay más errores; corrija el archivo y vuelva a revisarlo.' if len(errors) > 30 else ''))
        if not rules:
            raise ValueError('La hoja Reglas está vacía. Diligencie al menos una fila; la hoja Ejemplo no se importa.')
        return rules
    finally:
        book.close()


def ensure_table(table):
    if table not in TABLES:
        raise ValueError('Seleccione una tabla de reglas admitida.')
    if sql is None:
        raise RuntimeError('Se requiere psycopg2 para consultar el repositorio.')


def fetch_existing(connection, table):
    ensure_table(table)
    with connection.cursor() as cursor:
        cursor.execute(sql.SQL('SELECT {} FROM public.{} ORDER BY version, orden NULLS LAST, componente, codigo_regla').format(
            sql.SQL(', ').join(sql.Identifier(column) for column in COLUMNS), sql.Identifier(table)))
        return [tuple(row) for row in cursor.fetchall()]


def make_plan(table, rules, existing):
    ensure_table(table)
    if not rules:
        raise ValueError('No hay reglas para incorporar.')
    keys = [rule.key for rule in rules]
    if len(keys) != len(set(keys)):
        raise ValueError('El archivo contiene reglas duplicadas.')
    major = re.search(r'_v(\d+)_', table)
    if major and any(not re.match(r'^v?' + major.group(1) + r'(?:[_.\s-]|$)', rule.values[0], re.I) for rule in rules):
        raise ValueError('La versión del archivo no corresponde a la tabla de destino V' + major.group(1) + '.')
    grouped = defaultdict(list)
    heads = {}
    for row in sorted(existing, key=lambda r: (r[0], r[7] is None, r[7] if r[7] is not None else 0, r[1], r[2])):
        grouped[row[:3]].append(row)
        heads.setdefault(row[0], row)
    metadata = {}
    for rule in rules:
        version = rule.values[0]
        if version in heads:
            metadata[version] = heads[version][8:10]
        else:
            descriptions = {r.values[8] for r in rules if r.values[0] == version and r.values[8] is not None}
            publications = {r.values[9] for r in rules if r.values[0] == version and r.values[9] is not None}
            if len(descriptions) > 1 or len(publications) > 1:
                raise ValueError('Los metadatos de la versión ' + version + ' no coinciden entre las filas.')
            metadata[version] = (next(iter(descriptions), ''), next(iter(publications), False))
    inserts, states = [], []
    for rule in rules:
        current = grouped.get(rule.key, [])
        if current:
            same = len(current) == 1 and current[0][:8] == rule.values[:8]
            states.append('Ya existe · no se modifica' if same else 'Existe con diferencias · no se modifica')
        else:
            inserts.append(rule.values[:8] + tuple(metadata[rule.values[0]]))
            states.append('Nueva')
    if table.startswith('reglas_cargue_snc_'):
        combined = list(existing) + inserts
        active_versions = {r[0] for r in combined if r[6]}
        publication = {version: head[9] for version, head in heads.items()}
        publication.update({version: values[1] for version, values in metadata.items()})
        if sum(bool(publication.get(version)) for version in active_versions) > 1:
            raise ValueError('Cargue SNC permite una sola versión publicada con reglas activas. La carga no despublica versiones existentes.')
    signature = hashlib.sha256(json.dumps(sorted(existing, key=repr), ensure_ascii=False, default=str).encode()).hexdigest()
    return UploadPlan(table, rules, signature, inserts, states)


def preview_import(connection, table, rules):
    return make_plan(table, rules, fetch_existing(connection, table))


def insert_import(connection, plan):
    """Caller owns transaction. Lock and recheck ensure the reviewed state still applies."""
    ensure_table(plan.table)
    with connection.cursor() as cursor:
        cursor.execute("SET LOCAL lock_timeout = '8s'")
        cursor.execute("SET LOCAL statement_timeout = '30s'")
        cursor.execute(sql.SQL('LOCK TABLE public.{} IN SHARE ROW EXCLUSIVE MODE').format(sql.Identifier(plan.table)))
    current = preview_import(connection, plan.table, plan.rules)
    if current.signature != plan.signature or current.inserts != plan.inserts:
        raise ValueError('El repositorio cambió después de la vista previa. Revise el archivo de nuevo antes de incorporarlo.')
    if not current.inserts:
        return 0
    with connection.cursor() as cursor:
        query = sql.SQL('INSERT INTO public.{} ({}) VALUES ({})').format(
            sql.Identifier(plan.table), sql.SQL(', ').join(sql.Identifier(c) for c in COLUMNS),
            sql.SQL(', ').join(sql.Placeholder() for _ in COLUMNS))
        cursor.executemany(query, current.inserts)
    return len(current.inserts)
