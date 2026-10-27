import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from qgis_control_calidad.review import build_cases, digest, load_review, restore_review, save_review
from qgis_control_calidad.report_exceptions import HEADERS, FIELDS, DEFAULTS


def cases():
    rule = SimpleNamespace(category='Administrativas', rule='r1', status='Correcta',
                           criterion_hash=digest('SELECT validacion'), identity_safe=True,
                           rows=[{'numero_predial': '00001', 'dato': 'valor'}])
    return build_cases([rule], {'model': 'CICA', 'schema': 'test'})


class ReportExceptionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'review.xlsx'

    def test_dt_and_operator_require_their_own_observation(self):
        for field, note in (('excepcion_dt', 'observacion_dt'), ('excepcion_operador', 'observacion_operador')):
            row = cases()[0]
            setattr(row, field, 'Sí')
            for value in ('', '  ', '\u200b'):
                setattr(row, note, value)
                with self.assertRaises(ValueError):
                    save_review(self.path, [row])
            setattr(row, note, 'Justificación válida')
            save_review(self.path, [row])

    def test_dt_keeps_operator_marks_and_notes_when_saving(self):
        row = cases()[0]
        row.excepcion_dt = 'Sí'; row.observacion_dt = 'Validado DT'
        for flag, note in (('Sí', 'Observación operador'), ('No', 'Comentario operador')):
            row.excepcion_operador, row.observacion_operador = flag, note
            save_review(self.path, [row])
            restored = cases()
            restore_review(restored, load_review(self.path))
            self.assertEqual(restored[0].exception_tracking(), row.exception_tracking())

    def test_export_fields_validation_protection_and_gray_operator(self):
        save_review(self.path, cases())
        book = load_workbook(self.path); sheet = book['Revisión']
        self.assertEqual(tuple(sheet.cell(1, col).value for col in range(9, 13)), HEADERS)
        self.assertEqual(tuple(sheet.cell(2, col).value or '' for col in range(9, 13)), DEFAULTS)
        self.assertEqual(sheet['I2'].fill.fgColor.rgb, '00FFC7CE')
        self.assertEqual(sheet['K2'].fill.fgColor.rgb, '00FFC7CE')
        self.assertEqual(sheet['I2'].font.color.rgb, '009C0006')
        self.assertTrue(sheet.protection.sheet)
        self.assertTrue(sheet['A2'].protection.locked)
        for col in ('I', 'J', 'K', 'L'):
            self.assertFalse(sheet[f'{col}2'].protection.locked)
        validations = list(sheet.data_validations.dataValidation)
        self.assertEqual(len(validations), 2)
        for rule in validations:
            self.assertEqual(rule.errorStyle, 'stop')
            self.assertTrue(rule.showErrorMessage)
            self.assertLessEqual(len(rule.formula1), 255)
        self.assertEqual([rule.type for rule in validations], ['list', 'list'])
        self.assertEqual([str(rule.sqref) for rule in validations], ['I2', 'K2'])
        self.assertTrue(all(rule.formula1 == '"Sí,No"' for rule in validations))
        self.assertGreater(len(sheet.conditional_formatting), 0)
        book.close()

    def test_domains_are_inline_lists_with_visible_dropdown_and_no_external_names(self):
        save_review(self.path, cases())
        book = load_workbook(self.path); sheet = book['Revisión']
        self.assertNotIn('_DominiosExcepcion', book.sheetnames)
        for validation in sheet.data_validations.dataValidation:
            self.assertEqual(validation.type, 'list')
            self.assertFalse(validation.showDropDown)
            self.assertEqual(validation.formula1, '"Sí,No"')
        book.close()

    def test_dt_hides_only_operator_cells_and_preserves_flag_colors(self):
        save_review(self.path, cases())
        book = load_workbook(self.path); sheet = book['Revisión']
        rules = [(str(area.sqref), rule) for area in sheet.conditional_formatting
                 for rule in sheet.conditional_formatting[area]]
        flags = [rule for area, rule in rules if area == 'I2']
        self.assertEqual({tuple(rule.formula): rule.dxf.fill.fgColor.rgb for rule in flags},
                         {('$I2="Sí"',): '00C6EFCE', ('$I2="No"',): '00FFC7CE'})
        self.assertEqual({tuple(rule.formula): rule.dxf.font.color.rgb for rule in flags},
                         {('$I2="Sí"',): '00006100', ('$I2="No"',): '009C0006'})
        gray = next(rule for area, rule in rules if area == 'K2:L2')
        self.assertTrue(gray.stopIfTrue)
        self.assertEqual(gray.dxf.fill.fgColor.rgb, '00FFFFFF')
        self.assertEqual(gray.dxf.font.color.rgb, '00FFFFFF')
        self.assertLess(gray.priority, min(rule.priority for area, rule in rules if area == 'K2'))
        self.assertFalse(any(area.startswith('A2:') for area, _rule in rules))
        book.close()

    def test_observation_columns_expand_to_content_and_wrap_long_text(self):
        current = cases()
        current[0].observacion_dt = 'Justificación detallada ' * 5
        current[0].observacion_operador = 'Texto operador ' * 50
        save_review(self.path, current)
        book = load_workbook(self.path); sheet = book['Revisión']
        self.assertGreater(sheet.column_dimensions['J'].width, 42)
        self.assertGreater(sheet.column_dimensions['L'].width, sheet.column_dimensions['J'].width)
        self.assertLessEqual(sheet.column_dimensions['L'].width, 255)
        self.assertTrue(sheet['J2'].alignment.wrap_text)
        self.assertTrue(sheet['L2'].alignment.wrap_text)
        self.assertGreater(sheet.row_dimensions[2].height, 30)
        self.assertEqual(sheet['L2'].value, current[0].observacion_operador)
        book.close()

    def test_recovery_preserves_dt_or_operator_only_when_case_unchanged(self):
        for flag, note in (('excepcion_dt', 'observacion_dt'), ('excepcion_operador', 'observacion_operador')):
            old = cases(); setattr(old[0], flag, 'Sí'); setattr(old[0], note, 'Texto de justificación')
            save_review(self.path, old)
            new = cases(); counts = restore_review(new, load_review(self.path))
            self.assertEqual(counts['incorporadas'], 1)
            self.assertEqual(new[0].exception_tracking(), old[0].exception_tracking())

    def test_import_invalid_observation_and_operator_block_are_reported(self):
        save_review(self.path, cases())
        book = load_workbook(self.path); sheet = book['Revisión']
        sheet['I2'] = 'Sí'; sheet['J2'] = 'Justificación DT'
        sheet['K2'] = 'Sí'; sheet['L2'] = 'Intento operador'
        book.save(self.path); book.close()
        new = cases(); counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['operador_bloqueado'], 0)
        self.assertEqual(counts['operador_aprobadas_dt'], 1)
        self.assertEqual(new[0].exception_tracking(), ('Sí', 'Justificación DT', 'Sí', 'Intento operador'))
        book = load_workbook(self.path); sheet = book['Revisión']
        sheet['J2'] = ' '; sheet['K2'] = 'No'; sheet['L2'] = ''
        book.save(self.path); book.close()
        new = cases(); counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['excepciones_invalidas'], 1)
        self.assertEqual(new[0].excepcion_dt, 'No')

    def test_existing_operator_observation_is_not_silently_replaced(self):
        old = cases(); old[0].excepcion_operador = 'Sí'; old[0].observacion_operador = 'Anterior'
        save_review(self.path, old)
        new = cases(); new[0].observacion_operador = 'Actual'
        counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['conflictos'], 1)
        self.assertEqual(new[0].observacion_operador, 'Actual')


if __name__ == '__main__':
    unittest.main()
