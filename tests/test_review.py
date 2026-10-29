import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from qgis_control_calidad.review import (
    ADJUST, EXCEPTION, PENDING, build_cases, digest, load_review, restore_review, save_review, review_statistics,
)

CONTEXT = {'model': 'CICA', 'environment': 'Validaciones del modelo', 'database': 'test', 'schema': 'municipio'}


def result(rows, rule='regla_1', sql='SELECT id, numero_predial, dato FROM predio', safe=True):
    return SimpleNamespace(category='Administrativas', rule=rule, status='Correcta', rows=rows,
                           criterion_hash=digest(sql), identity_safe=safe, findings=len(rows),
                           detail='', description='Descripción del control')


def cases(rows, **kwargs):
    return build_cases([result(rows, **kwargs)], CONTEXT)


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'revision.xlsx'
        self.rows = [{'id': 2, 'numero_predial': '000001', 'dato': Decimal('12.30')},
                     {'id': 3, 'numero_predial': '000002', 'dato': Decimal('15.10')}]

    def test_reordering_rows_and_columns_preserves_review_and_date(self):
        old = cases(self.rows)
        old[0].set_review(EXCEPTION, 'Caso justificado', '2026-10-26T09:15:00-05:00')
        save_review(self.path, old)
        new = cases([dict(reversed(list(row.items()))) for row in reversed(self.rows)])
        counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(new[1].tracking(), old[0].tracking())
        self.assertEqual(new[0].estado_revision_usuario, PENDING)

    def test_reordering_excel_rows_and_fragmented_identity_preserves_review(self):
        rows = [dict(self.rows[0], a='x' * 20000, b='y' * 20000), self.rows[1]]
        old = cases(rows); old[0].set_review(EXCEPTION, 'Identidad completa')
        save_review(self.path, old)
        book = load_workbook(self.path); sheet = book['Revisión']
        values = list(sheet.iter_rows(min_row=2, values_only=True))
        for index, values_row in enumerate(reversed(values), 2):
            for column, value in enumerate(values_row, 1):
                sheet.cell(index, column, value)
        self.assertGreater(book['_Identidades'].max_row, 3)
        book.save(self.path); book.close()
        new = cases(rows)
        counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(new[0].observacion_usuario, 'Identidad completa')

    def test_legacy_matricula_group_can_recover_without_npn(self):
        old = cases([{'matricula': '00104749', 'conteo': 2, 'cond_1_12': 'Fail'}], rule='regla_1_12')
        self.assertTrue(old[0].recoverable)
        old[0].excepcion_dt = 'Sí'
        old[0].observacion_dt = 'Matrícula verificada por DT'
        save_review(self.path, old)
        book = load_workbook(self.path)
        book['Revisión']['H2'] = 'No'  # Previous versions did not recognize matrícula.
        book.save(self.path); book.close()
        current = cases([{'matricula': '00104749', 'conteo': 2, 'cond_1_12': 'Fail'}], rule='regla_1_12')
        counts = restore_review(current, load_review(self.path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(current[0].excepcion_dt, 'Sí')
        changed = cases([{'matricula': '00104749', 'conteo': 3, 'cond_1_12': 'Fail'}], rule='regla_1_12')
        self.assertEqual(restore_review(changed, load_review(self.path))['incorporadas'], 0)

    def test_statistics_deduct_valid_exceptions_without_changing_automatic_errors(self):
        rows = self.rows + [dict(self.rows[0], id=4, numero_predial='000003')]
        automatic = result(rows)
        error = result([], rule='error_tecnico'); error.status = 'Error'; error.findings = '—'
        current = cases(rows)
        current[0].excepcion_dt = 'Sí'; current[0].observacion_dt = 'DT verificada'
        current[1].excepcion_operador = 'Sí'; current[1].observacion_operador = 'Operador verificado'
        groups, totals = review_statistics([automatic, error], current)
        self.assertEqual(totals, dict(rules=2, findings=3, pending=2, dt=1, operator=1, errors=1))
        self.assertEqual(automatic.findings, 3)
        self.assertEqual(error.status, 'Error')
        save_review(self.path, current, [automatic, error])
        book = load_workbook(self.path)
        stats = book['Estadística de revisión']
        self.assertEqual(list(stats.values)[-1], ('TOTAL', 2, 3, 1, 1, 2, 1))
        book.close()

    def test_operator_exceptions_are_visible_but_do_not_discount_errors(self):
        current = cases(self.rows)
        current[0].excepcion_operador = 'Sí'; current[0].observacion_operador = 'Propuesta'
        _, totals = review_statistics([result(self.rows)], current)
        self.assertEqual(totals['operator'], 1)
        self.assertEqual(totals['dt'], 0)
        self.assertEqual(totals['pending'], 2)

    def test_statistics_ignore_unjustified_exceptions_and_do_not_count_twice(self):
        current = cases(self.rows)
        current[0].excepcion_dt = 'Sí'; current[0].observacion_dt = 'Justificada DT'
        current[0].excepcion_operador = 'Sí'; current[0].observacion_operador = 'También operador'
        current[1].excepcion_dt = 'Sí'; current[1].observacion_dt = '  '
        _, totals = review_statistics([result(self.rows)], current)
        self.assertEqual(totals['dt'], 1)
        self.assertEqual(totals['operator'], 1)
        self.assertEqual(totals['pending'], 1)

    def test_relevant_value_sql_rule_category_or_context_change_invalidates(self):
        old = cases(self.rows[:1])
        old[0].set_review(EXCEPTION, 'Justificada')
        save_review(self.path, old)
        records = load_review(self.path)
        changed = [cases([dict(self.rows[0], dato=Decimal('12.31'))]),
                   cases(self.rows[:1], sql='SELECT otro_criterio'),
                   cases(self.rows[:1], rule='regla_2'),
                   build_cases([result(self.rows[:1])], dict(CONTEXT, schema='otro'))]
        r = result(self.rows[:1]); r.category = 'Juridicas'
        changed.append(build_cases([r], CONTEXT))
        for current in changed:
            counts = restore_review(current, records)
            self.assertEqual(counts['incorporadas'], 0)
            self.assertEqual(counts['sin_correspondencia'], 1)
            self.assertEqual(current[0].estado_revision_usuario, PENDING)

    def test_duplicate_source_or_current_is_ambiguous_even_if_same_review(self):
        old = cases(self.rows[:1]); old[0].set_review(EXCEPTION, 'Justificada')
        save_review(self.path, old)
        records = load_review(self.path)
        counts = restore_review(cases([self.rows[0], self.rows[0]]), records)
        self.assertEqual(counts['ambiguas'], 1)
        counts = restore_review(cases(self.rows[:1]), records * 2)
        self.assertEqual(counts['ambiguas'], 2)
        self.assertEqual(counts['incorporadas'], 0)

    def test_position_fields_do_not_identify_cases(self):
        a = cases([dict(self.rows[0], row_number=1)])
        b = cases([dict(self.rows[0], row_number=89)])
        self.assertEqual(a[0].key, b[0].key)
        a[0].set_review(EXCEPTION, 'Observación')
        save_review(self.path, a)
        self.assertEqual(restore_review(b, load_review(self.path))['incorporadas'], 1)

    def test_missing_object_id_or_sql_or_duplicate_sql_columns_not_recoverable(self):
        for current in (cases([{'cantidad': 5}]), cases([{'codigo_regla': 'regla_1', 'cantidad': 5}]), cases(self.rows[:1], safe=False)):
            self.assertFalse(current[0].recoverable)
        r = result(self.rows[:1]); r.criterion_hash = ''
        self.assertFalse(build_cases([r], CONTEXT)[0].recoverable)
        self.assertTrue(cases([{'t_ili_tid': 'stable-001', 'dato': 'valor'}])[0].recoverable)

    def test_exception_requires_visible_observation_at_edit_and_save(self):
        current = cases(self.rows[:1])
        for text in ('', '  \t\n', '\u200b\ufeff'):
            with self.assertRaises(ValueError):
                current[0].set_review(EXCEPTION, text)
            current[0].estado_revision_usuario = EXCEPTION
            current[0].observacion_usuario = text
            with self.assertRaises(ValueError):
                save_review(self.path, current)
            self.assertFalse(self.path.exists())
            current[0].estado_revision_usuario = PENDING

    def test_invalid_exception_in_excel_becomes_pending_and_is_reported(self):
        old = cases(self.rows[:1]); old[0].set_review(EXCEPTION, 'Texto válido')
        save_review(self.path, old)
        book = load_workbook(self.path)
        book['Revisión']['F2'] = '   '
        book.save(self.path); book.close()
        new = cases(self.rows[:1])
        counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['excepciones_invalidas'], 1)
        self.assertEqual(new[0].estado_revision_usuario, PENDING)

    def test_existing_tracking_is_not_silently_overwritten(self):
        old = cases(self.rows[:1]); old[0].set_review(EXCEPTION, 'Anterior')
        save_review(self.path, old)
        new = cases(self.rows[:1]); new[0].set_review(ADJUST, 'Observación actual')
        previous = new[0].tracking()
        counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['conflictos'], 1)
        self.assertEqual(new[0].tracking(), previous)

    def test_all_rows_text_identifiers_formula_safety_and_automatic_summary(self):
        data = [dict(self.rows[0], nota='=HYPERLINK("url")'), self.rows[1]]
        current = cases(data)
        save_review(self.path, current, [result(data), result([], rule='sin_hallazgos')])
        book = load_workbook(self.path)
        sheet = book['Revisión']
        self.assertEqual(sheet.max_row, 3)
        columns = {cell.value: cell.column for cell in sheet[1]}
        identifier = sheet.cell(2, columns['Numero_Predial_Nacional'])
        self.assertEqual(identifier.value, '000001')
        self.assertEqual(identifier.data_type, 's')
        self.assertEqual(sheet.cell(2, columns['dato:nota']).data_type, 's')
        self.assertEqual(book['Reglas & Componente'].max_row, 3)
        book.close()
        self.assertEqual(len(load_review(self.path)), 2)

    def test_excel_changes_to_case_data_prevent_recovery(self):
        old = cases(self.rows[:1]); old[0].set_review(EXCEPTION, 'Observación')
        save_review(self.path, old)
        book = load_workbook(self.path); sheet = book['Revisión']
        cols = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, cols['Numero_Predial_Nacional'], '999999')
        book.save(self.path); book.close()
        new = cases(self.rows[:1]); counts = restore_review(new, load_review(self.path))
        self.assertEqual(counts['datos_invalidos'], 1)
        self.assertEqual(new[0].estado_revision_usuario, PENDING)

    def test_failed_save_does_not_replace_existing_file(self):
        old = cases(self.rows[:1]); save_review(self.path, old)
        original = self.path.read_bytes()
        bad = cases([dict(self.rows[0], dato='x' * 33000)])
        with self.assertRaises(ValueError):
            save_review(self.path, bad)
        self.assertEqual(self.path.read_bytes(), original)

    def test_canonical_nested_dates_and_binary_values_roundtrip(self):
        rows = [dict(self.rows[0], fecha=date(2026, 10, 26), blob=b'\x00\x12', extra={'z': [Decimal('1.20'), '001']})]
        old = cases(rows); old[0].set_review(ADJUST, 'Ajustar')
        save_review(self.path, old)
        new = cases(rows)
        self.assertEqual(restore_review(new, load_review(self.path))['incorporadas'], 1)


if __name__ == '__main__':
    unittest.main()
