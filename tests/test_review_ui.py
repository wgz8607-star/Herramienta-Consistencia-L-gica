import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from qgis.core import QgsApplication
from qgis.PyQt.QtGui import QCloseEvent
from qgis.PyQt.QtWidgets import QMessageBox, QPushButton
from openpyxl import load_workbook
from qgis_control_calidad import dialog as module
from qgis_control_calidad.review import build_cases, digest
from qgis_control_calidad.report_exceptions import HEADERS

app = QgsApplication.instance() or QgsApplication([], False)
app.initQgis()


class ReviewUiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.owner = module.ControlCalidadDialog()
        self.addCleanup(self.cleanup)
        self.result = module.RuleResult('regla_1', 'Administrativas', 'Correcta', 2,
                                       rows=[{'numero_predial': '00001', 'valor': 'dato A'},
                                             {'numero_predial': '00002', 'valor': 'dato B'}],
                                       criterion_hash=digest('SELECT criterio'))
        self.owner.results = [self.result]
        self.owner.review_context = {'model': 'CICA', 'schema': 'test'}
        self.owner.review_cases = build_cases(self.owner.results, self.owner.review_context)
        self.owner.finished()

    def cleanup(self):
        self.owner.review_dirty = False
        self.owner.close()

    def test_excel_task_keeps_event_loop_alive(self):
        from threading import Event
        from qgis.PyQt.QtCore import QTimer
        released = Event()
        QTimer.singleShot(20, released.set)
        self.assertTrue(self.owner.report_task(released.wait, 1))

    def test_export_failure_restores_controls_and_retains_diagnostic(self):
        partial = Path(self.directory.name) / 'incompleto.xlsx'
        def fail():
            self.owner.exporting = True
            self.owner.left_tabs.setEnabled(False)
            self.owner.export_button.setEnabled(False)
            partial.write_bytes(b'partial')
            self.owner.export_created_paths.append(partial)
            raise PermissionError('Archivo bloqueado por Excel')
        with patch.object(self.owner, 'build_detail_reports', side_effect=fail), patch.object(QMessageBox, 'exec_', return_value=QMessageBox.Ok):
            self.owner.export_excel()
        self.assertFalse(self.owner.exporting)
        self.assertTrue(self.owner.left_tabs.isEnabled())
        self.assertTrue(self.owner.export_button.isEnabled())
        self.assertIn('Archivo bloqueado por Excel', self.owner.incidents_browser.toPlainText())
        self.assertIn('No se generó', self.owner.status.text())
        self.assertFalse(partial.exists())
        self.assertEqual(len(self.owner.review_cases), 2)

    def test_record_review_window_is_removed(self):
        self.assertFalse(hasattr(self.owner, 'review_viewer'))
        self.assertFalse(hasattr(self.owner, 'open_review_viewer'))
        self.assertNotIn('Revisar hallazgos', [button.text() for button in self.owner.findChildren(QPushButton)])

    def test_each_export_has_four_editable_exception_fields(self):
        self.owner.review_cases[0].excepcion_dt = 'Sí'
        self.owner.review_cases[0].observacion_dt = 'Documentado por DT'
        self.owner.execution_info['Identificador de entrega'] = 'no_exportar'
        self.owner.execution_info['Versión plugin'] = 'version-ejecucion'
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        files = list(Path(self.directory.name).glob('*.xlsx'))
        self.assertEqual(len(files), 2)
        for path in files:
            book = load_workbook(path)
            self.assertNotIn('Sheet', book.sheetnames)
            self.assertNotIn('Estadística de revisión', book.sheetnames)
            self.assertTrue(book.security.lockStructure)
            for report_sheet in book:
                if report_sheet.sheet_state == 'visible':
                    self.assertTrue(report_sheet.protection.sheet)
                    self.assertFalse(report_sheet.protection.formatColumns)
                    self.assertTrue(report_sheet.protection.formatRows)
                    self.assertTrue(report_sheet.protection.insertRows)
                    self.assertTrue(report_sheet.protection.deleteColumns)
                    self.assertTrue(report_sheet.freeze_panes.startswith('B'))
            if 'regla_1' in book:
                self.assertEqual(book['regla_1'].freeze_panes, 'B5')
                self.assertTrue(book['regla_1']['A5'].protection.locked)
            self.assertNotIn('Contexto de validación', book.sheetnames)
            self.assertNotIn('Contexto', book.sheetnames)
            if 'regla_1' in book.sheetnames:
                self.assertTrue(book['Detalle Validación']['A1'].value.startswith('Detalle Validación'))
                self.assertNotIn('Versión plugin', [cell.value for cell in book['Detalle Validación']['A']])
            if 'regla_1' not in book.sheetnames:
                self.assertNotIn('Revisión', book.sheetnames)
                self.assertNotIn('_Identidades', book.sheetnames)
                self.assertIn('Resumen ejecutivo', book.sheetnames)
                summary = book['Resumen ejecutivo']
                version_row = next(cell.row for cell in summary['A'] if cell.value == 'Versión plugin')
                self.assertEqual(summary.cell(version_row, 2).value, 'version-ejecucion')
                header = next(cell.row for cell in summary['A'] if cell.value == 'Componente')
                metadata = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(2, header)}
                self.assertEqual(metadata['Modelo'], 'CICA')
                self.assertEqual(metadata['Esquema'], 'test')
                self.assertIn('Ubicación territorial', metadata)
                self.assertNotIn('Identificador de entrega', metadata)
                self.assertEqual(summary.freeze_panes, 'B' + str(header + 1))
                self.assertTrue(summary.auto_filter.ref.startswith('A' + str(header) + ':'))
                self.assertGreater(version_row, header)
                book.close()
                continue
            self.assertIn('Revisión', book.sheetnames)
            self.assertEqual(book['Revisión'].sheet_state, 'hidden')
            sheet = book['Revisión']
            self.assertEqual(sheet.max_row, 3)
            for header in HEADERS:
                self.assertIn(header, [cell.value for cell in sheet[1]])
            if 'regla_1' not in book.sheetnames:
                self.assertEqual(sheet['I2'].value, 'Sí')
                self.assertEqual(sheet['J2'].value, 'Documentado por DT')
            self.assertTrue(sheet.protection.sheet)
            if 'regla_1' in book.sheetnames:
                detail = book['regla_1']
                cols = {cell.value: cell.column for cell in detail[4]}
                self.assertEqual(detail.cell(5, cols[HEADERS[0]]).value, 'Sí')
                self.assertEqual(detail.cell(5, cols[HEADERS[0]]).fill.fgColor.rgb, '00C6EFCE')
                self.assertEqual(detail.cell(5, cols[HEADERS[0]]).font.color.rgb, '00006100')
                self.assertEqual(detail.cell(6, cols[HEADERS[0]]).fill.fgColor.rgb, '00FFC7CE')
                self.assertEqual(detail.cell(5, cols[HEADERS[1]]).value, 'Documentado por DT')
                for header in HEADERS:
                    self.assertFalse(detail.cell(5, cols[header]).protection.locked)
                self.assertTrue(sheet['I2'].data_type == 'f')
                self.assertTrue(sheet['I2'].protection.locked)
                self.assertTrue(detail.protection.sheet)
                self.assertEqual(len(detail.data_validations.dataValidation), 2)
                for validation in detail.data_validations.dataValidation:
                    self.assertEqual(validation.type, 'list')
                    self.assertFalse(validation.showDropDown)
                    self.assertEqual(validation.formula1, '"Sí,No"')
            book.close()
        self.assertEqual(self.result.status, 'Correcta')

    def test_component_report_can_recover_its_operator_review(self):
        from qgis_control_calidad.review import load_review, restore_review
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        package = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(package)
        sheet = book['regla_1']
        cols = {cell.value: cell.column for cell in sheet[4]}
        sheet.cell(5, cols['Excepción operador'], 'Sí')
        sheet.cell(5, cols['Observación operador'], 'Justificación del operador')
        book.save(package); book.close()
        new = build_cases(self.owner.results, self.owner.review_context)
        counts = restore_review(new, load_review(package))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(new[0].excepcion_operador, 'Sí')
        self.assertEqual(new[0].observacion_operador, 'Justificación del operador')

    def test_rule_sheet_reordering_does_not_change_recovery(self):
        from qgis_control_calidad.review import load_review, restore_review
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path); sheet = book['regla_1']
        columns = {cell.value: cell.column for cell in sheet[4]}
        sheet.cell(5, columns['Excepción DT'], 'Sí')
        sheet.cell(5, columns['Observación DT'], 'Justificado DT')
        rows = list(sheet.iter_rows(min_row=5, max_row=6, values_only=True))
        for index, values in enumerate(reversed(rows), 5):
            for column, value in enumerate(values, 1):
                sheet.cell(index, column).value = value
        book.save(path); book.close()
        current = build_cases(self.owner.results, self.owner.review_context)
        counts = restore_review(current, load_review(path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(current[0].excepcion_dt, 'Sí')
        self.assertEqual(current[1].excepcion_dt, 'No')

    def test_modified_rule_data_prevents_recovery(self):
        from qgis_control_calidad.review import load_review, restore_review
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path); sheet = book['regla_1']
        columns = {cell.value: cell.column for cell in sheet[4]}
        sheet.cell(5, columns['Excepción DT'], 'Sí')
        sheet.cell(5, columns['Observación DT'], 'Justificado DT')
        sheet.cell(5, columns['Numero_Predial_Nacional'], 'Alterado')
        book.save(path); book.close()
        current = build_cases(self.owner.results, self.owner.review_context)
        counts = restore_review(current, load_review(path))
        self.assertEqual(counts['datos_invalidos'], 1)
        self.assertEqual(counts['incorporadas'], 0)

    def test_excel_normalized_consolidation_formulas_do_not_block_source_review(self):
        from qgis_control_calidad.review import load_review, restore_review
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path); sheet = book['regla_1']
        columns = {cell.value: cell.column for cell in sheet[4]}
        sheet.cell(5, columns['Excepción DT'], 'Sí')
        sheet.cell(5, columns['Observación DT'], 'Justificación DT')
        for col in range(9, 13):
            cell = book['Revisión'].cell(2, col)
            cell.value = cell.value.replace("'regla_1'!", 'regla_1!')
        book.save(path); book.close()
        current = build_cases(self.owner.results, self.owner.review_context)
        counts = restore_review(current, load_review(path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(current[0].excepcion_dt, 'Sí')

    def test_five_detail_buttons_load_and_reload_independently(self):
        from qgis_control_calidad.review import save_review
        components = ('Administrativas', 'Economicas', 'Estructura', 'Fisicas', 'Juridicas')
        self.owner.results = [module.RuleResult('r1', c, 'Correcta', 1,
            rows=[{'numero_predial': '00001', 'dato': c}], criterion_hash=digest(c)) for c in components]
        self.owner.review_cases = build_cases(self.owner.results, self.owner.review_context)
        self.owner.finished()
        self.assertEqual(set(self.owner.detail_review_rows), set(components))
        paths = {}
        for index, component in enumerate(components):
            path = str(Path(self.directory.name) / (component + '.xlsx'))
            cases = build_cases([self.owner.results[index]], self.owner.review_context)
            cases[0].excepcion_dt = 'Sí'; cases[0].observacion_dt = 'DT: ' + component
            save_review(path, cases)
            paths[component] = path
            load, state = self.owner.detail_review_rows[component]
            self.assertIn('No cargado', state.text())
            with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(path, 'Excel')), patch.object(module.QMessageBox, 'information'):
                load.click()
            self.assertIn('Cargado', state.text())
        self.assertTrue(all(case.excepcion_dt == 'Sí' for case in self.owner.review_cases))
        component = 'Administrativas'
        book = load_workbook(paths[component]); book['Revisión']['I2'] = 'No'; book['Revisión']['J2'] = ''
        book.save(paths[component]); book.close()
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(paths[component], 'Excel')), patch.object(module.QMessageBox, 'information'):
            self.owner.detail_review_rows[component][0].click()
        self.assertEqual(self.owner.review_cases[0].excepcion_dt, 'No')
        self.assertTrue(all(case.excepcion_dt == 'Sí' for case in self.owner.review_cases[1:]))
        self.assertIn('Cargado', self.owner.detail_review_rows[component][1].text())
        # The file is not held open and can be edited and saved in place.
        book = load_workbook(paths[component]); book['Revisión']['I2'] = 'Sí'; book['Revisión']['J2'] = 'DT corregida'
        book.save(paths[component]); book.close()
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(paths[component], 'Excel')), patch.object(module.QMessageBox, 'information'):
            self.owner.detail_review_rows[component][0].click()
        self.assertEqual(self.owner.review_cases[0].observacion_dt, 'DT corregida')

    def test_failed_detail_load_is_visible_without_losing_loaded_review(self):
        from qgis_control_calidad.review import save_review
        path = str(Path(self.directory.name) / 'admin.xlsx')
        cases = build_cases(self.owner.results, self.owner.review_context)
        cases[0].excepcion_dt = 'Sí'; cases[0].observacion_dt = 'DT válida'
        save_review(path, cases)
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(path, 'Excel')), patch.object(module.QMessageBox, 'information'):
            self.owner.recover_detail('Administrativas')
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(path + '.missing', 'Excel')), patch.object(module.QMessageBox, 'warning'):
            self.owner.recover_detail('Administrativas')
        self.assertIn('Error de carga', self.owner.detail_review_rows['Administrativas'][1].text())
        self.assertEqual(self.owner.review_cases[0].observacion_dt, 'DT válida')

    def test_statistics_refresh_on_load_reload_clear_and_discard(self):
        from qgis_control_calidad.review import save_review
        path = str(Path(self.directory.name) / 'admin.xlsx')
        cases = build_cases(self.owner.results, self.owner.review_context)
        cases[0].excepcion_dt = 'Sí'; cases[0].observacion_dt = 'DT válida'
        save_review(path, cases)
        self.assertEqual(self.owner.table.item(0, 5).text(), '2')
        def load():
            with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(path, 'Excel')), patch.object(module.QMessageBox, 'information'):
                self.owner.recover_detail('Administrativas')
        load()
        self.assertEqual(self.owner.table.item(0, 2).text(), '2')
        self.assertEqual(self.owner.table.item(0, 5).text(), '1')
        self.assertEqual(self.owner.table.item(0, 3).text(), '1')
        self.assertIn('1 errores', self.owner.summary.text())
        cases[1].excepcion_operador = 'Sí'; cases[1].observacion_operador = 'Operador válida'
        save_review(path, cases)
        load()
        self.assertEqual(self.owner.table.item(0, 5).text(), '1')
        self.assertEqual(self.owner.table.item(0, 3).text(), '1')
        self.assertEqual(self.owner.table.item(0, 4).text(), '1')
        for case in cases:
            case.excepcion_dt = 'No'; case.observacion_dt = ''
            case.excepcion_operador = 'No'; case.observacion_operador = ''
        save_review(path, cases)
        load()
        self.assertEqual(self.owner.table.item(0, 5).text(), '2')
        self.assertEqual(self.owner.table.item(0, 3).text(), '0')
        cases[0].excepcion_dt = 'Sí'; cases[0].observacion_dt = 'DT válida'
        save_review(path, cases)
        load()
        with patch.object(module.QMessageBox, 'question', return_value=QMessageBox.Discard):
            self.owner.confirm_review_changes()
        self.assertEqual(self.owner.table.item(0, 5).text(), '2')
        self.assertEqual(self.owner.table.item(0, 3).text(), '0')
        self.assertEqual(self.owner.table.item(0, 6).text(), '0')

    def test_manual_connection_fields_are_not_clipped_after_mode_changes(self):
        self.owner.show()
        app.processEvents()
        for mode in (1, 0, 1):
            self.owner.connection_mode.setCurrentIndex(mode)
            app.processEvents()
            app.processEvents()
            page = self.owner.connection_stack.currentWidget()
            self.assertGreaterEqual(page.height(), page.layout().minimumSize().height())
            controls = self.owner.schema_controls
            self.assertGreaterEqual(controls.height(), controls.minimumSizeHint().height())
            self.assertLessEqual(controls.y() + controls.height(), page.height())
            self.assertTrue(self.owner.schema.isVisible())

    def test_minimize_keeps_reviews_and_pending_changes(self):
        from qgis.PyQt.QtCore import Qt
        self.assertTrue(self.owner.windowFlags() & Qt.WindowMinimizeButtonHint)
        self.owner.review_cases[0].observacion_dt = 'No perder al minimizar'
        self.owner.review_dirty = True
        self.owner.show()
        self.owner.showMinimized()
        app.processEvents()
        self.assertTrue(self.owner.isMinimized())
        self.assertEqual(self.owner.review_cases[0].observacion_dt, 'No perder al minimizar')
        self.assertTrue(self.owner.review_dirty)
        self.owner.showNormal()
        app.processEvents()
        self.assertFalse(self.owner.isMinimized())

    def test_plugin_action_restores_minimized_window(self):
        from qgis_control_calidad.plugin import ControlCalidadPlugin
        plugin = ControlCalidadPlugin(None)
        plugin.dialog = MagicMock()
        plugin.dialog.isMinimized.return_value = True
        plugin.open_dialog()
        plugin.dialog.showNormal.assert_called_once()
        plugin.dialog.show.assert_not_called()
        plugin.dialog.activateWindow.assert_called_once()

    def test_detail_import_warns_only_dt_counts_and_operator_does_not_discount(self):
        from qgis_control_calidad.review import save_review
        path = str(Path(self.directory.name) / 'operator.xlsx')
        current = build_cases(self.owner.results, self.owner.review_context)
        current[0].excepcion_operador = 'Sí'; current[0].observacion_operador = 'Propuesta operador'
        save_review(path, current)
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(path, 'Excel')), patch.object(module.QMessageBox, 'information') as notice:
            self.owner.recover_detail('Administrativas')
        self.assertIn(module.EXCEPTION_NOTICE, notice.call_args.args[2])
        self.assertEqual(module.EXCEPTION_NOTICE, 'Solo las excepciones validadas y marcadas por DT descuentan errores.')
        self.assertEqual(self.owner.table.item(0, 4).text(), '1')
        self.assertEqual(self.owner.table.item(0, 5).text(), '2')

    def test_recovery_summary_has_separate_role_counts_and_incidents(self):
        from qgis_control_calidad.review import restore_review
        counts = restore_review([], [])
        counts.update(dt_incorporadas=4, dt_sin_observacion=1,
                      operador_incorporadas=2, operador_sin_observacion=3,
                      operador_aprobadas_dt=1, operador_bloqueado=2)
        message = self.owner.recovery_message(counts)
        self.assertTrue(message.startswith('Resumen de excepciones'))
        self.assertIn('DT\n• Incorporadas: 4\n• Descartadas sin observación: 1', message)
        self.assertIn('Operador\n• Incorporadas: 2\n• Descartadas sin observación: 3', message)
        self.assertIn('• Aprobadas por DT: 1', message)
        self.assertIn('Incidencias de carga', message)
        self.assertNotIn('omitidos por prioridad DT', message)
        self.assertTrue(message.endswith(module.EXCEPTION_NOTICE))

    def test_operator_followup_survives_import_and_new_detail_export(self):
        from qgis_control_calidad.review import load_review, restore_review
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        original = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(original); sheet = book['regla_1']
        cols = {cell.value: cell.column for cell in sheet[4]}
        sheet.cell(5, cols['Excepción DT'], 'Sí')
        sheet.cell(5, cols['Observación DT'], 'Validada DT')
        for row, note in ((5, 'Propuesta aprobada'), (6, 'Propuesta pendiente')):
            sheet.cell(row, cols['Excepción operador'], 'Sí')
            sheet.cell(row, cols['Observación operador'], note)
        book.save(original); book.close()
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(str(original), 'Excel')), patch.object(module.QMessageBox, 'information'):
            self.owner.recover_detail('Administrativas')
        self.assertEqual(self.owner.review_cases[0].excepcion_operador, 'Sí')
        self.assertEqual(self.owner.review_cases[0].observacion_operador, 'Propuesta aprobada')
        self.assertEqual(self.owner.table.item(0, 3).text(), '1')
        self.assertEqual(self.owner.table.item(0, 4).text(), '2')
        self.assertEqual(self.owner.table.item(0, 5).text(), '1')
        output = Path(self.directory.name) / 'nueva_entrega'
        output.mkdir()
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=str(output)):
            self.owner.export_excel()
        regenerated = next(output.glob('detalle_Administrativas_*.xlsx'))
        fresh = build_cases(self.owner.results, self.owner.review_context)
        counts = restore_review(fresh, load_review(regenerated))
        self.assertEqual(counts['operador_incorporadas'], 2)
        self.assertEqual(counts['operador_aprobadas_dt'], 1)
        self.assertEqual(fresh[0].exception_tracking(), ('Sí', 'Validada DT', 'Sí', 'Propuesta aprobada'))
        self.assertEqual(fresh[1].exception_tracking(), ('No', '', 'Sí', 'Propuesta pendiente'))

    def test_modified_detail_blocks_export_until_reloaded(self):
        from qgis_control_calidad.review import save_review
        path = Path(self.directory.name) / 'review.xlsx'
        current = build_cases(self.owner.results, self.owner.review_context)
        current[0].excepcion_dt = 'Sí'; current[0].observacion_dt = 'Inicial'
        save_review(path, current)
        def load():
            with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(str(path), 'Excel')), patch.object(module.QMessageBox, 'information'):
                self.owner.recover_detail('Administrativas')
        load()
        self.assertTrue(self.owner.export_button.isEnabled())
        self.assertNotEqual(self.owner.detail_review_table.item(0, 2).text(), '—')
        book = load_workbook(path); book['Revisión']['J2'] = 'Modificada'
        book.save(path); book.close()
        self.owner.check_review_files(force=True)
        self.assertFalse(self.owner.export_button.isEnabled())
        self.assertIn('Archivo modificado', self.owner.detail_review_rows['Administrativas'][1].text())
        self.assertFalse(self.owner.review_file_alert.isHidden())
        with patch.object(module.QFileDialog, 'getExistingDirectory') as chooser, patch.object(module.QMessageBox, 'warning'):
            self.owner.export_excel()
        chooser.assert_not_called()
        self.assertEqual(self.owner.review_cases[0].observacion_dt, 'Inicial')
        load()
        self.assertEqual(self.owner.review_cases[0].observacion_dt, 'Modificada')
        self.assertTrue(self.owner.export_button.isEnabled())
        self.assertTrue(self.owner.review_file_alert.isHidden())

    def test_missing_detail_preserves_review_but_blocks_export(self):
        from qgis_control_calidad.review import save_review
        path = Path(self.directory.name) / 'review.xlsx'
        current = build_cases(self.owner.results, self.owner.review_context)
        current[0].observacion_operador = 'Conservar'
        save_review(path, current)
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(str(path), 'Excel')), patch.object(module.QMessageBox, 'information'):
            self.owner.recover_detail('Administrativas')
        path.unlink()
        self.owner.check_review_files(force=True)
        self.assertIn('no disponible', self.owner.detail_review_rows['Administrativas'][1].text())
        self.assertEqual(self.owner.review_cases[0].observacion_operador, 'Conservar')
        self.assertFalse(self.owner.export_button.isEnabled())

    def test_total_row_and_kpis_match_original_and_dt_only_statistics(self):
        self.owner.review_cases[0].excepcion_dt = 'Sí'; self.owner.review_cases[0].observacion_dt = 'DT'
        self.owner.review_cases[1].excepcion_operador = 'Sí'; self.owner.review_cases[1].observacion_operador = 'Operador'
        self.owner.refresh_review_statistics()
        last = self.owner.table.rowCount()-1
        self.assertEqual(self.owner.table.item(last, 0).text(), 'TOTAL')
        self.assertEqual(self.owner.table.item(last, 5).text(), '1')
        self.assertEqual(self.owner.kpi_values['findings'].text(), '2')
        self.assertEqual(self.owner.kpi_values['dt'].text(), '1')
        self.assertEqual(self.owner.kpi_values['pending'].text(), '1')

    def test_exported_observations_are_wide_literal_text(self):
        note = '= Texto literal con tildes y justificación extensa ' * 3
        self.owner.review_cases[0].observacion_dt = note
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path); sheet = book['regla_1']
        columns = {cell.value: cell.column for cell in sheet[4]}
        cell = sheet.cell(5, columns['Observación DT'])
        self.assertEqual(cell.value, note)
        self.assertEqual(cell.data_type, 's')
        self.assertTrue(cell.alignment.wrap_text)
        self.assertGreater(sheet.column_dimensions[cell.column_letter].width, 42)
        book.close()

    def test_help_opens_local_guide_inside_plugin(self):
        self.assertEqual(self.owner.main_tabs.currentIndex(), 0)
        self.owner.help_button.click()
        self.assertEqual(self.owner.main_tabs.currentIndex(), 1)
        self.assertEqual(self.owner.main_tabs.tabText(1), 'Ayuda')
        self.assertIn('Validador Consistencia Lógica Alfanumérica', self.owner.help_browser.toPlainText())
        self.assertFalse(self.owner.help_browser.openExternalLinks())
        self.assertFalse(self.owner.help_browser.openLinks())
        self.assertTrue(self.owner.help_browser.isReadOnly())
        self.assertFalse(hasattr(module, 'QDesktopServices'))

    def test_help_missing_guide_stays_inside_plugin(self):
        with patch.object(module.Path, 'read_text', side_effect=OSError('missing')):
            self.owner.load_help_content()
        self.assertIn('Reinstale el ZIP', self.owner.help_browser.toPlainText())

    def test_help_search_and_contents_navigation(self):
        self.owner.help_search.setText('Alcance')
        self.owner.find_help_text()
        self.assertEqual(self.owner.help_browser.textCursor().selectedText().casefold(), 'alcance')
        self.owner.navigate_help(module.QUrl('#8-interpretar-los-resultados'))
        self.assertIn('Interpretar los resultados', self.owner.help_browser.textCursor().selectedText())

    def test_release_buttons_are_not_present_in_details_or_references(self):
        self.assertNotIn('Liberar', [button.text() for button in self.owner.findChildren(QPushButton)])
        self.assertFalse(hasattr(self.owner, 'release_reference_button'))
        for load, status in self.owner.detail_review_rows.values():
            self.assertIn(load.text(), ('Cargar', 'Recargar'))
            self.assertIn('cargado', status.text().casefold())

    def test_consolidated_review_controls_and_save_prompt_are_removed(self):
        labels = [button.text() for button in self.owner.findChildren(QPushButton)]
        self.assertNotIn('Guardar revisión en Excel', labels)
        self.assertNotIn('Cargar revisión consolidada', labels)
        self.assertFalse(hasattr(self.owner, 'save_review_excel'))
        self.assertFalse(hasattr(self.owner, 'recover_review'))
        self.owner.review_dirty = True
        with patch.object(module.QMessageBox, 'question', return_value=QMessageBox.Cancel) as question:
            self.assertFalse(self.owner.confirm_review_changes())
        buttons = question.call_args.args[3]
        self.assertFalse(buttons & QMessageBox.Save)
        self.assertTrue(buttons & QMessageBox.Cancel)

    def test_statistical_summary_cannot_be_loaded_as_review_detail(self):
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('resumen_control_calidad_*.xlsx'))
        previous = [case.exception_tracking() for case in self.owner.review_cases]
        with patch.object(module.QFileDialog, 'getOpenFileName', return_value=(str(path), 'Excel')), patch.object(module.QMessageBox, 'warning') as warning:
            self.owner.recover_detail('Administrativas')
        warning.assert_called_once()
        self.assertEqual(self.owner.review_imports.sources, {})
        self.assertEqual(previous, [case.exception_tracking() for case in self.owner.review_cases])

    def test_reselect_same_model_and_switch_model_clear_environment_and_sql(self):
        with patch.object(self.owner, 'open_sql_version_dialog') as chooser:
            self.owner.environment.setCurrentText(module.SNC_ENVIRONMENT)
            self.owner.selected_sql_version = {'codigo': 'v1'}
            self.owner.model.activated.emit(self.owner.model.currentIndex())
            self.assertEqual(self.owner.environment.currentIndex(), 0)
            self.assertIsNone(self.owner.selected_sql_version)
            self.assertFalse(self.owner.choose_sql_version.isEnabled())
            chooser.reset_mock()
            self.owner.model.setCurrentText('CICA')
            chooser.assert_not_called()
        self.assertEqual(self.owner.sql_version_label.text(), 'Seleccione una versión de SQL')

    def test_discard_restores_dt_and_operator_fields(self):
        self.owner.review_cases[0].excepcion_dt = 'Sí'
        self.owner.review_cases[0].observacion_dt = 'Importado'
        self.owner.review_dirty = True
        with patch.object(module.QMessageBox, 'question', return_value=QMessageBox.Cancel):
            event = QCloseEvent(); self.owner.closeEvent(event)
            self.assertFalse(event.isAccepted())
        with patch.object(module.QMessageBox, 'question', return_value=QMessageBox.Discard):
            event = QCloseEvent(); self.owner.closeEvent(event)
            self.assertTrue(event.isAccepted())
        self.assertEqual(self.owner.review_cases[0].exception_tracking(), ('No', '', 'No', ''))

    @patch.object(module.psycopg2, 'connect')
    def test_worker_captures_sql_fingerprint_and_preserves_duplicate_columns(self, connect):
        cursor = MagicMock()
        connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value = cursor
        cursor.description = [SimpleNamespace(name='id'), SimpleNamespace(name='id')]
        cursor.fetchall.return_value = [(1, 2)]
        worker = module.RuleWorker({}, 'test', '', '', [('Administrativas', 'r1', 'SELECT * FROM {esquema}predio', '')])
        captured = []; worker.result.connect(captured.append); worker.run()
        self.assertEqual(captured[0].rows, [{'id': 1, 'id__2': 2}])
        self.assertFalse(captured[0].identity_safe)
        self.assertEqual(captured[0].criterion_hash, digest('SELECT * FROM test.predio'))

    def test_predio_aliases_export_as_one_column_and_restore_with_original_identity(self):
        from qgis_control_calidad.review import load_review, restore_review
        npn = '000000000000000000000000000001'
        rows = [{'numero_predial_nacional': npn, 'numero_predial': npn, 'Predio': '123', 'valor': 'A'}]
        result = module.RuleResult('unificada', 'Administrativas', 'Correcta', 1, rows=rows, criterion_hash=digest('criterio'))
        self.owner.results = [result]
        self.owner.review_cases = build_cases([result], self.owner.review_context)
        case = self.owner.review_cases[0]
        original_key = case.key
        case.excepcion_dt = 'Sí'; case.observacion_dt = 'Justificada'
        self.owner.review_snapshot = [case.tracking() + case.exception_tracking()]
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path)
        headers = [cell.value for cell in book['unificada'][4]]
        self.assertEqual(headers.count('Numero_Predial_Nacional'), 1)
        self.assertNotIn('Predio', headers)
        self.assertNotIn('numero_predial', headers)
        self.assertNotIn('numero_predial_nacional', headers)
        self.assertIn('valor', headers)
        visible = [cell.value for cell in book['unificada'][4] if not book['unificada'].column_dimensions[cell.column_letter].hidden]
        self.assertEqual(visible, ['Numero_Predial_Nacional', 'valor', *HEADERS])
        self.assertEqual(book['unificada'].cell(5, 1).value, npn)
        self.assertEqual(book['unificada'].cell(5, 1).data_type, 's')
        review_headers = [cell.value for cell in book['Revisión'][1]]
        self.assertIn('Numero_Predial_Nacional', review_headers)
        self.assertNotIn('dato:Predio', review_headers)
        self.assertNotIn('dato:numero_predial_nacional', review_headers)
        book.close()
        current = build_cases([result], self.owner.review_context)
        self.assertEqual(current[0].key, original_key)
        counts = restore_review(current, load_review(path))
        self.assertEqual(counts['incorporadas'], 1)
        self.assertEqual(current[0].excepcion_dt, 'Sí')
        self.assertEqual(current[0].observacion_dt, 'Justificada')


    def test_rule_tabs_use_declared_identifiers_and_natural_order_per_component(self):
        codes = ['regla_1_52', 'regla_1_10', 'regla_1_2', 'regla_1_9', 'regla_1_1']
        self.owner.results = [module.RuleResult(code, 'Administrativas', 'Correcta', 1,
                             rows=[{'numero_predial': '00001'}], criterion_hash=digest(code)) for code in codes]
        self.owner.results.append(module.RuleResult('regla_1_1', 'Juridicas', 'Correcta', 1,
                                  rows=[{'numero_predial': '00002'}], criterion_hash=digest('juridica')))
        self.owner.review_cases = build_cases(self.owner.results, self.owner.review_context)
        self.owner.finished()
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path)
        expected = ['1_1', '1_2', '1_9', '1_10', '1_52']
        self.assertEqual([name for name in book.sheetnames if name in expected], expected)
        index = book['Detalle Validación']
        header = next(row[0].row for row in index if row[0].value == 'Componente')
        self.assertEqual([index.cell(row, 4).value for row in range(header + 1, header + 6)], ['↪ Ir a ' + name for name in expected])
        link = index.cell(header + 5, 4)
        self.assertEqual(link.hyperlink.location, "'1_52'!A1")
        self.assertIsNone(link.hyperlink.target)
        self.assertEqual(link.font.underline, 'single')
        self.assertEqual(link.font.color.rgb, '000066B3')
        self.assertEqual(index.freeze_panes, 'B' + str(header + 1))
        self.assertEqual(index.auto_filter.ref, f'A{header}:D{header + 5}')
        metadata = {row[0].value: row[1].value for row in index.iter_rows(min_row=2, max_row=header-1)}
        self.assertEqual(metadata['Modelo'], 'CICA')
        self.assertEqual(metadata['Esquema'], 'test')
        self.assertNotIn('Identificador de entrega', metadata)
        self.assertNotIn('1_3', book.sheetnames)
        book.close()
        other = next(Path(self.directory.name).glob('detalle_Juridicas_*.xlsx'))
        book = load_workbook(other)
        self.assertIn('1_1', book.sheetnames)
        self.assertNotIn('1_1_2', book.sheetnames)
        book.close()


    def test_rule_details_show_sql_columns_in_order_and_detect_changed_attributes(self):
        from qgis_control_calidad.review import load_review, restore_review
        npn = '000000000000000000000000000001'
        data = {'matricula': '00123', 'numero_predial_nacional': npn, 'cantidad': 2, 'cond_1_12': 'Fail'}
        result = module.RuleResult('regla_1_12', 'Administrativas', 'Correcta', 1,
                                   rows=[data], criterion_hash=digest('criterio completo'))
        self.owner.results = [result]
        self.owner.review_cases = build_cases([result], self.owner.review_context)
        self.owner.review_cases[0].excepcion_dt = 'Sí'
        self.owner.review_cases[0].observacion_dt = 'Justificación'
        self.owner.finished()
        with patch.object(module.QFileDialog, 'getExistingDirectory', return_value=self.directory.name):
            self.owner.export_excel()
        path = next(Path(self.directory.name).glob('detalle_Administrativas_*.xlsx'))
        book = load_workbook(path); sheet = book['1_12']
        visible = [cell.value for cell in sheet[4] if not sheet.column_dimensions[cell.column_letter].hidden]
        self.assertEqual(visible, ['matricula', 'Numero_Predial_Nacional', 'cantidad', 'cond_1_12', *HEADERS])
        self.assertEqual([sheet.cell(5, column).value for column in range(1, 5)], ['00123', npn, '2', 'Fail'])
        book.close()
        self.assertEqual(restore_review(build_cases([result], self.owner.review_context), load_review(path))['incorporadas'], 1)
        book = load_workbook(path); book['1_12'].cell(5, 3, '99'); book.save(path); book.close()
        self.assertEqual(restore_review(build_cases([result], self.owner.review_context), load_review(path))['datos_invalidos'], 1)


if __name__ == '__main__':
    unittest.main()
