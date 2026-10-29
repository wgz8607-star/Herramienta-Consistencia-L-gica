from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from openpyxl import load_workbook
from psycopg2 import sql
from qgis.core import QgsApplication
from qgis.PyQt.QtWidgets import QMessageBox
from cargador_reglas.repositorio import (
    COLUMNS, create_template, read_rules, make_plan, insert_import, RuleUpload,
)
from cargador_reglas import servicio as ui
from qgis_control_calidad.dialog import ControlCalidadDialog

app = QgsApplication.instance() or QgsApplication([], False)
app.initQgis()
TABLE = 'reglas_cargue_snc_cica'


def values(code='001', version='V1_SNC', active=False, published=False, description='Versión de prueba'):
    return (version, 'Cargue SNC', code, 'CONTROL',
            "SELECT numero_predial, 'Sin NPN' AS observacion\nFROM {esquema}.cica_predio\nWHERE numero_predial IS NULL;",
            'Verifica NPN', active, 100, description, published)


def render(query):
    if isinstance(query, str):
        return query
    if isinstance(query, sql.SQL):
        return query.string
    if isinstance(query, sql.Identifier):
        return '.'.join('"' + value.replace('"', '""') + '"' for value in query.strings)
    if isinstance(query, sql.Placeholder):
        return '%s'
    return ''.join(render(part) for part in query)


class RepositoryCursor:
    def __init__(self, connection):
        self.connection = connection
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def execute(self, query, params=None):
        self.connection.events.append(render(query))

    def fetchall(self):
        return list(self.connection.records)

    def executemany(self, query, records):
        self.connection.events.append(render(query))
        self.rowcount = 0
        for index, record in enumerate(records):
            self.connection.pending.append(record)
            self.rowcount += 1
            if self.connection.fail_insert and index == 0:
                raise RuntimeError('Fallo de inserción')


class RepositoryConnection:
    def __init__(self, records=(), fail_insert=False):
        self.records = list(records); self.pending = []
        self.fail_insert = fail_insert; self.events = []; self.closed = False

    def set_session(self, **settings):
        self.events.append(settings)

    def cursor(self):
        return RepositoryCursor(self)

    def __enter__(self):
        return self

    def __exit__(self, kind, value, traceback):
        if kind is None:
            self.records.extend(self.pending); self.events.append('COMMIT')
        else:
            self.events.append('ROLLBACK')
        self.pending.clear()

    def close(self):
        self.closed = True


class RuleUploadTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'reglas.xlsx'

    def test_template_is_blank_with_all_columns_and_separate_example(self):
        create_template(self.path)
        book = load_workbook(self.path)
        self.assertEqual(tuple(cell.value for cell in book['Reglas'][1]), COLUMNS)
        self.assertEqual(book['Reglas']['C2'].number_format, '@')
        self.assertIn('Ejemplo', book.sheetnames)
        book.close()
        with self.assertRaisesRegex(ValueError, 'vacía'):
            read_rules(self.path)

    def test_roundtrip_exports_all_rows_flags_and_multiline_sql(self):
        records = [values('001'), values('002', 'V2_SNC', active=True, published=True)]
        create_template(self.path, records)
        parsed = read_rules(self.path)
        self.assertEqual([r.values for r in parsed], records)
        self.assertEqual(parsed[0].values[2], '001')
        self.assertIn('\nFROM {esquema}', parsed[0].values[4])

    def test_export_does_not_silently_truncate_long_sql(self):
        record = list(values()); record[4] = 'SELECT ' + 'x' * 32768
        with self.assertRaisesRegex(ValueError, 'evitar truncar'):
            create_template(self.path, [record])
        self.assertFalse(self.path.exists())

    def test_empty_description_and_metadata_are_supported_on_existing_rows(self):
        record = list(values()); record[5] = None; record[8] = None; record[9] = None
        create_template(self.path, [record])
        parsed = read_rules(self.path)[0]
        self.assertEqual(parsed.values[5], '')
        self.assertIsNone(parsed.values[9])

    def test_invalid_rows_and_formulas_are_rejected_as_a_whole(self):
        create_template(self.path, [values(), values('002')])
        book = load_workbook(self.path); book['Reglas']['E3'] = '=1+1'; book.save(self.path); book.close()
        with self.assertRaisesRegex(ValueError, 'Fila 3.*fórmulas'):
            read_rules(self.path)
        create_template(self.path, [values()])
        book = load_workbook(self.path); book['Reglas']['G2'] = 'Quizá'; book.save(self.path); book.close()
        with self.assertRaisesRegex(ValueError, 'Sí o No'):
            read_rules(self.path)

    def test_duplicate_codes_in_same_version_and_component_are_rejected(self):
        create_template(self.path, [values(), values()])
        with self.assertRaisesRegex(ValueError, 'duplicada'):
            read_rules(self.path)

    def test_existing_rules_are_preserved_and_new_rules_keep_header_metadata(self):
        existing = [values(active=True, published=True, description='Publicación original')]
        rules = [RuleUpload(2, values()), RuleUpload(3, values('002', description='No modificar'))]
        plan = make_plan(TABLE, rules, existing)
        self.assertEqual(len(plan.inserts), 1)
        self.assertIn('diferencias', plan.states[0])
        self.assertEqual(plan.inserts[0][8:], ('Publicación original', True))
        connection = RepositoryConnection(existing)
        with connection:
            self.assertEqual(insert_import(connection, plan), 1)
        self.assertEqual(connection.records[0], existing[0])
        self.assertEqual(len(connection.records), 2)
        self.assertTrue(any(query.startswith('INSERT INTO') for query in connection.events if isinstance(query, str)))
        self.assertFalse(any(query.startswith('UPDATE') or query.startswith('DELETE') for query in connection.events if isinstance(query, str)))

    def test_new_version_defaults_to_unpublished(self):
        record = list(values()); record[8:] = [None, None]
        plan = make_plan(TABLE, [RuleUpload(2, tuple(record))], [])
        self.assertEqual(plan.inserts[0][8:], ('', False))

    def test_metadata_conflicts_and_two_published_snc_versions_are_rejected(self):
        with self.assertRaisesRegex(ValueError, 'metadatos'):
            make_plan(TABLE, [RuleUpload(2, values()), RuleUpload(3, values('002', description='Otra'))], [])
        with self.assertRaisesRegex(ValueError, 'una sola versión'):
            make_plan(TABLE, [RuleUpload(2, values('002', 'V2_SNC', active=True, published=True))], [values(active=True, published=True)])

    def test_date_publication_is_allowed_in_v2_table(self):
        for version in ('29_10_2026', '2026-10-29', 'V2_29_10_2026'):
            with self.subTest(version=version):
                plan = make_plan('reglas_interno_v2_29_07_2024',
                                 [RuleUpload(2, values(version=version))], [])
                self.assertEqual(len(plan.inserts), 1)
                self.assertEqual(plan.inserts[0][0], version)

    def test_target_version_and_table_are_checked(self):
        with self.assertRaisesRegex(ValueError, 'destino V2'):
            make_plan('reglas_cica_v2_29_07_2024', [RuleUpload(2, values(version='V3_07_09_2026'))], [])
        with self.assertRaisesRegex(ValueError, 'admitida'):
            make_plan('tabla; DROP TABLE otra', [RuleUpload(2, values())], [])

    def test_repository_change_after_preview_rejects_write(self):
        plan = make_plan(TABLE, [RuleUpload(2, values())], [])
        connection = RepositoryConnection([values('otro')])
        with self.assertRaisesRegex(ValueError, 'cambió'):
            with connection:
                insert_import(connection, plan)
        self.assertEqual(len(connection.records), 1)
        self.assertIn('ROLLBACK', connection.events)

    def test_apply_is_atomic_and_does_not_execute_rule_sql(self):
        rules = [RuleUpload(2, values()), RuleUpload(3, values('002'))]
        plan = make_plan(TABLE, rules, [])
        connection = RepositoryConnection(fail_insert=True)
        with patch.object(ui.psycopg2, 'connect', return_value=connection):
            with self.assertRaisesRegex(RuntimeError, 'inserción'):
                ui.apply_plan({}, plan)
        self.assertEqual(connection.records, [])
        self.assertEqual(connection.pending, [])
        self.assertTrue(connection.closed)
        self.assertNotIn(values()[4], connection.events)
        self.assertIn('ROLLBACK', connection.events)

    def test_export_existing_reads_without_filtering_active_or_published(self):
        records = [values(), values('002', active=True, published=True)]
        connection = RepositoryConnection(records)
        with patch.object(ui.psycopg2, 'connect', return_value=connection):
            self.assertEqual(ui.export_existing({}, TABLE, self.path), 2)
        self.assertEqual([rule.values for rule in read_rules(self.path)], records)
        select = next(event for event in connection.events if isinstance(event, str) and event.startswith('SELECT'))
        self.assertNotIn('WHERE', select)
        self.assertIn({'readonly': True}, connection.events)
        self.assertTrue(connection.closed)

    def test_plugin_has_no_administration_tab(self):
        owner = ControlCalidadDialog(); self.addCleanup(owner.close)
        self.assertEqual([owner.main_tabs.tabText(index) for index in range(owner.main_tabs.count())], ['Validación', 'Ayuda'])
        self.assertFalse(hasattr(owner, 'rule_upload_page'))
