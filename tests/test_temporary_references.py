import csv
from pathlib import Path
import os
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from psycopg2 import sql
from qgis.core import QgsApplication
from qgis.PyQt.QtWidgets import QDialog
from qgis_control_calidad.reference_data import prepare_reference, install_reference, REQUIRED
from qgis_control_calidad.reference import ReferenceDialog
from qgis_control_calidad.dialog import RuleWorker, ControlCalidadDialog

app = QgsApplication.instance() or QgsApplication([], False)
app.initQgis()


def render(statement):
    if isinstance(statement, str):
        return statement
    if isinstance(statement, sql.SQL):
        return statement.string
    if isinstance(statement, sql.Identifier):
        return '.'.join('"' + s.replace('"', '""') + '"' for s in statement.strings)
    return ''.join(render(part) for part in statement)


class FakeCursor:
    def __init__(self, connection):
        self.connection = connection
        self.description = None

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def execute(self, statement):
        query = render(statement)
        self.connection.events.append(('sql', query))
        if query.startswith('CREATE TEMP TABLE'):
            self.connection.tables.add(query.split('"')[1])
        elif query == 'FAIL_RULE':
            raise RuntimeError('Fallo de una regla')
        elif query.startswith('SELECT'):
            if 'pg_temp.reference_r1' in query:
                assert 'reference_r1' in self.connection.tables
            if 'pg_temp.reference_r2' in query:
                assert 'reference_r2' in self.connection.tables
            self.description = [SimpleNamespace(name='numero_predial')]

    def copy_expert(self, statement, stream):
        self.connection.events.append(('copy', statement, list(csv.reader(stream))))

    def fetchall(self):
        return [('00001',)]


class FakeConnection:
    def __init__(self):
        self.tables = set()
        self.retained = set()
        self.events = []
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, exception, value, traceback):
        self.rollback() if exception else self.commit()

    def cursor(self):
        return FakeCursor(self)

    def commit(self):
        self.events.append(('commit',))
        self.retained = set(self.tables)

    def rollback(self):
        self.events.append(('rollback',))
        self.tables = set(self.retained)

    def close(self):
        self.events.append(('close',))
        self.tables.clear()
        self.closed = True


class TemporaryReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def excel(self, kind, name=None):
        path = Path(self.temp.name) / (name or kind + '.xlsx')
        book = Workbook(); sheet = book.active
        headers = sorted(REQUIRED[kind])
        sheet.append(headers)
        values = ['00001' if h == 'numero_predial' else '' if h == 'destino_economico' else '1' for h in headers]
        sheet.append(values)
        book.save(path); book.close()
        return path

    def prepared(self, kind):
        reference = prepare_reference(kind, self.excel(kind))
        self.addCleanup(reference.close)
        return reference

    def test_prepare_is_local_snapshot_and_preserves_zeroes_and_empty_strings(self):
        reference = self.prepared('R1')
        with reference.csv_path.open(newline='', encoding='utf-8') as stream:
            row = next(csv.reader(stream))
        self.assertEqual(row[reference.columns.index('numero_predial')], '00001')
        self.assertEqual(row[reference.columns.index('destino_economico')], '')
        self.assertTrue(reference.csv_path.read_text().startswith('"'))
        same = prepare_reference('R1', self.excel('R1', 'otro_nombre.xlsx'))
        self.addCleanup(same.close)
        self.assertEqual(same.fingerprint, reference.fingerprint)
        original = reference.csv_path
        reference.close()
        self.assertFalse(original.exists())

    def test_missing_and_duplicate_headers_are_rejected(self):
        path = self.excel('R1')
        book = Workbook(); book.active.append(['NUMERO_PREDIAL', 'numero_predial']); book.save(path); book.close()
        with self.assertRaisesRegex(ValueError, 'duplicados'):
            prepare_reference('R1', path)

    def test_resolve_all_reference_marker_forms_before_source_schema(self):
        refs = {'R1': self.prepared('R1'), 'R2': self.prepared('R2')}
        worker = RuleWorker({}, 'municipio', '', '', [], references=refs)
        for kind in ('r1', 'r2'):
            for pattern in ('{esquema_%s}.{tabla_%s}', '{esquema_%s}{punto}{tabla_%s}', '{esquema_%s}{tabla_%s}'):
                source = pattern % (kind, kind)
                self.assertEqual(worker.resolve_sql('SELECT * FROM ' + source), 'SELECT * FROM pg_temp.reference_' + kind)
        self.assertEqual(worker.resolve_sql('SELECT * FROM {esquema}{punto}ilc_predio'), 'SELECT * FROM municipio.ilc_predio')
        with self.assertRaises(ValueError):
            RuleWorker({}, 'municipio', '', '', []).resolve_sql('SELECT * FROM {esquema_r1}.{tabla_r1}')
        self.assertEqual(RuleWorker({}, 'municipio', '', 'r1_publicado', []).resolve_sql('SELECT * FROM {esquema_r1}.{tabla_r1}'), 'SELECT * FROM public.r1_publicado')

    def test_worker_uses_one_connection_temporaries_survive_rule_failure_and_close(self):
        refs = {'R1': self.prepared('R1'), 'R2': self.prepared('R2')}
        connection = FakeConnection()
        worker = RuleWorker({}, 'test', '', '', [
            ('Prueba', 'fallida', 'FAIL_RULE', ''),
            ('Prueba', 'r1', 'SELECT * FROM {esquema_r1}{punto}{tabla_r1}', ''),
            ('Prueba', 'r2', 'SELECT * FROM {esquema_r2}.{tabla_r2}', ''),
        ], references=refs)
        results, errors = [], []
        worker.result.connect(results.append); worker.fatal_error.connect(errors.append)
        with patch('qgis_control_calidad.dialog.psycopg2.connect', return_value=connection) as connect, patch.object(sql.Composed, 'as_string', lambda self, conn: render(self)):
            worker.run()
        connect.assert_called_once()
        self.assertEqual([r.status for r in results], ['Error', 'Correcta', 'Correcta'])
        self.assertEqual(errors, [])
        self.assertTrue(connection.closed)
        self.assertFalse(connection.tables)
        statements = [entry[1] for entry in connection.events if entry[0] == 'sql']
        self.assertEqual(sum(q.startswith('CREATE TEMP TABLE') for q in statements), 2)
        self.assertTrue(all('public.' not in q and 'DROP TABLE' not in q for q in statements))
        self.assertLess(connection.events.index(('commit',)), connection.events.index(('sql', 'FAIL_RULE')))

    def test_setup_failure_closes_connection_and_skips_rules(self):
        connection = FakeConnection()
        worker = RuleWorker({}, 'test', '', '', [('Prueba', 'r1', 'SELECT 1', '')], references={'R1': self.prepared('R1')})
        errors, results = [], []
        worker.fatal_error.connect(errors.append); worker.result.connect(results.append)
        with patch('qgis_control_calidad.dialog.psycopg2.connect', return_value=connection), patch('qgis_control_calidad.dialog.install_reference', side_effect=RuntimeError('TEMP denied')):
            worker.run()
        self.assertTrue(connection.closed)
        self.assertEqual(results, [])
        self.assertEqual(errors, ['TEMP denied'])

    def test_bottom_section_has_three_modes_and_future_mode_is_disabled(self):
        widget = ControlCalidadDialog()
        self.addCleanup(widget.close)
        self.assertEqual(widget.reference_mode.count(), 3)
        self.assertFalse(widget.reference_mode.model().item(2).isEnabled())
        widget.reference_mode.setCurrentIndex(1)
        self.assertEqual(widget.reference_pages.currentIndex(), 1)
        widget.reference_mode.setCurrentIndex(2)
        self.assertEqual(widget.reference_mode.currentIndex(), 0)
        self.assertFalse(hasattr(widget, 'published_r2'))

    def test_prepare_dialog_reads_both_files_without_database_access(self):
        dialog = ReferenceDialog()
        dialog.r1.setText(str(self.excel('R1')))
        dialog.r2.setText(str(self.excel('R2')))
        with patch('qgis_control_calidad.dialog.psycopg2.connect') as connect:
            dialog.validate()
        connect.assert_not_called()
        self.assertTrue(dialog.accept_button.isEnabled())
        self.assertEqual(set(dialog.references), {'R1', 'R2'})
        paths = [ref.csv_path for ref in dialog.references.values()]
        dialog.done(QDialog.Rejected)
        self.assertTrue(all(not path.exists() for path in paths))

    def test_dialog_cancel_removes_local_snapshots(self):
        dialog = ReferenceDialog()
        reference = self.prepared('R1')
        dialog.references = {'R1': reference}
        path = reference.csv_path
        dialog.done(QDialog.Rejected)
        self.assertFalse(path.exists())

    def test_existing_public_mode_does_not_install_tables(self):
        connection = FakeConnection()
        worker = RuleWorker({}, 'test', '', 'r1_existente', [('Prueba', 'r1', 'SELECT * FROM {esquema_r1}.{tabla_r1}', '')])
        with patch('qgis_control_calidad.dialog.psycopg2.connect', return_value=connection), patch('qgis_control_calidad.dialog.install_reference') as install:
            worker.run()
        install.assert_not_called()
        self.assertIn(('sql', 'SELECT * FROM public.r1_existente'), connection.events)
        self.assertTrue(connection.closed)


if __name__ == '__main__':
    unittest.main()
